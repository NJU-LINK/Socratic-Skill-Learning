from __future__ import annotations

from pathlib import Path

import openpyxl


answer_sheet = "master"
range_str = "A1:AH1"
style_mutable_ranges = ["master!A1:AH1"]
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _is_yellow(cell) -> bool:
    fill = cell.fill
    return (
        fill.fill_type == "solid"
        and fill.fgColor.type == "rgb"
        and str(fill.fgColor.rgb).upper().endswith("FFFF00")
    )


def _expected_values() -> list[object]:
    source = openpyxl.load_workbook(INPUT, data_only=False)["Test"]
    selected = [
        source.cell(row, 1).value
        for row in range(1, source.max_row + 1)
        if _is_yellow(source.cell(row, 1))
        and source.cell(row, 1).value not in (None, "")
    ]
    return selected + [None] * (34 - len(selected))


EXPECTED = _expected_values()


def test_style_selected_values_are_compacted_and_transposed() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False)["master"]
    assert [sheet.cell(1, column).value for column in range(1, 35)] == EXPECTED
    assert all(sheet.cell(1, column).data_type != "f" for column in range(1, 35))


def test_populated_output_cells_have_required_style() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False)["master"]
    populated = sum(value is not None for value in EXPECTED)
    for column in range(1, populated + 1):
        cell = sheet.cell(1, column)
        assert _is_yellow(cell)
        assert (
            cell.border.left.style,
            cell.border.right.style,
            cell.border.top.style,
            cell.border.bottom.style,
        ) == ("thin", "thin", "thin", "thin")


def test_unused_target_cells_are_unformatted_blanks() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False)["master"]
    populated = sum(value is not None for value in EXPECTED)
    for column in range(populated + 1, 35):
        cell = sheet.cell(1, column)
        assert cell.value is None
        assert cell.fill.fill_type is None
        assert all(
            side is None or side.style is None
            for side in (
                cell.border.left,
                cell.border.right,
                cell.border.top,
                cell.border.bottom,
            )
        )
