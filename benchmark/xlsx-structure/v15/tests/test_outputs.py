from __future__ import annotations

from copy import copy
from pathlib import Path

import openpyxl


answer_sheet = "MyResult"
range_str = "A1:O20"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _source_contract() -> tuple[list[list[object]], set[object], dict[object, int]]:
    workbook = openpyxl.load_workbook(INPUT, data_only=False)
    ids_sheet = workbook["Main unique ID"]
    source = workbook["Result what i am getting"]
    desired_ids = [
        ids_sheet.cell(row, 1).value
        for row in range(1, ids_sheet.max_row + 1)
        if ids_sheet.cell(row, 1).value not in (None, "")
    ]
    source_rows = {
        source.cell(row, 1).value: row
        for row in range(2, source.max_row + 1)
        if source.cell(row, 1).value not in (None, "")
    }
    expected = [[source.cell(1, column).value for column in range(1, 16)]]
    missing = set()
    for identifier in desired_ids:
        if identifier in source_rows:
            row = source_rows[identifier]
            expected.append(
                [source.cell(row, column).value for column in range(1, 16)]
            )
        else:
            expected.append([identifier] + [None] * 14)
            missing.add(identifier)
    return expected, missing, source_rows


EXPECTED, MISSING_IDS, SOURCE_ROWS = _source_contract()


def _style_without_fill_equal(left, right) -> bool:
    return (
        copy(left.font) == copy(right.font)
        and copy(left.border) == copy(right.border)
        and copy(left.alignment) == copy(right.alignment)
        and left.number_format == right.number_format
        and copy(left.protection) == copy(right.protection)
    )


def _is_yellow(cell) -> bool:
    return (
        cell.fill.fill_type == "solid"
        and cell.fill.fgColor.type == "rgb"
        and str(cell.fill.fgColor.rgb).upper().endswith("FFFF00")
    )


def test_reconciled_table_values_and_order() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False)["MyResult"]
    actual = [
        [sheet.cell(row, column).value for column in range(1, 16)]
        for row in range(1, 21)
    ]
    assert actual == EXPECTED
    assert all(
        sheet.cell(row, column).data_type != "f"
        for row in range(1, 21)
        for column in range(1, 16)
    )


def test_existing_rows_and_header_keep_source_formatting() -> None:
    source = openpyxl.load_workbook(INPUT, data_only=False)[
        "Result what i am getting"
    ]
    result = openpyxl.load_workbook(OUTPUT, data_only=False)["MyResult"]
    for column in range(1, 16):
        assert source.cell(1, column)._style == result.cell(1, column)._style
    for output_row in range(2, 21):
        identifier = result.cell(output_row, 1).value
        if identifier in MISSING_IDS:
            continue
        source_row = SOURCE_ROWS[identifier]
        for column in range(1, 16):
            assert (
                source.cell(source_row, column)._style
                == result.cell(output_row, column)._style
            )


def test_inserted_rows_use_base_formatting_and_yellow_fill() -> None:
    source = openpyxl.load_workbook(INPUT, data_only=False)[
        "Result what i am getting"
    ]
    result = openpyxl.load_workbook(OUTPUT, data_only=False)["MyResult"]
    for row in range(2, 21):
        if result.cell(row, 1).value not in MISSING_IDS:
            continue
        for column in range(1, 16):
            cell = result.cell(row, column)
            assert _style_without_fill_equal(cell, source.cell(2, column))
            assert _is_yellow(cell)


def test_result_column_widths_match_source() -> None:
    source = openpyxl.load_workbook(INPUT, data_only=False)[
        "Result what i am getting"
    ]
    result = openpyxl.load_workbook(OUTPUT, data_only=False)["MyResult"]
    for column in range(1, 16):
        letter = openpyxl.utils.get_column_letter(column)
        assert (
            result.column_dimensions[letter].width
            == source.column_dimensions[letter].width
        )
