from __future__ import annotations

import ast
import json
import math
import os
from copy import copy
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter, range_boundaries


TEST_ROOT = Path(os.environ.get("TEST_DIR", "/tests"))
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")
MAX_VALUE_GROUPS = 8


def _literal_assignment(tree: ast.AST, name: str) -> str:
    values: list[str] = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Assign) or not isinstance(node.value, ast.Constant):
            continue
        if not isinstance(node.value.value, str):
            continue
        if any(isinstance(target, ast.Name) and target.id == name for target in node.targets):
            values.append(node.value.value)
    if len(values) != 1:
        raise AssertionError(f"Expected exactly one literal {name} assignment, found {values!r}")
    return values[0]


def _literal_string_list_assignment(tree: ast.AST, name: str) -> list[str]:
    values: list[list[str]] = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Assign):
            continue
        if not any(isinstance(target, ast.Name) and target.id == name for target in node.targets):
            continue
        value = ast.literal_eval(node.value)
        if not isinstance(value, list) or not all(isinstance(item, str) for item in value):
            raise AssertionError(f"{name} must be a literal list of strings")
        values.append(value)
    if len(values) > 1:
        raise AssertionError(f"Expected at most one {name} assignment")
    return values[0] if values else []


def _style_mutable_cells(tree: ast.AST) -> set[tuple[str, str]]:
    mutable: set[tuple[str, str]] = set()
    for entry in _literal_string_list_assignment(tree, "style_mutable_ranges"):
        if "!" not in entry:
            raise AssertionError(
                f"style_mutable_ranges entry must be SHEET!RANGE, got {entry!r}"
            )
        sheet_name, cell_range = entry.split("!", 1)
        min_column, min_row, max_column, max_row = range_boundaries(cell_range)
        mutable.update(
            (sheet_name, f"{get_column_letter(column)}{row}")
            for row in range(min_row, max_row + 1)
            for column in range(min_column, max_column + 1)
        )
    return mutable


def _contract() -> tuple[str, str, list[list[tuple[str, object]]]]:
    tree = ast.parse((TEST_ROOT / "test_outputs.py").read_text(encoding="utf-8"))
    sheet_name = _literal_assignment(tree, "answer_sheet")
    answer_range = _literal_assignment(tree, "range_str")
    expected = json.loads((TEST_ROOT / "expected.json").read_text(encoding="utf-8"))["values"]
    min_column, min_row, max_column, max_row = range_boundaries(answer_range)
    coordinates = [
        f"{get_column_letter(column)}{row}"
        for row in range(min_row, max_row + 1)
        for column in range(min_column, max_column + 1)
    ]
    if len(coordinates) != len(expected):
        raise AssertionError(
            f"{answer_range} has {len(coordinates)} cells but expected.json has {len(expected)} values"
        )
    cases = list(zip(coordinates, expected))
    chunk_size = max(1, math.ceil(len(cases) / MAX_VALUE_GROUPS))
    groups = [cases[index:index + chunk_size] for index in range(0, len(cases), chunk_size)]
    return sheet_name, answer_range, groups


def _value_equal(actual: object, expected: object) -> bool:
    if actual == expected:
        return True
    if (actual is None and expected == "") or (actual == "" and expected is None):
        return True
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return abs(float(actual) - float(expected)) < 1e-6
    return str(actual) == str(expected)


def _style_equal(before: object, after: object) -> bool:
    return (
        copy(before.font) == copy(after.font)
        and copy(before.fill) == copy(after.fill)
        and copy(before.border) == copy(after.border)
        and copy(before.alignment) == copy(after.alignment)
        and before.number_format == after.number_format
        and copy(before.protection) == copy(after.protection)
    )


ANSWER_SHEET, ANSWER_RANGE, EXPECTED_GROUPS = _contract()
TEST_TREE = ast.parse((TEST_ROOT / "test_outputs.py").read_text(encoding="utf-8"))
STYLE_MUTABLE_CELLS = _style_mutable_cells(TEST_TREE)


@pytest.mark.parametrize(
    "expected_group",
    EXPECTED_GROUPS,
    ids=[f"answer-values-{index + 1:02d}" for index in range(len(EXPECTED_GROUPS))],
)
def test_answer_value_group(expected_group: list[tuple[str, object]]) -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    workbook = openpyxl.load_workbook(OUTPUT, data_only=True)
    sheet_name = workbook.sheetnames[0] if ANSWER_SHEET == "None" else ANSWER_SHEET
    assert sheet_name in workbook.sheetnames
    sheet = workbook[sheet_name]
    mismatches = [
        (coordinate, sheet[coordinate].value, expected)
        for coordinate, expected in expected_group
        if not _value_equal(sheet[coordinate].value, expected)
    ]
    assert not mismatches, f"Answer mismatches: {mismatches[:8]!r}"


def test_input_sheets_and_unrequested_values_are_preserved() -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    source = openpyxl.load_workbook(INPUT, data_only=False)
    result = openpyxl.load_workbook(OUTPUT, data_only=False)
    missing = sorted(set(source.sheetnames) - set(result.sheetnames))
    assert not missing, f"Output removed input sheets: {missing}"

    answer_sheet = source.sheetnames[0] if ANSWER_SHEET == "None" else ANSWER_SHEET
    min_column, min_row, max_column, max_row = range_boundaries(ANSWER_RANGE)
    changes: list[tuple[str, str, object, object]] = []
    for sheet_name in source.sheetnames:
        source_sheet = source[sheet_name]
        result_sheet = result[sheet_name]
        scan_rows = max(source_sheet.max_row, result_sheet.max_row)
        scan_columns = max(source_sheet.max_column, result_sheet.max_column)
        for row in range(1, scan_rows + 1):
            for column in range(1, scan_columns + 1):
                in_answer = (
                    sheet_name == answer_sheet
                    and min_row <= row <= max_row
                    and min_column <= column <= max_column
                )
                if in_answer:
                    continue
                before = source_sheet.cell(row, column).value
                after = result_sheet.cell(row, column).value
                if not _value_equal(after, before):
                    changes.append(
                        (sheet_name, f"{get_column_letter(column)}{row}", before, after)
                    )
                    if len(changes) >= 8:
                        break
            if len(changes) >= 8:
                break
        if len(changes) >= 8:
            break
    assert not changes, f"Values outside the declared answer range changed: {changes!r}"


def test_existing_workbook_structure_and_styles_are_preserved() -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    source = openpyxl.load_workbook(INPUT, data_only=False)
    result = openpyxl.load_workbook(OUTPUT, data_only=False)
    problems: list[str] = []

    for sheet_name in source.sheetnames:
        source_sheet = source[sheet_name]
        result_sheet = result[sheet_name]
        if source_sheet.auto_filter.ref != result_sheet.auto_filter.ref:
            problems.append(
                f"{sheet_name}: auto-filter {source_sheet.auto_filter.ref!r} -> "
                f"{result_sheet.auto_filter.ref!r}"
            )
        if source_sheet.freeze_panes != result_sheet.freeze_panes:
            problems.append(
                f"{sheet_name}: freeze panes {source_sheet.freeze_panes!r} -> "
                f"{result_sheet.freeze_panes!r}"
            )
        source_merges = sorted(str(item) for item in source_sheet.merged_cells.ranges)
        result_merges = sorted(str(item) for item in result_sheet.merged_cells.ranges)
        if source_merges != result_merges:
            problems.append(f"{sheet_name}: merged ranges changed")

        for key, dimension in source_sheet.column_dimensions.items():
            other = result_sheet.column_dimensions.get(key)
            if other is None or (
                dimension.width,
                dimension.hidden,
                dimension.outlineLevel,
                dimension.collapsed,
            ) != (
                other.width,
                other.hidden,
                other.outlineLevel,
                other.collapsed,
            ):
                problems.append(f"{sheet_name}: column {key} dimensions changed")
        for key, dimension in source_sheet.row_dimensions.items():
            other = result_sheet.row_dimensions.get(key)
            if other is None or (
                dimension.height,
                dimension.hidden,
                dimension.outlineLevel,
                dimension.collapsed,
            ) != (
                other.height,
                other.hidden,
                other.outlineLevel,
                other.collapsed,
            ):
                problems.append(f"{sheet_name}: row {key} dimensions changed")

        for row in range(1, source_sheet.max_row + 1):
            for column in range(1, source_sheet.max_column + 1):
                before = source_sheet.cell(row, column)
                after = result_sheet.cell(row, column)
                coordinate = f"{get_column_letter(column)}{row}"
                if (sheet_name, coordinate) in STYLE_MUTABLE_CELLS:
                    continue
                if not _style_equal(before, after):
                    problems.append(
                        f"{sheet_name}!{coordinate}: style changed"
                    )
                    if len(problems) >= 8:
                        break
            if len(problems) >= 8:
                break
        if len(problems) >= 8:
            break
    assert not problems, f"Workbook preservation failures: {problems!r}"
