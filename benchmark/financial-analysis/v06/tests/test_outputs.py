from pathlib import Path

from openpyxl import load_workbook


OUTPUT = Path("/app/msft_analysis.xlsx")
HEADER = ["Fiscal Year", "Revenue", "Net_Income", "Net_Margin_Pct"]
EXPECTED = {
    "FY2021": (168088, 61271, 36.45),
    "FY2022": (198270, 72738, 36.69),
    "FY2023": (211915, 72361, 34.15),
    "FY2024": (245122, 88136, 35.96),
    "FY2025": (281724, 101832, 36.15),
}


def output_rows():
    assert OUTPUT.is_file(), f"{OUTPUT} not found"
    workbook = load_workbook(OUTPUT, data_only=True)
    sheet = workbook.active
    assert [sheet.cell(1, column).value for column in range(1, 5)] == HEADER
    rows = [
        tuple(sheet.cell(row, column).value for column in range(1, 5))
        for row in range(2, sheet.max_row + 1)
        if any(sheet.cell(row, column).value is not None for column in range(1, 5))
    ]
    return rows


def indexed_rows():
    indexed = {}
    for year, revenue, net_income, margin in output_rows():
        assert year not in indexed, f"duplicate fiscal year: {year}"
        indexed[year] = (revenue, net_income, margin)
    return indexed


def test_workbook_has_exact_annual_row_set():
    indexed = indexed_rows()
    assert set(indexed) == set(EXPECTED)
    assert len(indexed) == 5


def test_all_source_revenues_are_preserved():
    indexed = indexed_rows()
    for year, (revenue, _, _) in EXPECTED.items():
        assert indexed[year][0] == revenue


def test_all_source_net_incomes_are_preserved():
    indexed = indexed_rows()
    for year, (_, net_income, _) in EXPECTED.items():
        assert indexed[year][1] == net_income


def test_all_net_margins_are_correct():
    indexed = indexed_rows()
    for year, (_, _, expected_margin) in EXPECTED.items():
        assert abs(float(indexed[year][2]) - expected_margin) < 0.011


def test_non_financial_department_rows_are_excluded():
    years = set(indexed_rows())
    assert not years.intersection({"Department", "Engineering", "Sales", "Marketing", "HR"})


def test_output_contains_no_undeclared_fifth_column():
    workbook = load_workbook(OUTPUT, data_only=True)
    sheet = workbook.active
    assert all(
        sheet.cell(row, 5).value is None
        for row in range(1, sheet.max_row + 1)
    )
