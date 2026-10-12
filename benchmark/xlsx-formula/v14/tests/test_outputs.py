from __future__ import annotations

import math
from pathlib import Path

import openpyxl


answer_sheet = "None"
range_str = "E3:E15"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _expected() -> tuple[list[str], list[int]]:
    sheet = openpyxl.load_workbook(INPUT, data_only=False).active
    report_date = sheet["B1"].value
    categories = {
        int(sheet.cell(row, 8).value): str(sheet.cell(row, 9).value)
        for row in range(3, 12)
        if sheet.cell(row, 8).value is not None
    }
    category_values = []
    actual_ages = []
    for row in range(3, 16):
        birth_date = sheet.cell(row, 3).value
        lookup_age = math.floor(((report_date - birth_date).days / 365) + 0.5)
        category_values.append(categories[lookup_age])
        age = report_date.year - birth_date.year
        if (report_date.month, report_date.day) < (birth_date.month, birth_date.day):
            age -= 1
        actual_ages.append(age)
    return category_values, actual_ages


EXPECTED_CATEGORIES, EXPECTED_AGES = _expected()


def test_categories_use_rounded_age_lookup() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=True).active
    assert [sheet.cell(row, 5).value for row in range(3, 16)] == EXPECTED_CATEGORIES


def test_completed_actual_ages_and_header() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False).active
    assert sheet["P2"].value == "Actual Age"
    assert [sheet.cell(row, 16).value for row in range(3, 16)] == EXPECTED_AGES
    assert all(sheet.cell(row, 16).data_type != "f" for row in range(2, 16))


def test_actual_age_formatting() -> None:
    source = openpyxl.load_workbook(INPUT, data_only=False).active
    result = openpyxl.load_workbook(OUTPUT, data_only=False).active

    def fill_signature(cell: openpyxl.cell.cell.Cell) -> tuple[object, ...]:
        fill = cell.fill
        return (
            fill.fill_type,
            fill.fgColor.type,
            fill.fgColor.rgb,
            fill.fgColor.indexed,
            fill.fgColor.theme,
            fill.fgColor.tint,
            fill.bgColor.type,
            fill.bgColor.rgb,
            fill.bgColor.indexed,
            fill.bgColor.theme,
            fill.bgColor.tint,
        )

    for row in range(2, 16):
        target = result.cell(row, 16)
        assert fill_signature(target) == fill_signature(source.cell(row, 15))
        assert target.font.bold is True
        assert target.alignment.horizontal == "center"
        if row == 2:
            assert target.alignment.vertical == "top"
        else:
            assert target.number_format == "0"


def test_existing_cell_formatting_is_preserved() -> None:
    source = openpyxl.load_workbook(INPUT, data_only=False).active
    result = openpyxl.load_workbook(OUTPUT, data_only=False).active
    mismatches = []
    for row in range(1, source.max_row + 1):
        for column in range(1, source.max_column + 1):
            coordinate = source.cell(row, column).coordinate
            if source[coordinate]._style != result[coordinate]._style:
                mismatches.append(coordinate)
    assert not mismatches, f"Existing formatting changed: {mismatches[:10]}"
