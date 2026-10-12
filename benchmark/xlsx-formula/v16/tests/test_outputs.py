from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import openpyxl


answer_sheet = "Sheet1"
range_str = "H2:H18"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _expected() -> list[object]:
    workbook = openpyxl.load_workbook(INPUT, data_only=True)
    orders = workbook["Sheet1"]
    inventory = workbook["Sheet2"]

    stock: dict[tuple[str, str], float] = defaultdict(float)
    for row in range(2, inventory.max_row + 1):
        order_number = inventory.cell(row, 1).value
        item_code = inventory.cell(row, 2).value
        if order_number not in (None, "") and item_code not in (None, ""):
            stock[(str(order_number), str(item_code))] += float(
                inventory.cell(row, 4).value or 0
            )

    records = []
    for row in range(2, orders.max_row + 1):
        records.append(
            (
                str(orders.cell(row, 1).value),
                str(orders.cell(row, 4).value),
                orders.cell(row, 3).value,
                float(orders.cell(row, 2).value),
                row,
                float(orders.cell(row, 7).value or 0),
            )
        )
    records.sort(key=lambda case: (case[0], case[1], case[2], case[3], case[4]))

    remaining = dict(stock)
    by_row = {}
    for order_number, item_code, _date, _line, row, requested in records:
        key = (order_number, item_code)
        allocated = min(requested, remaining.get(key, 0.0))
        remaining[key] = remaining.get(key, 0.0) - allocated
        by_row[row] = int(allocated) if allocated.is_integer() else allocated
    return [by_row[row] for row in range(2, 19)]


EXPECTED = _expected()


def test_fifo_inventory_allocation() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=True)["Sheet1"]
    assert [sheet.cell(row, 8).value for row in range(2, 19)] == EXPECTED
