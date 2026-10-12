from __future__ import annotations

from pathlib import Path

import openpyxl


answer_sheet = "None"
range_str = "R2:R20"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _expected() -> list[str]:
    sheet = openpyxl.load_workbook(INPUT, data_only=False).active
    expected = []
    for row in range(2, 21):
        status = sheet.cell(row, 19)
        difference = sheet.cell(row, 21).value
        if status.data_type == "e":
            expected.append("Upload")
        elif difference is not None and -1 <= float(difference) <= 1:
            expected.append("Do not Upload")
        else:
            expected.append("to Check")
    return expected


EXPECTED = _expected()


def test_error_and_difference_classification() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=True).active
    assert [sheet.cell(row, 18).value for row in range(2, 21)] == EXPECTED
