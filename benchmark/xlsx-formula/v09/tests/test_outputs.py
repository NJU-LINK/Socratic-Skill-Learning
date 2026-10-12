from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import openpyxl


answer_sheet = "None"
range_str = "E2:E7"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _expected() -> list[object]:
    sheet = openpyxl.load_workbook(INPUT, data_only=True).active
    source: dict[str, list[object]] = defaultdict(list)
    for row in range(2, sheet.max_row + 1):
        product = sheet.cell(row, 1).value
        if product not in (None, ""):
            source[str(product)].append(sheet.cell(row, 2).value)
    seen: dict[str, int] = defaultdict(int)
    expected = []
    for row in range(2, 8):
        product = str(sheet.cell(row, 4).value)
        index = seen[product]
        expected.append(source[product][index] if index < len(source[product]) else None)
        seen[product] += 1
    return expected


EXPECTED = _expected()


def test_occurrence_aware_values() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=True).active
    assert [sheet.cell(row, 5).value for row in range(2, 8)] == EXPECTED


def test_required_target_formatting() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=False).active
    for row in range(2, 10):
        cell = sheet.cell(row, 5)
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.type == "rgb"
        assert str(cell.fill.fgColor.rgb).upper().endswith("FCE4D6")
        assert cell.font.name == "Calibri"
        assert cell.number_format == "0;-0;;@"
