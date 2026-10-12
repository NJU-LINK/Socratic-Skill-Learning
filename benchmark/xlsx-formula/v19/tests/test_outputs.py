from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import openpyxl


answer_sheet = "Foglio1"
range_str = "G2:P18"
INPUT = Path("/app/input.xlsx")
OUTPUT = Path("/app/output.xlsx")


def _expected() -> list[list[object]]:
    workbook = openpyxl.load_workbook(INPUT, data_only=True)
    data_sheet = workbook["DATI"]
    output_sheet = workbook["Foglio1"]
    by_identifier: dict[object, list[object]] = defaultdict(list)
    for row in range(2, data_sheet.max_row + 1):
        identifier = data_sheet.cell(row, 11).value
        value = data_sheet.cell(row, 10).value
        if identifier not in (None, "") and value not in (None, ""):
            by_identifier[identifier].append(value)

    expected = []
    for row in range(2, 19):
        latest = by_identifier[output_sheet.cell(row, 5).value][-10:]
        expected.append([None] * (10 - len(latest)) + latest)
    return expected


EXPECTED = _expected()


def test_complete_last_ten_matrix() -> None:
    sheet = openpyxl.load_workbook(OUTPUT, data_only=True)["Foglio1"]
    actual = [
        [sheet.cell(row, column).value for column in range(7, 17)]
        for row in range(2, 19)
    ]
    assert actual == EXPECTED
