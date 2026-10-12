from __future__ import annotations

import ast
import json
import math
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula


TEST_ROOT = Path(os.environ.get("TEST_DIR", "/tests"))
OUTPUT = Path("/app/output.xlsx")
INPUT = Path("/app/input.xlsx")
MAX_VALUE_GROUPS = 8


def _literal_assignment(tree: ast.AST, name: str) -> str:
    values: list[str] = []
    for node in ast.walk(tree):
        if not isinstance(node, ast.Assign) or not isinstance(node.value, ast.Constant):
            continue
        if not isinstance(node.value.value, str):
            continue
        if any(isinstance(target, ast.Name) and target.id == name for target in node.targets):
            values.append(node.value.value)
    if len(values) != 1:
        raise AssertionError(f"Expected exactly one literal {name} assignment, found {values!r}")
    return values[0]


def _value_equal(actual: object, expected: object) -> bool:
    if actual == expected:
        return True
    if actual is None and expected == "":
        return True
    if actual == "" and expected is None:
        return True
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return abs(float(actual) - float(expected)) < 1e-6
    return str(actual) == str(expected)


def _preserved_content(value: object) -> object:
    if value is None or value == "":
        return None
    if isinstance(value, ArrayFormula):
        return ("array-formula", value.ref, value.text)
    return value


def _expected_groups() -> tuple[str, str, list[list[tuple[str, object]]]]:
    source = (TEST_ROOT / "test_outputs.py").read_text(encoding="utf-8")
    tree = ast.parse(source)
    sheet_name = _literal_assignment(tree, "answer_sheet")
    answer_range = _literal_assignment(tree, "range_str")
    expected = json.loads((TEST_ROOT / "expected.json").read_text(encoding="utf-8"))["values"]

    min_col, min_row, max_col, max_row = range_boundaries(answer_range)
    coordinates = [
        f"{get_column_letter(column)}{row}"
        for row in range(min_row, max_row + 1)
        for column in range(min_col, max_col + 1)
    ]
    if len(coordinates) != len(expected):
        raise AssertionError(
            f"{answer_range} contains {len(coordinates)} cells but expected.json has {len(expected)} values"
        )
    cases = list(zip(coordinates, expected))
    if not cases:
        raise AssertionError("The scorer contains no target cells")
    chunk_size = max(1, math.ceil(len(cases) / MAX_VALUE_GROUPS))
    groups = [cases[index:index + chunk_size] for index in range(0, len(cases), chunk_size)]
    return sheet_name, answer_range, groups


ANSWER_SHEET, ANSWER_RANGE, EXPECTED_GROUPS = _expected_groups()


@pytest.mark.parametrize(
    "expected_group",
    EXPECTED_GROUPS,
    ids=[f"values-{index + 1:02d}" for index in range(len(EXPECTED_GROUPS))],
)
def test_expected_value_group(expected_group: list[tuple[str, object]]) -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    workbook = openpyxl.load_workbook(OUTPUT, data_only=True)
    sheet_name = workbook.sheetnames[0] if ANSWER_SHEET == "None" else ANSWER_SHEET
    assert sheet_name in workbook.sheetnames
    sheet = workbook[sheet_name]
    mismatches = [
        (coordinate, sheet[coordinate].value, expected)
        for coordinate, expected in expected_group
        if not _value_equal(sheet[coordinate].value, expected)
    ]
    assert not mismatches, f"Value mismatches: {mismatches[:8]!r}"


def test_output_preserves_all_input_sheets() -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    input_workbook = openpyxl.load_workbook(INPUT, read_only=True)
    output_workbook = openpyxl.load_workbook(OUTPUT, read_only=True)
    missing = sorted(set(input_workbook.sheetnames) - set(output_workbook.sheetnames))
    assert not missing, f"Output removed input sheets: {missing}"


def test_target_cells_contain_no_formulas() -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    workbook = openpyxl.load_workbook(OUTPUT, data_only=False)
    sheet_name = workbook.sheetnames[0] if ANSWER_SHEET == "None" else ANSWER_SHEET
    sheet = workbook[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(ANSWER_RANGE)
    formulas = []
    for row in range(min_row, max_row + 1):
        for column in range(min_col, max_col + 1):
            cell = sheet.cell(row, column)
            if cell.data_type == "f" or isinstance(cell.value, ArrayFormula):
                formulas.append(cell.coordinate)
    assert not formulas, f"Target cells must contain computed values, not formulas: {formulas[:12]!r}"


def test_output_preserves_non_target_cell_contents() -> None:
    assert OUTPUT.exists(), "output.xlsx not found"
    input_workbook = openpyxl.load_workbook(INPUT, data_only=False)
    output_workbook = openpyxl.load_workbook(OUTPUT, data_only=False)
    target_sheet = input_workbook.sheetnames[0] if ANSWER_SHEET == "None" else ANSWER_SHEET
    min_col, min_row, max_col, max_row = range_boundaries(ANSWER_RANGE)

    mismatches: list[tuple[str, str, object, object]] = []
    for sheet_name in input_workbook.sheetnames:
        if sheet_name not in output_workbook.sheetnames:
            continue
        input_sheet = input_workbook[sheet_name]
        output_sheet = output_workbook[sheet_name]
        for row in input_sheet.iter_rows():
            for input_cell in row:
                in_target = (
                    sheet_name == target_sheet
                    and min_row <= input_cell.row <= max_row
                    and min_col <= input_cell.column <= max_col
                )
                if in_target:
                    continue
                output_value = output_sheet[input_cell.coordinate].value
                if _preserved_content(output_value) != _preserved_content(input_cell.value):
                    mismatches.append(
                        (sheet_name, input_cell.coordinate, output_value, input_cell.value)
                    )
                    if len(mismatches) >= 12:
                        break
            if len(mismatches) >= 12:
                break
        if len(mismatches) >= 12:
            break

    assert not mismatches, f"Non-target workbook contents changed: {mismatches!r}"
