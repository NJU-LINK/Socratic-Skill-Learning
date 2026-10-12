import json, os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

def load():
    return Presentation("/app/output.pptx")

def load_report():
    try:
        with open("/app/report.json") as f:
            return json.load(f)
    except:
        return {}

def test_output_exists():
    assert os.path.exists("/app/output.pptx"), "output.pptx not found"

def test_no_empty_output():
    prs = load()
    assert len(prs.slides) > 0, "No slides in output"


def test_change_text():
    prs = load()
    found = False
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and "Q-Material" in shape.text:
            found = True
            break
    assert found, "Text change not found on slide 0"


def test_notes_added():
    prs = load()
    ns = prs.slides[0].notes_slide
    assert "updated" in ns.notes_text_frame.text.lower(), f"Notes missing expected content, got: {ns.notes_text_frame.text[:60]}"
