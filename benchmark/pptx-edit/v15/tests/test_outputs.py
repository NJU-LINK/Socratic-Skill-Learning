import json, os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

def load():
    return Presentation("/app/output.pptx")

def load_report():
    try:
        with open("/app/report.json") as f:
            return json.load(f)
    except:
        return {}

def test_output_exists():
    assert os.path.exists("/app/output.pptx"), "output.pptx not found"


def test_title_bold():
    prs = load()
    last = prs.slides[-1]
    found = False
    for shape in last.shapes:
        if shape.has_text_frame and "Summary" in shape.text:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.bold:
                        found = True
    assert found, "Bold text 'Summary' not found on new slide"


def test_no_empty_output():
    prs = load()
    assert len(prs.slides) == 11, f"Expected 11 slides, got {len(prs.slides)}"




def test_change_title():
    prs = load()
    found = False
    for shape in prs.slides[-1].shapes:
        if shape.has_text_frame and "Summary" in shape.text:
            found = True
            break
    assert found, "Title 'Summary' not found on new slide"
