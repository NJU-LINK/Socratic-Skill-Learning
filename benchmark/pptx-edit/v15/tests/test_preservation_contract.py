from pptx.util import Inches, Pt

from pptx_contract import assert_slide_exact, load_presentations


def test_all_ten_original_slides_are_exactly_preserved():
    before, after = load_presentations()
    assert len(before.slides) == 10
    assert len(after.slides) == 11
    for index in range(10):
        assert_slide_exact(before.slides[index], after.slides[index])


def test_new_slide_has_only_the_requested_textbox():
    _, after = load_presentations()
    added = after.slides[-1]
    assert added.slide_layout.name == "Slide 1 master"
    assert len(added.shapes) == 1

    box = added.shapes[0]
    assert box.has_text_frame
    assert (box.left, box.top, box.width, box.height) == (
        Inches(1),
        Inches(1),
        Inches(8),
        Inches(2),
    )
    assert box.text == "Summary"
    runs = [
        run
        for paragraph in box.text_frame.paragraphs
        for run in paragraph.runs
    ]
    assert len(runs) == 1
    assert runs[0].font.bold is True
    assert runs[0].font.size == Pt(36)
