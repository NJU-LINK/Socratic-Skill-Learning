from pptx import Presentation

from pptx_contract import (
    INPUT,
    OUTPUT,
    assert_package_resources_preserved,
)


def test_output_is_a_valid_presentation():
    assert OUTPUT.is_file()
    presentation = Presentation(OUTPUT)
    assert presentation.slides


def test_page_dimensions_are_preserved():
    before = Presentation(INPUT)
    after = Presentation(OUTPUT)
    assert (after.slide_width, after.slide_height) == (
        before.slide_width,
        before.slide_height,
    )


def test_media_theme_master_and_layout_resources_are_preserved():
    assert_package_resources_preserved()
