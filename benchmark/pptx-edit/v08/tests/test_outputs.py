import json, os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

def load():
    return Presentation("/app/output.pptx")

def load_report():
    try:
        with open("/app/report.json") as f:
            return json.load(f)
    except:
        return {}

def test_output_exists():
    assert os.path.exists("/app/output.pptx"), "output.pptx not found"

def test_no_empty_output():
    prs = load()
    assert len(prs.slides) == 10, f"Expected 10 slides, got {len(prs.slides)}"

def test_change_text():
    prs = load()
    found = False
    for shape in prs.slides[3].shapes:
        if shape.has_text_frame and "Hadal Zone" in shape.text:
            found = True
            break
    assert found, "Text change not found on slide 4"

def test_slide_deleted():
    prs = load()
    all_text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text += shape.text
    assert "哈达尔区的独特之处" not in all_text, "Deleted slide's content still present"
