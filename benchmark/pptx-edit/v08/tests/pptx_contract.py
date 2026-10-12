from copy import deepcopy
from collections import Counter
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation


INPUT = Path("/app/input.pptx")
OUTPUT = Path("/app/output.pptx")

DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
SEMANTIC_XML_PREFIXES = (
    "ppt/theme/",
    "ppt/slideMasters/",
    "ppt/slideLayouts/",
)
PAYLOAD_PACKAGE_PREFIXES = (
    "ppt/media/",
    "ppt/embeddings/",
    "ppt/fonts/",
)
NOTES_RELATIONSHIP = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/notesSlide"
)


def load_presentations() -> tuple[Presentation, Presentation]:
    return Presentation(INPUT), Presentation(OUTPUT)


def canonical_xml(element) -> bytes:
    root = etree.fromstring(element.xml.encode())
    return etree.tostring(root, method="c14n")


def normalized_text_target_xml(element, *, allow_bold_change: bool) -> bytes:
    root = etree.fromstring(deepcopy(element).xml.encode())
    texts = root.xpath(
        ".//*[local-name()='t' and namespace-uri()=$namespace]",
        namespace=DRAWING_NS,
    )
    assert texts, "target shape has no DrawingML text nodes"
    for index, text in enumerate(texts):
        text.text = f"__TARGET_TEXT_{index}__"
    if allow_bold_change:
        for node in root.iter():
            node.attrib.pop("b", None)
        for run_properties in root.xpath(
            ".//*[local-name()='rPr' and namespace-uri()=$namespace]",
            namespace=DRAWING_NS,
        ):
            if not run_properties.attrib and len(run_properties) == 0:
                run_properties.getparent().remove(run_properties)
    return etree.tostring(root, method="c14n")


def title_shape_index(slide) -> int:
    title = slide.shapes.title
    if title is not None:
        for index, shape in enumerate(slide.shapes):
            if shape.element is title.element:
                return index
    for index, shape in enumerate(slide.shapes):
        if shape.has_text_frame and shape.text.strip():
            return index
    raise AssertionError("slide has no non-empty text shape")


def canonical_part_digest(part) -> str:
    blob = part.blob
    if blob.lstrip().startswith(b"<"):
        blob = etree.tostring(etree.fromstring(blob), method="c14n")
    return sha256(blob).hexdigest()


def slide_relationship_signature(slide, *, ignore_notes: bool = False) -> list[tuple]:
    relationships = []
    for relationship in slide.part.rels.values():
        if ignore_notes and relationship.reltype == NOTES_RELATIONSHIP:
            continue
        target = (
            relationship.target_ref
            if relationship.is_external
            else canonical_part_digest(relationship.target_part)
        )
        relationships.append(
            (
                relationship.rId,
                relationship.reltype,
                relationship.is_external,
                target,
            )
        )
    return sorted(relationships)


def assert_slide_exact(before, after) -> None:
    assert canonical_xml(after.element) == canonical_xml(before.element)
    assert after.slide_layout.name == before.slide_layout.name
    assert slide_relationship_signature(after) == slide_relationship_signature(before)


def assert_slide_with_title_edit(
    before,
    after,
    *,
    allow_bold_change: bool,
    allow_notes_change: bool = False,
) -> None:
    assert len(after.shapes) == len(before.shapes)
    assert after.slide_layout.name == before.slide_layout.name
    assert slide_relationship_signature(
        after,
        ignore_notes=allow_notes_change,
    ) == slide_relationship_signature(
        before,
        ignore_notes=allow_notes_change,
    )
    before_target = title_shape_index(before)
    after_target = title_shape_index(after)
    assert after_target == before_target, "title shape position changed"

    for index, (before_shape, after_shape) in enumerate(zip(before.shapes, after.shapes)):
        if index == before_target:
            assert normalized_text_target_xml(
                after_shape.element,
                allow_bold_change=allow_bold_change,
            ) == normalized_text_target_xml(
                before_shape.element,
                allow_bold_change=allow_bold_change,
            ), "target title changed beyond allowed text/bold properties"
        else:
            assert canonical_xml(after_shape.element) == canonical_xml(before_shape.element), (
                f"non-target shape {index} changed"
            )


def semantic_xml_hashes(archive: ZipFile, prefix: str) -> Counter:
    return Counter(
        sha256(
            etree.tostring(
                etree.fromstring(archive.read(name)),
                method="c14n",
            )
        ).hexdigest()
        for name in archive.namelist()
        if name.startswith(prefix) and name.endswith(".xml")
    )


def payload_hashes(archive: ZipFile, prefix: str) -> Counter:
    return Counter(
        sha256(archive.read(name)).hexdigest()
        for name in archive.namelist()
        if name.startswith(prefix) and not name.endswith("/")
    )


def assert_package_resources_preserved() -> None:
    with ZipFile(INPUT) as before_zip, ZipFile(OUTPUT) as after_zip:
        before_names = before_zip.namelist()
        after_names = after_zip.namelist()
        notes_master_added = sum(
            name.startswith("ppt/notesMasters/") and name.endswith(".xml")
            for name in after_names
        ) > sum(
            name.startswith("ppt/notesMasters/") and name.endswith(".xml")
            for name in before_names
        )

        for prefix in SEMANTIC_XML_PREFIXES:
            before_resources = semantic_xml_hashes(before_zip, prefix)
            after_resources = semantic_xml_hashes(after_zip, prefix)
            if prefix == "ppt/theme/" and notes_master_added:
                assert before_resources <= after_resources, (
                    "existing presentation themes changed while adding notes"
                )
                assert sum((after_resources - before_resources).values()) == 1, (
                    "adding notes introduced unexpected extra themes"
                )
            else:
                assert after_resources == before_resources, (
                    f"semantic package resources changed under {prefix}"
                )

        before_presentation = Presentation(INPUT)
        after_presentation = Presentation(OUTPUT)
        slides_deleted = len(after_presentation.slides) < len(before_presentation.slides)
        for prefix in PAYLOAD_PACKAGE_PREFIXES:
            before_payload = payload_hashes(before_zip, prefix)
            after_payload = payload_hashes(after_zip, prefix)
            if slides_deleted:
                assert after_payload <= before_payload, (
                    f"new or modified payload resource under {prefix}"
                )
            else:
                assert after_payload == before_payload, (
                    f"payload resources changed under {prefix}"
                )
