import json, os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

def load():
    return Presentation("/app/output.pptx")

def load_report():
    try:
        with open("/app/report.json") as f:
            return json.load(f)
    except:
        return {}

def test_output_exists():
    assert os.path.exists("/app/output.pptx"), "output.pptx not found"

def test_no_empty_output():
    prs = load()
    assert len(prs.slides) == 10, f"Expected 10 slides, got {len(prs.slides)}"


def test_title_bold():
    prs = load()
    slide = prs.slides[0]
    found = False
    for shape in slide.shapes:
        if shape.has_text_frame and "深海探索" in shape.text:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.bold:
                        found = True
    assert found, "Bold text '深海探索' not found on slide 1"


def test_change_title():
    prs = load()
    found = False
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame and "深海探索" in shape.text:
            found = True
            break
    assert found, "Title '深海探索' not found on slide 0"

def test_notes_added():
    prs = load()
    ns = prs.slides[0].notes_slide
    assert len(ns.notes_text_frame.text) > 0
