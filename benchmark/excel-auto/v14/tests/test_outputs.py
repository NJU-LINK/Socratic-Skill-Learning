from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pypdf import PdfReader


DATA = Path("/app/data")
SOURCE_NAMES = ["Invoice.pdf", "Invoice_2.pdf"]


def _extract_source(name: str) -> tuple[str, str, float, float]:
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(DATA / name).pages
    )
    item_match = None
    for raw_line in text.splitlines():
        line = " ".join(raw_line.split())
        item_match = re.fullmatch(
            r"\d+(?:\s+TB)?\s+(.+?)\s+([\d,]+\.\d{2})\s+([\d,]+\.\d{2})",
            line,
        )
        if item_match:
            break
    assert item_match, f"No line item found in {name}"
    total_match = re.search(r"TOTAL DUE\s+([\d,]+\.\d{2})", text, re.I)
    assert total_match, f"No TOTAL DUE found in {name}"
    return (
        name,
        item_match.group(1),
        float(item_match.group(3).replace(",", "")),
        float(total_match.group(1).replace(",", "")),
    )


EXPECTED = [_extract_source(name) for name in SOURCE_NAMES]
EXPECTED_COMBINED = sum(case[3] for case in EXPECTED)

sys.path.insert(0, "/app")
from solve import run


def test_workbook_schema_and_cardinality() -> None:
    run()
    sheet = openpyxl.load_workbook(DATA / "invoices.xlsx", data_only=True).active
    assert [sheet.cell(1, column).value for column in range(1, 4)] == ["Invoice", "Item", "Cost"]
    assert sheet.max_row == 4
    assert sheet.max_column == 3


@pytest.mark.parametrize(("row_number", "expected"), list(enumerate(EXPECTED, start=2)))
def test_workbook_line_item(row_number: int, expected: tuple[str, str, float, float]) -> None:
    run()
    sheet = openpyxl.load_workbook(DATA / "invoices.xlsx", data_only=True).active
    name, item, line_total, _total_due = expected
    assert sheet.cell(row_number, 1).value == name
    assert sheet.cell(row_number, 2).value == item
    assert abs(float(sheet.cell(row_number, 3).value) - line_total) < 0.01


def test_workbook_combined_total_due() -> None:
    run()
    sheet = openpyxl.load_workbook(DATA / "invoices.xlsx", data_only=True).active
    assert sheet["A4"].value == "Total"
    assert sheet["B4"].value == "Combined TOTAL DUE"
    assert abs(float(sheet["C4"].value) - EXPECTED_COMBINED) < 0.01


@pytest.mark.parametrize("expected", EXPECTED)
def test_docx_reports_each_invoice_total(expected: tuple[str, str, float, float]) -> None:
    run()
    text = "\n".join(paragraph.text for paragraph in Document(DATA / "report.docx").paragraphs)
    name, _item, _line_total, total_due = expected
    assert name in text
    assert f"{total_due:.2f}" in text


def test_docx_reports_combined_total() -> None:
    run()
    text = "\n".join(paragraph.text for paragraph in Document(DATA / "report.docx").paragraphs)
    assert "Combined TOTAL DUE" in text
    assert f"{EXPECTED_COMBINED:.2f}" in text
