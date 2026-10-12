from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from docx import Document


sys.path.insert(0, "/app")
from solve import run


DATA = Path("/app/data")


def _expected_difference() -> int:
    workbook = openpyxl.load_workbook(DATA / "company_revenues.xlsx", data_only=True)
    sheet = workbook.active
    return int(sheet.cell(row=2, column=2).value - sheet.cell(row=sheet.max_row, column=2).value)


EXPECTED_DIFFERENCE = _expected_difference()


def test_numeric_output_exists() -> None:
    run()
    assert (DATA / "improvement.txt").exists()


def test_numeric_output_is_exact() -> None:
    run()
    assert (DATA / "improvement.txt").read_text(encoding="utf-8").strip() == str(EXPECTED_DIFFERENCE)


def test_docx_output_exists_and_is_readable() -> None:
    run()
    report = DATA / "reports.docx"
    assert report.exists()
    document = Document(report)
    assert any(paragraph.text.strip() for paragraph in document.paragraphs)


def test_docx_reports_the_same_difference() -> None:
    run()
    document = Document(DATA / "reports.docx")
    text = " ".join(paragraph.text for paragraph in document.paragraphs)
    assert str(EXPECTED_DIFFERENCE) in text
