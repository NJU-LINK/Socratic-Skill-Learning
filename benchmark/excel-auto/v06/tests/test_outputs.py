from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest
from pypdf import PdfReader


DATA = Path("/app/data")
SOURCE = DATA / "financial_report_2.xlsx"
OUTPUT = DATA / "financial/report.pdf"


def _expected() -> tuple[str, dict[str, float]]:
    workbook = openpyxl.load_workbook(SOURCE, data_only=False)
    sheet = workbook["BUDGET TOOL"]
    company = str(sheet["B2"].value)
    investment = sum(float(sheet.cell(row=row, column=3).value) for row in range(11, 14))
    benefits = [
        sum(float(sheet.cell(row=row, column=column).value or 0) for row in range(17, 22))
        for column in range(3, 6)
    ]
    depreciation = investment / 3
    costs = [
        sum(float(sheet.cell(row=row, column=column).value or 0) for row in range(25, 29))
        + depreciation
        for column in range(3, 6)
    ]
    metrics = {"Total initial investment": investment}
    metrics.update({f"Benefits Year {year}": value for year, value in enumerate(benefits, 1)})
    metrics.update({f"Costs Year {year}": value for year, value in enumerate(costs, 1)})
    return company, metrics


EXPECTED_COMPANY, EXPECTED_METRICS = _expected()

sys.path.insert(0, "/app")
from solve import run


def _pdf_text() -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(OUTPUT).pages)


def test_pdf_exists_and_is_readable() -> None:
    run()
    assert OUTPUT.exists()
    assert _pdf_text().strip()


def test_pdf_identifies_company() -> None:
    run()
    assert EXPECTED_COMPANY in _pdf_text()


@pytest.mark.parametrize(("label", "expected"), list(EXPECTED_METRICS.items()))
def test_pdf_reports_metric(label: str, expected: float) -> None:
    run()
    match = re.search(rf"{re.escape(label)}\s*:\s*([-0-9,.]+)", _pdf_text(), re.I)
    assert match, f"Missing metric: {label}"
    actual = float(match.group(1).replace(",", ""))
    assert abs(actual - expected) < 0.01
