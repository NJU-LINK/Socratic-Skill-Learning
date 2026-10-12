from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pypdf import PdfReader


sys.path.insert(0, "/app")
from solve import run


DATA = Path("/app/data")


def _expected_metrics() -> dict[str, float]:
    workbook = openpyxl.load_workbook(DATA / "score.xlsx", data_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    first = [float(row[1]) for row in rows]
    second = [float(row[2]) for row in rows]
    return {
        "average of midterm1": sum(first) / len(first),
        "average of midterm2": sum(second) / len(second),
        "students who improved": float(sum(b > a for a, b in zip(first, second))),
        "pass rate midterm1": 100 * sum(score >= 60 for score in first) / len(first),
        "pass rate midterm2": 100 * sum(score >= 60 for score in second) / len(second),
    }


EXPECTED_METRICS = _expected_metrics()


def _docx_text() -> str:
    document = Document(DATA / "teaching.docx")
    return " ".join(paragraph.text for paragraph in document.paragraphs)


def _pdf_text() -> str:
    reader = PdfReader(DATA / "teaching.pdf")
    return " ".join(page.extract_text() or "" for page in reader.pages)


def _reported_value(text: str, label: str) -> float:
    match = re.search(rf"{re.escape(label)}\s*:\s*([0-9]+(?:\.[0-9]+)?)", text, re.I)
    assert match, f"Missing metric {label!r}"
    return float(match.group(1))


def test_docx_exists() -> None:
    run()
    assert (DATA / "teaching.docx").exists()


def test_pdf_exists() -> None:
    run()
    assert (DATA / "teaching.pdf").exists()


@pytest.mark.parametrize("label", list(EXPECTED_METRICS))
def test_docx_metric(label: str) -> None:
    run()
    assert abs(_reported_value(_docx_text(), label) - EXPECTED_METRICS[label]) < 0.11


@pytest.mark.parametrize("label", list(EXPECTED_METRICS))
def test_pdf_metric(label: str) -> None:
    run()
    assert abs(_reported_value(_pdf_text(), label) - EXPECTED_METRICS[label]) < 0.11
