from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pypdf import PdfReader


DATA = Path("/app/data")


def _expected() -> list[tuple[str, float, str]]:
    text = "\n".join(
        page.extract_text() or ""
        for page in PdfReader(DATA / "financial_report.pdf").pages
    )
    cases = [
        ("Total revenue", r"Quarterly revenue of \$([0-9.]+) billion", "USD billions"),
        ("Data Center revenue", r"Record Data Center revenue of \$([0-9.]+) billion", "USD billions"),
        (
            "GAAP diluted earnings per share",
            r"GAAP earnings per diluted share for the quarter were \$([0-9.]+)",
            "USD per diluted share",
        ),
    ]
    expected = []
    for label, pattern, unit in cases:
        match = re.search(pattern, text, re.I)
        assert match, f"Source PDF does not contain {label}"
        expected.append((label, float(match.group(1)), unit))
    return expected


EXPECTED = _expected()

sys.path.insert(0, "/app")
from solve import run


def test_workbook_schema_and_cardinality() -> None:
    run()
    workbook = openpyxl.load_workbook(DATA / "report.xlsx", data_only=True)
    assert workbook.sheetnames == ["Q1 FY24"]
    sheet = workbook["Q1 FY24"]
    assert [sheet.cell(1, column).value for column in range(1, 4)] == ["Metric", "Value", "Unit"]
    assert sheet.max_row == 4
    assert sheet.max_column == 3


@pytest.mark.parametrize(("row_number", "expected"), list(enumerate(EXPECTED, start=2)))
def test_workbook_metric(row_number: int, expected: tuple[str, float, str]) -> None:
    run()
    sheet = openpyxl.load_workbook(DATA / "report.xlsx", data_only=True)["Q1 FY24"]
    label, value, unit = expected
    assert sheet.cell(row_number, 1).value == label
    assert abs(float(sheet.cell(row_number, 2).value) - value) < 1e-9
    assert sheet.cell(row_number, 3).value == unit


@pytest.mark.parametrize("expected", EXPECTED)
def test_docx_reports_metric(expected: tuple[str, float, str]) -> None:
    run()
    text = "\n".join(paragraph.text for paragraph in Document(DATA / "summary.docx").paragraphs)
    label, value, unit = expected
    assert label in text
    assert f"{value:.2f}" in text
    assert unit in text
