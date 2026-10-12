from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest
from icalendar import Calendar
from pypdf import PdfReader


sys.path.insert(0, "/app")
from solve import run


DATA = Path("/app/data")


def _expected_cases() -> list[tuple[str, str, int]]:
    cases: list[tuple[str, str, int]] = []
    for person in ("Tom", "Bob"):
        calendar = Calendar.from_ical((DATA / f"calendar/{person}.ics").read_bytes())
        counts: dict[str, int] = {}
        for component in calendar.walk():
            if component.name != "VEVENT":
                continue
            summary = str(component.get("summary")).lower().strip()
            category = next(
                (
                    label
                    for label in ("class", "sleep", "nap", "dinner", "lunch")
                    if label in summary
                ),
                summary,
            )
            counts[category] = counts.get(category, 0) + 1
        cases.extend((person, category, count) for category, count in sorted(counts.items()))
    return cases


def test_workbook_has_exact_schema_and_rows() -> None:
    run()
    workbook = openpyxl.load_workbook(DATA / "events.xlsx", data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Person", "Event Type", "Count")
    assert rows[1:] == _expected_cases()


@pytest.mark.parametrize(("person", "category", "count"), _expected_cases())
def test_workbook_contains_each_event_count(person: str, category: str, count: int) -> None:
    run()
    workbook = openpyxl.load_workbook(DATA / "events.xlsx", data_only=True)
    rows = list(workbook.active.iter_rows(min_row=2, values_only=True))
    assert (person, category, count) in rows


@pytest.mark.parametrize(("person", "category", "count"), _expected_cases())
def test_pdf_contains_each_event_count(person: str, category: str, count: int) -> None:
    run()
    text = "\n".join(
        page.extract_text() or "" for page in PdfReader(DATA / "daily_event.pdf").pages
    )
    section_start = text.find(f"{person}:")
    assert section_start >= 0
    other_person = "Bob" if person == "Tom" else "Tom"
    section_end = text.find(f"{other_person}:", section_start + len(person) + 1)
    section = text[section_start:] if section_end < 0 else text[section_start:section_end]
    assert f"{category}: {count}" in section
