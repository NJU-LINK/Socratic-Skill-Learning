#!/usr/bin/env python3
"""
Reusable solver template for xlsx-structure block-sort tasks.

Usage:
    python3 block_sort_solver.py <input_xlsx> <output_xlsx> <sheet_name> <start_row> <end_row> <start_col> <end_col> <sort_col> [desc]

Example:
    python3 block_sort_solver.py /app/input.xlsx /app/output.xlsx Sheet1 1 85 1 8 4 desc

Arguments:
    input_xlsx   - Path to input workbook
    output_xlsx  - Path to save output workbook
    sheet_name   - Name of the sheet to transform
    start_row    - First row of target range (1-based)
    end_row      - Last row of target range (1-based)
    start_col    - First column of target range (1-based, A=1)
    end_col      - Last column of target range (1-based, H=8)
    sort_col     - Column index to sort by (1-based, D=4)
    desc         - Optional: 'desc' for descending (default), 'asc' for ascending
"""

import sys
import re
import openpyxl


def to_number(val):
    """Convert cell value to float, handling None, int, float, string-numbers."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        val = val.strip()
        if val == '':
            return 0.0
        try:
            return float(val)
        except ValueError:
            return 0.0
    return 0.0


def col_letters_to_idx(letters):
    """Convert column letters (A, B, ..., Z, AA, AB, ...) to 1-based index."""
    idx = 0
    for ch in letters:
        idx = idx * 26 + (ord(ch) - ord('A') + 1)
    return idx


def eval_formula_with_deps(expr, row_vals, evaluated_cache):
    """
    Evaluate a formula string, resolving dependencies within the row.
    row_vals: dict {col_idx: original_value}
    evaluated_cache: dict {col_idx: evaluated_numeric_value}
    """
    if not isinstance(expr, str) or not expr.startswith('='):
        return to_number(expr)

    expr_body = expr[1:]  # strip '='

    # Find all cell references
    refs = set(re.findall(r'([A-Z]+)(\d+)', expr_body))

    # Resolve dependencies: if a referenced column has a formula, evaluate it first
    for col_letters, row_num in refs:
        col_idx = col_letters_to_idx(col_letters)
        if col_idx not in evaluated_cache:
            ref_val = row_vals.get(col_idx)
            if isinstance(ref_val, str) and ref_val.startswith('='):
                evaluated_cache[col_idx] = eval_formula_with_deps(ref_val, row_vals, evaluated_cache)

    # Substitute all references with their evaluated values
    def replace_ref(match):
        col_letters = match.group(1)
        col_idx = col_letters_to_idx(col_letters)
        if col_idx in evaluated_cache:
            return str(evaluated_cache[col_idx])
        val = row_vals.get(col_idx)
        return str(to_number(val))

    numeric_expr = re.sub(r'([A-Z]+)\d+', replace_ref, expr_body)
    result = eval(numeric_expr)
    return round(result, 10)


def is_blank_row(ws, row, start_col, end_col):
    """Check if a row is blank (all cells None or empty string)."""
    for c in range(start_col, end_col + 1):
        v = ws.cell(row=row, column=c).value
        if v is not None and v != '':
            return False
    return True


def detect_blocks(ws, start_row, end_row, start_col, end_col):
    """Detect blocks of non-blank rows separated by blank rows."""
    blocks = []
    current_block = []
    for r in range(start_row, end_row + 1):
        if is_blank_row(ws, r, start_col, end_col):
            if current_block:
                blocks.append(current_block)
                current_block = []
        else:
            vals = {c: ws.cell(row=r, column=c).value for c in range(start_col, end_col + 1)}
            current_block.append((r, vals))
    if current_block:
        blocks.append(current_block)
    return blocks


def evaluate_formulas(blocks, start_col, end_col):
    """Evaluate all formula strings in the target range."""
    for block in blocks:
        for orig_row, vals in block:
            eval_cache = {}
            for c in range(start_col, end_col + 1):
                v = vals[c]
                if isinstance(v, str) and v.startswith('='):
                    eval_cache[c] = eval_formula_with_deps(v, vals, eval_cache)
                    vals[c] = eval_cache[c]


def sort_blocks(blocks, sort_col, descending=True):
    """Sort each block by the specified column (stable sort)."""
    for block in blocks:
        block.sort(key=lambda x: to_number(x[1][sort_col]), reverse=descending)


def write_results(ws, blocks, start_row, end_row, start_col, end_col):
    """Write sorted blocks back to the worksheet with blank row separators."""
    output_row = start_row
    for block in blocks:
        for orig_row, vals in block:
            for c in range(start_col, end_col + 1):
                ws.cell(row=output_row, column=c).value = vals[c]
            output_row += 1
        # Blank row separator between blocks
        if block is not blocks[-1]:
            for c in range(start_col, end_col + 1):
                ws.cell(row=output_row, column=c).value = None
            output_row += 1

    # Clear remaining rows in target range
    for r in range(output_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            ws.cell(row=r, column=c).value = None


def solve(input_path, output_path, sheet_name, start_row, end_row,
          start_col, end_col, sort_col, descending=True):
    """Main solver function."""
    wb = openpyxl.load_workbook(input_path)
    ws = wb[sheet_name]

    # Detect blocks
    blocks = detect_blocks(ws, start_row, end_row, start_col, end_col)
    print(f"Detected {len(blocks)} blocks")
    for i, b in enumerate(blocks):
        print(f"  Block {i}: rows {b[0][0]}-{b[-1][0]}, {len(b)} data rows")

    # Evaluate formulas
    evaluate_formulas(blocks, start_col, end_col)

    # Sort blocks
    sort_blocks(blocks, sort_col, descending)

    # Write results
    write_results(ws, blocks, start_row, end_row, start_col, end_col)

    wb.save(output_path)
    print(f"Saved to {output_path}")


if __name__ == '__main__':
    if len(sys.argv) < 9:
        print(__doc__)
        sys.exit(1)

    input_xlsx = sys.argv[1]
    output_xlsx = sys.argv[2]
    sheet_name = sys.argv[3]
    start_row = int(sys.argv[4])
    end_row = int(sys.argv[5])
    start_col = int(sys.argv[6])
    end_col = int(sys.argv[7])
    sort_col = int(sys.argv[8])
    descending = True
    if len(sys.argv) > 9 and sys.argv[9].lower() == 'asc':
        descending = False

    solve(input_xlsx, output_xlsx, sheet_name, start_row, end_row,
          start_col, end_col, sort_col, descending)
