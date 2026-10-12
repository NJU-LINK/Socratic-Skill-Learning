#!/usr/bin/env python3
"""
xlsx_transform.py — Reusable helpers for xlsx-structure tasks.

Quick start:
    import sys; sys.path.insert(0, "/app/scripts")
    from xlsx_transform import *

    wb = load_workbook("/app/input.xlsx")
    ws = wb["Sheet1"]
    blocks = detect_blocks(ws, col_start=1, col_end=8)
    # ... evaluate, sort, verify, save
"""

import os
import re
import sys
from copy import copy
from typing import List, Tuple, Dict, Any, Optional

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def load_workbook(path: str = "/app/input.xlsx"):
    """Load workbook preserving formulas and styles (NOT data_only)."""
    return openpyxl.load_workbook(path)


# ---------------------------------------------------------------------------
# Topology detection
# ---------------------------------------------------------------------------

def _is_row_blank(ws, row: int, col_start: int, col_end: int) -> bool:
    """Return True if every cell in [col_start, col_end] of *row* is None/empty."""
    for c in range(col_start, col_end + 1):
        v = ws.cell(row=row, column=c).value
        if v is not None and v != "":
            return False
    return True


def detect_blocks(
    ws,
    col_start: int = 1,
    col_end: int = 8,
    row_start: int = 1,
    row_end: Optional[int] = None,
) -> List[Tuple[int, int]]:
    """
    Detect contiguous non-blank row blocks.

    A row is *blank* if every cell in [col_start, col_end] is None/empty.
    Returns list of (first_row, last_row) for each block.
    """
    if row_end is None:
        row_end = ws.max_row or 1
    blocks: List[Tuple[int, int]] = []
    start = None
    for r in range(row_start, row_end + 1):
        if _is_row_blank(ws, r, col_start, col_end):
            if start is not None:
                blocks.append((start, r - 1))
                start = None
        else:
            if start is None:
                start = r
    if start is not None:
        blocks.append((start, row_end))
    return blocks


def find_blank_rows(
    ws,
    col_start: int = 1,
    col_end: int = 8,
    row_start: int = 1,
    row_end: Optional[int] = None,
) -> List[int]:
    """Return row numbers that are completely blank in the given column range."""
    if row_end is None:
        row_end = ws.max_row or 1
    return [
        r for r in range(row_start, row_end + 1)
        if _is_row_blank(ws, r, col_start, col_end)
    ]


# ---------------------------------------------------------------------------
# Formula evaluation
# ---------------------------------------------------------------------------

_FORMULA_RE = re.compile(r'^\s*=\s*(.+)$', re.DOTALL)
_CELL_REF_RE = re.compile(r'([A-Za-z]+)(\d+)')


def _resolve_refs(expr: str, row: int, ws, _depth: int = 0) -> str:
    """Replace cell references in *expr* with their numeric values."""
    if _depth > 20:
        raise RecursionError("Formula reference depth exceeded")

    def _repl(m):
        col_letters, ref_row_str = m.group(1), m.group(2)
        ref_row = int(ref_row_str)
        c = column_index_from_string(col_letters)
        val = ws.cell(row=ref_row, column=c).value
        if val is None:
            return "0"
        if isinstance(val, str) and _FORMULA_RE.match(val):
            val = eval_formula_cell(val, ref_row, ws, _depth + 1)
            return str(val) if val is not None else "0"
        # Numeric or other — coerce to number
        try:
            return str(float(val))
        except (ValueError, TypeError):
            return "0"

    return _CELL_REF_RE.sub(_repl, expr)


def eval_formula_cell(
    formula_str: str,
    row: int,
    ws,
    _depth: int = 0,
) -> Optional[float]:
    """
    Evaluate a simple row-local formula string like '=B1+C1'.

    Supports: +, -, *, /, unary minus, parentheses, cell references,
    numeric literals (int and float, including negative).
    Returns None if the string is not a parseable formula.
    """
    m = _FORMULA_RE.match(formula_str)
    if not m:
        # Not a formula; try as plain number
        try:
            return float(formula_str)
        except (ValueError, TypeError):
            return None

    expr = m.group(1)
    resolved = _resolve_refs(expr, row, ws, _depth)
    # Clean whitespace
    resolved = resolved.replace(" ", "")
    try:
        result = eval(resolved, {"__builtins__": {}}, {})
        return float(result)
    except Exception:
        return None


def is_formula(value) -> bool:
    """Check whether a cell value looks like a formula string."""
    return isinstance(value, str) and bool(_FORMULA_RE.match(value))


def evaluate_formulas_in_range(
    ws,
    rows: List[int],
    formula_cols: List[str],
    col_start: int = 1,
    col_end: int = 8,
):
    """
    Evaluate formula cells in the given columns and rows, replacing them
    with numeric results.  Process columns left-to-right so that dependent
    columns (e.g., F referencing D) see already-evaluated values.

    formula_cols: list of column letters, e.g. ['D', 'F']
    Also converts plain-string numbers to float.
    """
    sorted_cols = sorted(formula_cols, key=lambda c: column_index_from_string(c))
    for r in rows:
        for col_letter in sorted_cols:
            c = column_index_from_string(col_letter)
            cell = ws.cell(row=r, column=c)
            val = cell.value
            if val is None:
                continue
            if is_formula(val):
                result = eval_formula_cell(val, r, ws)
                if result is not None:
                    cell.value = result
            elif isinstance(val, str):
                # Try to convert plain text that looks numeric
                try:
                    cell.value = float(val)
                except (ValueError, TypeError):
                    pass


# ---------------------------------------------------------------------------
# Snapshot / restore
# ---------------------------------------------------------------------------

def _snapshot_cell(cell):
    """Capture value + style for a single cell."""
    return {
        "value": cell.value,
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
    }


def _restore_cell(cell, snap):
    """Write a snapshot dict back to a cell."""
    cell.value = snap["value"]
    cell.font = snap["font"]
    cell.fill = snap["fill"]
    cell.border = snap["border"]
    cell.alignment = snap["alignment"]
    cell.number_format = snap["number_format"]


def snapshot_range(
    ws,
    rows: List[int],
    col_start: int = 1,
    col_end: int = 8,
) -> Dict[Tuple[int, int], Dict[str, Any]]:
    """
    Capture value + style for every cell in the given rows and columns.
    Returns dict keyed by (row, col).
    """
    snap: Dict[Tuple[int, int], Dict[str, Any]] = {}
    for r in rows:
        for c in range(col_start, col_end + 1):
            snap[(r, c)] = _snapshot_cell(ws.cell(row=r, column=c))
    return snap


def write_snapshot(
    ws,
    data: Dict[Tuple[int, int], Dict[str, Any]],
    clear_first: bool = True,
    col_start: int = 1,
    col_end: int = 8,
    rows: Optional[List[int]] = None,
):
    """Write a snapshot dict back to the worksheet."""
    if clear_first and rows is not None:
        for r in rows:
            for c in range(col_start, col_end + 1):
                ws.cell(row=r, column=c).value = None

    for (r, c), d in data.items():
        _restore_cell(ws.cell(row=r, column=c), d)


# ---------------------------------------------------------------------------
# Sorting
# ---------------------------------------------------------------------------

def sort_block_rows(
    ws,
    block_rows: List[int],
    key_col: int,
    descending: bool = True,
    col_start: int = 1,
    col_end: int = 8,
):
    """
    Stable-sort the rows of a block by the numeric value in *key_col*.

    - Snapshots each row's data + styles.
    - Sorts by the evaluated value of key_col (None/missing → 0).
    - Writes the sorted rows back, preserving styles per original row.
    """
    # Snapshot
    row_snaps = []
    for r in block_rows:
        snap = {}
        for c in range(col_start, col_end + 1):
            snap[c] = _snapshot_cell(ws.cell(row=r, column=c))
        row_snaps.append(snap)

    # Build sort key
    def sort_key(item):
        idx, snap = item
        val = snap[key_col]["value"]
        if val is None:
            num = 0.0
        else:
            try:
                num = float(val)
            except (ValueError, TypeError):
                num = 0.0
        return -num if descending else num

    indexed = list(enumerate(row_snaps))
    indexed.sort(key=sort_key)  # Python sort is stable

    # Write back
    for new_pos, (_, snap) in enumerate(indexed):
        target_row = block_rows[new_pos]
        for c in range(col_start, col_end + 1):
            _restore_cell(ws.cell(row=target_row, column=c), snap[c])


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------

def verify_no_formulas(ws, rows: List[int], cols: List[str]) -> List[str]:
    """Return list of error strings for any remaining formula cells."""
    errors = []
    for r in rows:
        for col_letter in cols:
            c = column_index_from_string(col_letter)
            val = ws.cell(row=r, column=c).value
            if is_formula(val):
                errors.append(f"Unevaluated formula at {col_letter}{r}: {val}")
    return errors


def verify_sort_order(
    ws,
    blocks: List[List[int]],
    key_col: int,
    descending: bool = True,
) -> List[str]:
    """Return list of error strings for any sort-order violations."""
    errors = []
    for bi, block in enumerate(blocks):
        prev = None
        for r in block:
            val = ws.cell(row=r, column=key_col).value
            try:
                num = float(val) if val is not None else 0.0
            except (ValueError, TypeError):
                num = 0.0
            if prev is not None:
                if descending and num > prev + 1e-9:
                    errors.append(
                        f"Block {bi}: row {r} value {num} > prev {prev} (desc)"
                    )
                elif not descending and num < prev - 1e-9:
                    errors.append(
                        f"Block {bi}: row {r} value {num} < prev {prev} (asc)"
                    )
            prev = num
    return errors


def verify_block_structure(
    ws,
    expected_blank_rows: List[int],
    col_start: int = 1,
    col_end: int = 8,
) -> List[str]:
    """Verify that blank rows are still blank."""
    errors = []
    for r in expected_blank_rows:
        if not _is_row_blank(ws, r, col_start, col_end):
            errors.append(f"Row {r} should be blank but has data")
    return errors


def verify_sheet_preservation(wb, expected_sheets: List[str]) -> List[str]:
    """Check that all expected sheets still exist."""
    return [s for s in expected_sheets if s not in wb.sheetnames]


def verify_values_outside_range(
    ws_orig, ws_out,
    protected_rows: List[int],
    col_start: int = 1,
    col_end: int = 8,
) -> List[str]:
    """Verify that cells outside the target range are unchanged."""
    errors = []
    for r in protected_rows:
        for c in range(col_start, col_end + 1):
            ov = ws_orig.cell(row=r, column=c).value
            nv = ws_out.cell(row=r, column=c).value
            if ov != nv:
                cl = get_column_letter(c)
                errors.append(f"{cl}{r}: expected {ov!r}, got {nv!r}")
    return errors


# ---------------------------------------------------------------------------
# CLI entry (for quick testing)
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    # Quick demo: load, detect blocks, evaluate, sort, save
    input_path = sys.argv[1] if len(sys.argv) > 1 else "/app/input.xlsx"
    wb = load_workbook(input_path)
    ws = wb["Sheet1"]

    # Auto-detect column range from data
    max_col = ws.max_column or 8
    max_row = ws.max_row or 85

    blocks = detect_blocks(ws, col_start=1, col_end=min(max_col, 8))
    print(f"Detected {len(blocks)} blocks: {blocks}")

    # Evaluate formulas in D and F
    all_rows = []
    for start, end in blocks:
        all_rows.extend(range(start, end + 1))
    evaluate_formulas_in_range(ws, all_rows, ["D", "F"], col_start=1, col_end=min(max_col, 8))

    # Sort each block by column D (col 4) descending
    for start, end in blocks:
        sort_block_rows(ws, list(range(start, end + 1)), key_col=4,
                        descending=True, col_start=1, col_end=min(max_col, 8))

    output_path = "/app/output.xlsx"
    wb.save(output_path)
    print(f"Saved to {output_path}")

    # Verify
    errs = verify_no_formulas(ws, all_rows, ["D", "F"])
    errs += verify_sort_order(ws, [list(range(s,e+1)) for s,e in blocks], key_col=4)
    blank = find_blank_rows(ws, col_start=1, col_end=min(max_col, 8))
    errs += verify_block_structure(ws, blank, col_start=1, col_end=min(max_col, 8))
    if errs:
        print(f"ERRORS: {errs}")
    else:
        print("All checks passed.")
