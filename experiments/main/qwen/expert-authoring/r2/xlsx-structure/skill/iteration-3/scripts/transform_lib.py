#!/usr/bin/env python3
"""
Reusable building blocks for xlsx-structure transformations.

Import and compose these functions rather than writing from scratch.
Demonstrates: block detection, formula evaluation, style snapshot/restore,
row writing, and self-verification.

Usage as library:
    from transform_lib import (
        detect_blocks, evaluate_cell, evaluate_row,
        snapshot_row, write_evaluated_row, verify_no_formulas
    )

Usage as standalone (block-local sort example):
    python transform_lib.py [input.xlsx] [output.xlsx]
"""

import sys
import copy
import re
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter


# ---------------------------------------------------------------------------
# Block Detection
# ---------------------------------------------------------------------------
def detect_blocks(ws, min_r, max_r, min_c, max_c):
    """Return list of (start_row, end_row) for contiguous non-blank blocks.

    A row is blank if every cell in columns min_c..max_c is None.
    """
    blank_rows = []
    for r in range(min_r, max_r + 1):
        if all(ws.cell(r, c).value is None for c in range(min_c, max_c + 1)):
            blank_rows.append(r)

    blocks = []
    start = min_r
    for br in blank_rows:
        if br > start:
            blocks.append((start, br - 1))
        start = br + 1
    if start <= max_r and any(
        ws.cell(start, c).value is not None for c in range(min_c, max_c + 1)
    ):
        blocks.append((start, max_r))
    return blocks


# ---------------------------------------------------------------------------
# Cell Value Evaluation
# ---------------------------------------------------------------------------
_CELL_REF_RE = re.compile(r'([A-Z]+)(\d+)')


def evaluate_cell(cell_value, ref_values=None):
    """Evaluate a cell value to a number when possible.

    Handles:
      - None → None
      - int/float → float
      - str starting with '=' → evaluate as arithmetic expression
      - numeric-looking str → float
      - other str → return as-is (opaque text)

    ref_values: dict mapping cell references (e.g. 'B1') to numeric values.
                Used to substitute cell references in formulas.
    """
    if cell_value is None:
        return None
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    if isinstance(cell_value, str):
        s = cell_value.strip()
        if s.startswith('='):
            expr = s[1:]
            if ref_values:
                # Replace cell references longest-first to avoid partial matches
                for ref in sorted(ref_values, key=len, reverse=True):
                    expr = expr.replace(ref, str(ref_values[ref]))
            # Remove any remaining unreplaced cell references (treat as 0)
            expr = _CELL_REF_RE.sub('0', expr)
            try:
                return float(eval(expr, {"__builtins__": {}}, {}))
            except Exception:
                return None
        try:
            return float(s)
        except ValueError:
            return s
    return cell_value


def evaluate_row(ws, row, min_col, max_col):
    """Evaluate all cells in a row left-to-right.

    Returns (values_list, ref_values_dict).
    ref_values_dict maps cell references to evaluated numeric values,
    enabling downstream formulas in the same row to reference upstream results.
    """
    ref_values = {}
    values = []
    for c in range(min_col, max_col + 1):
        cell_val = ws.cell(row, c).value
        val = evaluate_cell(cell_val, ref_values)
        ref = f"{get_column_letter(c)}{row}"
        ref_values[ref] = val if isinstance(val, (int, float)) else 0
        values.append(val)
    return values, ref_values


# ---------------------------------------------------------------------------
# Style Snapshot / Restore
# ---------------------------------------------------------------------------
def snapshot_row(ws, row, min_col, max_col):
    """Capture value + full style for every cell in a row.

    Returns list of snapshot dicts, one per column.
    """
    cells = []
    for c in range(min_col, max_col + 1):
        cell = ws.cell(row, c)
        cells.append({
            'value':         cell.value,
            'font':          copy.copy(cell.font),
            'fill':          copy.copy(cell.fill),
            'border':        copy.copy(cell.border),
            'alignment':     copy.copy(cell.alignment),
            'number_format': cell.number_format,
            'protection':    copy.copy(cell.protection),
        })
    return cells


def write_evaluated_row(ws, row, evaluated_values, style_snapshots, min_col):
    """Write evaluated values to a row, preserving the style from snapshots.

    evaluated_values: list of evaluated cell values (from evaluate_row).
    style_snapshots: list of snapshot dicts (from snapshot_row of the SOURCE row).
    Styles travel WITH the data — use the snapshot from the row that is moving here.
    """
    for i, (val, snap) in enumerate(zip(evaluated_values, style_snapshots)):
        cell = ws.cell(row, min_col + i)
        # Use evaluated value if it was a formula/text-number, else keep original
        if isinstance(val, (int, float)) and isinstance(snap['value'], str):
            cell.value = val
        else:
            cell.value = snap['value']
        cell.font          = snap['font']
        cell.fill          = snap['fill']
        cell.border        = snap['border']
        cell.alignment     = snap['alignment']
        cell.number_format = snap['number_format']
        cell.protection    = snap['protection']


def clear_row(ws, row, min_col, max_col):
    """Set all cells in a row to None (for compaction/filter patterns)."""
    for c in range(min_col, max_col + 1):
        ws.cell(row, c).value = None


# ---------------------------------------------------------------------------
# Sort Helper
# ---------------------------------------------------------------------------
def sort_block_rows(block_rows, sort_keys, descending=True):
    """Stable-sort block_rows by sort_keys.

    block_rows: list of row indices [r1, r2, ...].
    sort_keys:  list of numeric keys aligned with block_rows.
    Returns: list of (original_index_in_block, row_index) in sorted order.
    """
    indexed = list(enumerate(zip(sort_keys, block_rows)))
    if descending:
        indexed.sort(key=lambda x: (-x[1][0], x[0]))
    else:
        indexed.sort(key=lambda x: (x[1][0], x[0]))
    return [(orig_idx, row_idx) for orig_idx, (_, row_idx) in indexed]


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------
def verify_no_formulas(ws, min_r, max_r, min_c, max_c):
    """Check that no unevaluated formulas remain in the working range."""
    issues = []
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith('='):
                issues.append((r, c, v))
    return issues


def verify_sort_order(ws, blocks, sort_col_idx, descending=True):
    """Check that sort column values are monotonic within each block."""
    issues = []
    for b_start, b_end in blocks:
        prev = float('inf') if descending else float('-inf')
        for r in range(b_start, b_end + 1):
            val = ws.cell(r, sort_col_idx).value
            if val is None:
                continue
            num = float(val) if isinstance(val, (int, float)) else None
            if num is not None:
                if descending and num > prev + 1e-9:
                    issues.append((r, num, prev))
                elif not descending and num < prev - 1e-9:
                    issues.append((r, num, prev))
                prev = num
    return issues


def verify_blank_rows(ws, blocks, min_r, max_r, min_c, max_c):
    """Check that separator rows between blocks are blank."""
    expected_blanks = set()
    block_set = set()
    for s, e in blocks:
        for r in range(s, e + 1):
            block_set.add(r)
    for r in range(min_r, max_r + 1):
        if r not in block_set:
            expected_blanks.add(r)
    issues = []
    for r in expected_blanks:
        if any(ws.cell(r, c).value is not None for c in range(min_c, max_c + 1)):
            issues.append(r)
    return issues


# ---------------------------------------------------------------------------
# Example: block-local sort (standalone usage)
# ---------------------------------------------------------------------------
def solve_block_sort(input_path, output_path, sheet_name="Sheet1",
                     work_range="A1:H85", sort_col="D", descending=True):
    """Complete block-local sort example.

    This function demonstrates how to compose the library functions.
    Adapt its parameters and logic for other transformation types.
    """
    wb = openpyxl.load_workbook(input_path)
    ws = wb[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(work_range)

    # Convert sort column letter to index
    sort_col_idx = 0
    for ch in sort_col.upper():
        sort_col_idx = sort_col_idx * 26 + (ord(ch) - ord('A') + 1)

    # 1. Detect blocks
    blocks = detect_blocks(ws, min_row, max_row, min_col, max_col)
    print(f"Detected {len(blocks)} blocks: {blocks}")

    # 2. Snapshot all styles in working range
    all_snapshots = {}
    for r in range(min_row, max_row + 1):
        all_snapshots[r] = snapshot_row(ws, r, min_col, max_col)

    # 3. Process each block
    for b_start, b_end in blocks:
        block_rows = list(range(b_start, b_end + 1))

        # Evaluate each row and extract sort keys
        evaluated = {}  # row_idx -> (values, ref_values)
        sort_keys = []
        for r in block_rows:
            values, refs = evaluate_row(ws, r, min_col, max_col)
            evaluated[r] = (values, refs)
            sort_col_ref = f"{sort_col}{r}"
            key = refs.get(sort_col_ref, 0.0)
            if not isinstance(key, (int, float)):
                try:
                    key = float(key)
                except (TypeError, ValueError):
                    key = 0.0
            sort_keys.append(key)

        # Sort
        sorted_order = sort_block_rows(block_rows, sort_keys, descending)

        # Write back: sorted_order[i] tells which original row goes to position i
        for new_pos, (orig_idx, orig_row) in enumerate(sorted_order):
            target_row = block_rows[new_pos]
            values, _ = evaluated[orig_row]
            write_evaluated_row(ws, target_row, values,
                                all_snapshots[orig_row], min_col)

    # 4. Save
    wb.save(output_path)
    print(f"Saved to {output_path}")

    # 5. Verify
    out_wb = openpyxl.load_workbook(output_path)
    out_ws = out_wb[sheet_name]

    formula_issues = verify_no_formulas(out_ws, min_row, max_row, min_col, max_col)
    if formula_issues:
        for r, c, v in formula_issues:
            print(f"  WARNING: Formula at row {r}, col {c}: {v}")
    else:
        print("  No unevaluated formulas.")

    sort_issues = verify_sort_order(out_ws, blocks, sort_col_idx, descending)
    if sort_issues:
        for r, val, prev in sort_issues:
            print(f"  WARNING: Sort violation at row {r}: {val} after {prev}")
    else:
        print("  Sort order correct.")

    blank_issues = verify_blank_rows(out_ws, blocks, min_row, max_row, min_col, max_col)
    if blank_issues:
        print(f"  WARNING: Non-blank separator rows: {blank_issues}")
    else:
        print("  Blank rows correct.")

    out_wb.close()
    print("Done.")


if __name__ == '__main__':
    input_path  = sys.argv[1] if len(sys.argv) > 1 else "/app/input.xlsx"
    output_path = sys.argv[2] if len(sys.argv) > 2 else "/app/output.xlsx"
    # Adapt these parameters to match the specific task
    solve_block_sort(
        input_path=input_path,
        output_path=output_path,
        sheet_name="Sheet1",
        work_range="A1:H85",
        sort_col="D",
        descending=True,
    )
