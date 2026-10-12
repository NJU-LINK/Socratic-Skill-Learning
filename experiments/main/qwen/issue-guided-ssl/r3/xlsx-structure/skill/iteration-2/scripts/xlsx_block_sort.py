#!/usr/bin/env python3
"""
xlsx_block_sort.py — Block-local sorting of xlsx data with formula evaluation.

Solves the v01 anchor task: sort variable-length blocks (separated by blank
rows) by evaluated numeric value of column D descending, with stable ordering
for ties. Evaluates formula strings in columns D-F using row-local arithmetic
with proper dependency order. Writes numeric values (not formula strings).
Preserves all workbook metadata outside A1:H85.

Two formula evaluators are provided:
  - evaluate_row_generic(): Auto-detects formulas and resolves dependencies
    topologically. Works for ANY column layout. RECOMMENDED for new variants.
  - evaluate_row_v01(): Hardcoded for the v01 anchor case (D=B+C, F=D+E).
    Faster but only works for that specific layout.

Usage:
    python xlsx_block_sort.py [input.xlsx] [output.xlsx]

Defaults: input=/app/input.xlsx, output=/app/output.xlsx
"""
import re
import sys
import openpyxl


def eval_formula(formula_str, local_vals):
    """Evaluate a row-local arithmetic formula string.

    Args:
        formula_str: cell value — either a formula string like '=B1+C1' or a
                     plain numeric value (int/float/None).
        local_vals: dict mapping column letters ('A'-'Z') to numeric values
                    for this row. Must include all referenced columns.

    Returns:
        float: the evaluated numeric result, rounded to 10 decimal places
               to remove floating-point noise.
    """
    if not isinstance(formula_str, str) or not formula_str.startswith('='):
        if formula_str is None:
            return 0.0
        return float(formula_str)

    expr = formula_str[1:]  # strip leading '='

    def replace_ref(match):
        col = match.group(1)
        return str(local_vals.get(col, 0.0))

    expr = re.sub(r'([A-Z]+)\d+', replace_ref, expr)
    result = float(eval(expr))
    return round(result, 10)


def detect_blocks(ws, min_row, max_row, min_col, max_col):
    """Detect variable-length blocks separated by blank rows.

    Returns:
        blocks: list of lists of (row_index, [cell_values])
        blank_positions: set of blank row indices within the range
        rows_data: list of (row_index, values_or_None) for all rows
    """
    blank_positions = set()
    rows_data = []

    for row_idx in range(min_row, max_row + 1):
        vals = []
        all_none = True
        for col_idx in range(min_col, max_col + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            vals.append(v)
            if v is not None:
                all_none = False
        if all_none:
            blank_positions.add(row_idx)
            rows_data.append((row_idx, None))
        else:
            rows_data.append((row_idx, vals))

    blocks = []
    current = []
    for row_idx, vals in rows_data:
        if vals is None:
            if current:
                blocks.append(current)
                current = []
        else:
            current.append((row_idx, vals))
    if current:
        blocks.append(current)

    return blocks, blank_positions, rows_data


def evaluate_row_generic(vals, min_col_idx=0):
    """Generic formula evaluator that auto-detects formulas and resolves dependencies.

    This is the RECOMMENDED approach for new variants. It works for any column layout
    by automatically detecting which cells contain formula strings and topologically
    sorting them based on their cell reference dependencies.

    Args:
        vals: list of cell values for the row (indexed by column position)
        min_col_idx: the column index that corresponds to column 'A' (usually 0)

    Returns:
        list of evaluated values with formulas replaced by numeric results.
        Non-formula cells retain their original values.

    Algorithm:
        1. Scan all cells: raw numeric values go into local map immediately.
        2. Identify formula cells and extract their column references.
        3. Iteratively evaluate formulas whose dependencies are all satisfied.
        4. Repeat until all formulas are resolved or max iterations reached.
    """
    # Build local value map from raw (non-formula) values.
    # Numeric values get their float value; None gets 0.0 (for formula refs);
    # text/other types are NOT added (they can't be referenced in arithmetic).
    local = {}
    for i, v in enumerate(vals):
        col_letter = chr(ord('A') + i + min_col_idx)
        if v is not None and not (isinstance(v, str) and v.startswith('=')):
            if isinstance(v, (int, float)):
                local[col_letter] = float(v)
        elif v is None:
            # None cells can be referenced by formulas as 0.0
            local[col_letter] = 0.0
        # Text values (non-formula strings) are NOT added to local map

    # Identify formula cells and their column dependencies
    formula_cells = []
    for i, v in enumerate(vals):
        col_letter = chr(ord('A') + i + min_col_idx)
        if isinstance(v, str) and v.startswith('='):
            refs = set(re.findall(r'([A-Z]+)\d+', v[1:]))
            formula_cells.append((i, col_letter, v, refs))

    # Topological evaluation: resolve formulas whose deps are all known
    evaluated = dict(local)
    remaining = list(formula_cells)
    max_iterations = len(remaining) * 2  # safety bound
    iteration = 0

    while remaining and iteration < max_iterations:
        iteration += 1
        for item in list(remaining):
            i, col_letter, formula_str, refs = item
            if all(ref in evaluated for ref in refs):
                val = eval_formula(formula_str, evaluated)
                evaluated[col_letter] = val
                remaining.remove(item)

    # Build result: only replace cells that were originally formula strings
    result = list(vals)
    formula_indices = {item[0] for item in formula_cells}
    for i in formula_indices:
        col_letter = chr(ord('A') + i + min_col_idx)
        if col_letter in evaluated:
            result[i] = evaluated[col_letter]

    return result


def evaluate_row_v01(vals):
    """Optimized evaluator for the v01 anchor case (D=B+C, F=D+E).

    This is faster than the generic evaluator but only works when:
    - B (idx 1), C (idx 2), E (idx 4) are raw numeric
    - D (idx 3) contains =B+C
    - F (idx 5) contains =D+E
    - A (idx 0) is text, G (idx 6) and H (idx 7) are integers

    For other variants, use evaluate_row_generic() instead.
    """
    b_val = float(vals[1]) if vals[1] is not None else 0.0
    c_val = float(vals[2]) if vals[2] is not None else 0.0
    e_val = float(vals[4]) if vals[4] is not None else 0.0

    local = {'B': b_val, 'C': c_val, 'E': e_val}
    d_eval = eval_formula(vals[3], local)
    local['D'] = d_eval
    f_eval = eval_formula(vals[5], local)

    return [vals[0], b_val, c_val, d_eval, e_val, f_eval, vals[6], vals[7]]


def block_sort_transform(blocks, sort_col_idx=3, descending=True, evaluator=None):
    """Sort each block by the evaluated value at sort_col_idx.

    Uses Python's stable sort to preserve original order for ties.

    Args:
        blocks: list of blocks, each block is list of (row_idx, vals)
        sort_col_idx: 0-based column index to sort by
        descending: True for descending order
        evaluator: function to evaluate a row. Defaults to evaluate_row_generic.
                   For v01 anchor case, pass evaluate_row_v01 for speed.

    Returns:
        list of blocks, each block is list of (row_idx, orig_vals, ev_vals)
    """
    if evaluator is None:
        evaluator = evaluate_row_generic

    sorted_blocks = []
    for block in blocks:
        evaluated = []
        for row_idx, vals in block:
            ev = evaluator(vals)
            evaluated.append((row_idx, vals, ev))

        evaluated.sort(key=lambda x: x[2][sort_col_idx], reverse=descending)
        sorted_blocks.append(evaluated)

    return sorted_blocks


def write_back(ws, sorted_blocks, rows_data, min_col=1):
    """Write sorted/evaluated values back to worksheet, preserving blank rows.

    Only modifies .value on cells — styles remain on their positions.
    """
    data_positions = [r for r, v in rows_data if v is not None]

    all_sorted_rows = []
    for block in sorted_blocks:
        for _, _, ev_vals in block:
            all_sorted_rows.append(ev_vals)

    for i, pos in enumerate(data_positions):
        ev_vals = all_sorted_rows[i]
        for col_idx, value in enumerate(ev_vals):
            ws.cell(row=pos, column=min_col + col_idx).value = value


def main(input_path='/app/input.xlsx', output_path='/app/output.xlsx'):
    # Load workbook (data_only=False to get formula strings)
    wb = openpyxl.load_workbook(input_path, data_only=False)
    ws = wb['Sheet1']

    # Configuration for the transformation range
    MIN_ROW, MAX_ROW = 1, 85
    MIN_COL, MAX_COL = 1, 8  # A-H
    SORT_COL_IDX = 3  # column D (0-based within the range)

    # Detect blocks
    blocks, blank_positions, rows_data = detect_blocks(
        ws, MIN_ROW, MAX_ROW, MIN_COL, MAX_COL
    )

    # Evaluate formulas and sort each block
    sorted_blocks = block_sort_transform(blocks, sort_col_idx=SORT_COL_IDX)

    # Write back
    write_back(ws, sorted_blocks, rows_data, min_col=MIN_COL)

    # Save (everything outside A1:H85 is untouched)
    wb.save(output_path)
    print(f"Saved {output_path}")


if __name__ == '__main__':
    inp = sys.argv[1] if len(sys.argv) > 1 else '/app/input.xlsx'
    out = sys.argv[2] if len(sys.argv) > 2 else '/app/output.xlsx'
    main(inp, out)
