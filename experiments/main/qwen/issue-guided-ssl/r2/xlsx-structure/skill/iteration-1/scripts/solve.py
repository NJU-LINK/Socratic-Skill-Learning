#!/usr/bin/env python3
"""
XLSX Block-Sort Solver

Reads an input workbook, detects variable-length blocks of non-blank rows
separated by blank rows within a target range, evaluates any row-local
arithmetic formulas in dependency order, sorts each block by a specified
column's evaluated numeric value (descending by default, stable), and writes
the result back in-place — preserving all workbook state outside the target
range.

Usage:
    python solve.py [--input INPUT] [--output OUTPUT] [--sheet SHEET]
                    [--range RANGE] [--sort-col SORT_COL] [--desc]
                    [--formula-cols FORMULA_COLS]

Defaults match the v01 task:
    --input /app/input.xlsx
    --output /app/output.xlsx
    --sheet Sheet1
    --range A1:H85
    --sort-col D
    --desc
    --formula-cols D,E,F  (evaluated in that dependency order)
"""

import argparse
import copy
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# ── Column letter → 0-based index ──────────────────────────────────────────
COL_LETTER_TO_IDX = {get_column_letter(i): i - 1 for i in range(1, 27)}


def parse_range(range_str):
    """Parse 'A1:H85' → (min_row, min_col_0based, max_row, max_col_0based)."""
    m = re.match(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', range_str)
    if not m:
        raise ValueError(f"Invalid range: {range_str}")
    min_col = column_index_from_string(m.group(1)) - 1
    max_col = column_index_from_string(m.group(3)) - 1
    return int(m.group(2)), min_col, int(m.group(4)), max_col


def is_blank_row(ws, row_num, min_col_0, max_col_0):
    """A row is blank if every cell in the column range is None."""
    for c in range(min_col_0 + 1, max_col_0 + 2):  # openpyxl is 1-based
        if ws.cell(row=row_num, column=c).value is not None:
            return False
    return True


def detect_blocks(ws, min_row, max_row, min_col_0, max_col_0):
    """
    Detect contiguous non-blank row blocks within [min_row, max_row].
    Returns list of (start_row, end_row) inclusive tuples.
    """
    blocks = []
    block_start = None
    for r in range(min_row, max_row + 1):
        blank = is_blank_row(ws, r, min_col_0, max_col_0)
        if not blank:
            if block_start is None:
                block_start = r
        else:
            if block_start is not None:
                blocks.append((block_start, r - 1))
                block_start = None
    if block_start is not None:
        blocks.append((block_start, max_row))
    return blocks


def read_block(ws, start_row, end_row, min_col_0, max_col_0):
    """Read a block of rows into a list of lists (raw cell values)."""
    rows = []
    for r in range(start_row, end_row + 1):
        row_data = []
        for c in range(min_col_0, max_col_0 + 1):
            row_data.append(ws.cell(row=r, column=c + 1).value)
        rows.append(row_data)
    return rows


def to_numeric(val):
    """Convert a cell value to float for arithmetic. None → 0."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        val = val.strip()
        try:
            return float(val)
        except ValueError:
            return 0.0
    return 0.0


def eval_formula(formula_str, row_values, col_map):
    """
    Evaluate a row-local arithmetic formula.
    formula_str: e.g., '=B1+C1' or 'B1+C1'
    row_values: list of values for this row (0-indexed by column)
    col_map: dict mapping column letter to 0-based index
    Returns float result, or None if not evaluable.
    """
    expr = formula_str.strip()
    if expr.startswith('='):
        expr = expr[1:]
    expr = expr.strip()

    def replace_ref(m):
        col_letter = m.group(1)
        if col_letter in col_map:
            val = row_values[col_map[col_letter]]
            return str(to_numeric(val))
        return m.group(0)  # leave unknown refs as-is

    expr = re.sub(r'([A-Z]+)\d+', replace_ref, expr)

    # Safe eval: only allow arithmetic characters
    if re.match(r'^[\d\s\+\-\*\/\(\)\.eE]+$', expr):
        try:
            return float(eval(expr))
        except Exception:
            return None
    return None


def is_formula(val):
    """Check if a value is a formula string or text-representing-formula."""
    if not isinstance(val, str):
        return False
    s = val.strip()
    if s.startswith('='):
        return True
    # Text that represents a formula: contains cell refs and operators
    # e.g., "B1+C1" without the '='
    if re.search(r'[A-Z]+\d+\s*[\+\-\*\/]', s) or re.search(r'[\+\-\*\/]\s*[A-Z]+\d+', s):
        return True
    return False


def evaluate_formulas(block_rows, formula_col_indices, col_letter_map):
    """
    Evaluate formula columns in dependency order.
    Updates block_rows in-place so later columns see earlier evaluated values.
    """
    for col_idx in formula_col_indices:
        for row in block_rows:
            val = row[col_idx]
            if is_formula(val):
                result = eval_formula(val, row, col_letter_map)
                if result is not None:
                    row[col_idx] = result


def stable_sort_block(block_rows, sort_col_idx, descending=True):
    """
    Stable sort rows by the numeric value in sort_col_idx.
    Uses enumerate to preserve original order for ties.
    """
    indexed = list(enumerate(block_rows))
    if descending:
        indexed.sort(key=lambda x: (-to_numeric(x[1][sort_col_idx]), x[0]))
    else:
        indexed.sort(key=lambda x: (to_numeric(x[1][sort_col_idx]), x[0]))
    return [row for _, row in indexed]


def solve(input_path, output_path, sheet_name, range_str, sort_col_letter,
          descending, formula_col_letters):
    """Main solver pipeline."""
    # Parse parameters
    min_row, min_col_0, max_row, max_col_0 = parse_range(range_str)
    sort_col_idx = column_index_from_string(sort_col_letter) - 1
    formula_col_indices = [column_index_from_string(c) - 1 for c in formula_col_letters]
    col_letter_map = {get_column_letter(i + 1): i for i in range(26)}

    # Load workbook (NOT data_only — we need formula strings)
    wb = openpyxl.load_workbook(input_path)
    ws = wb[sheet_name]

    # Detect blocks
    blocks = detect_blocks(ws, min_row, max_row, min_col_0, max_col_0)

    # Process each block
    for start_row, end_row in blocks:
        # Read raw values
        block_rows = read_block(ws, start_row, end_row, min_col_0, max_col_0)

        # Evaluate formulas in dependency order
        evaluate_formulas(block_rows, formula_col_indices, col_letter_map)

        # Stable sort by sort column
        sorted_rows = stable_sort_block(block_rows, sort_col_idx, descending)

        # Write back in-place (only .value changes, styles preserved)
        for i, row_data in enumerate(sorted_rows):
            for j, val in enumerate(row_data):
                ws.cell(row=start_row + i, column=min_col_0 + j + 1).value = val

    # Save
    wb.save(output_path)
    print(f"Saved to {output_path}")


def main():
    parser = argparse.ArgumentParser(description="XLSX Block-Sort Solver")
    parser.add_argument('--input', default='/app/input.xlsx')
    parser.add_argument('--output', default='/app/output.xlsx')
    parser.add_argument('--sheet', default='Sheet1')
    parser.add_argument('--range', default='A1:H85')
    parser.add_argument('--sort-col', default='D')
    parser.add_argument('--desc', action='store_true', default=True)
    parser.add_argument('--asc', action='store_true', default=False)
    parser.add_argument('--formula-cols', default='D,E,F')
    args = parser.parse_args()

    descending = not args.asc
    formula_cols = [c.strip() for c in args.formula_cols.split(',')]

    solve(args.input, args.output, args.sheet, args.range,
          args.sort_col, descending, formula_cols)


if __name__ == '__main__':
    main()
