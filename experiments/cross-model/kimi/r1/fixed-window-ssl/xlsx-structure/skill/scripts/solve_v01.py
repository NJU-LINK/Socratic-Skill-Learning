"""
Self-test: Solve v01 using only the packaged skill API.

This script:
1. Reads /app/input.xlsx
2. Uses xlsx_eval, blocks, rewrite modules
3. Sorts blank-separated blocks by column D descending
4. Writes /app/output_test.xlsx
5. Runs verify.run_verification
6. Compares against known-good /app/output.xlsx

Usage:
    cd /app
    python skill/xlsx-structure/scripts/solve_v01.py

Exit code 0 on success, 1 on failure.
"""

import sys
import os

# Add skill scripts to path
script_dir = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, script_dir)

from openpyxl import load_workbook
from xlsx_eval import evaluate_cell, clear_cache
from blocks import find_blank_separated_blocks
from rewrite import sort_block_by_key, rewrite_block
from verify import run_verification

INPUT_PATH = '/app/input.xlsx'
OUTPUT_PATH = '/app/output_test.xlsx'
REFERENCE_PATH = '/app/output.xlsx'


def main():
    print("=== XLSX Structure Skill Self-Test (v01) ===\n")
    
    # Clear any stale cache
    clear_cache()
    
    # Load workbook
    print(f"Loading {INPUT_PATH}...")
    wb = load_workbook(INPUT_PATH, data_only=False)
    ws = wb['Sheet1']
    
    # Find blocks
    print("Finding blank-separated blocks...")
    block_rows_list = find_blank_separated_blocks(ws, 1, 85, max_col=8)
    print(f"  Found {len(block_rows_list)} blocks: " + 
          ", ".join(f"{b[0]}-{b[-1]}" for b in block_rows_list))
    
    # Compute blank rows (complement of blocks within 1-85)
    all_rows = set(range(1, 86))
    block_rows_set = set()
    for b in block_rows_list:
        block_rows_set.update(b)
    blank_rows = sorted(all_rows - block_rows_set)
    print(f"  Blank separator rows: {blank_rows}")
    
    # Process each block
    print("\nProcessing blocks...")
    for i, block_rows in enumerate(block_rows_list, 1):
        print(f"  Block {i}: rows {block_rows[0]}-{block_rows[-1]} ({len(block_rows)} rows)")
        
        # Sort by column D descending
        sorted_values = sort_block_by_key(
            ws, block_rows,
            key_func=lambda r: evaluate_cell(ws, r, 4),  # Column D
            value_func=lambda r, c: evaluate_cell(ws, r, c) if c in [4, 5, 6] 
                                    else ws.cell(row=r, column=c).value,
            reverse=True
        )
        
        # Write back in place
        rewrite_block(ws, block_rows, sorted_values)
    
    # Save
    print(f"\nSaving to {OUTPUT_PATH}...")
    wb.save(OUTPUT_PATH)
    
    # Verify
    print("\nRunning verification...")
    blocks = [(b[0], b[-1]) for b in block_rows_list]
    
    try:
        run_verification(INPUT_PATH, OUTPUT_PATH, blocks, blank_rows,
                        'Sheet1', sort_key_col=4, reverse=True)
        print("  ✓ All verification checks passed")
    except AssertionError as e:
        print(f"  ✗ Verification failed: {e}")
        return 1
    
    # Compare with reference
    if os.path.exists(REFERENCE_PATH):
        print(f"\nComparing with reference {REFERENCE_PATH}...")
        ref_wb = load_workbook(REFERENCE_PATH, data_only=True)
        test_wb = load_workbook(OUTPUT_PATH, data_only=True)
        ref_ws = ref_wb['Sheet1']
        test_ws = test_wb['Sheet1']
        
        differences = []
        for r in range(1, 86):
            for c in range(1, 9):
                ref_v = ref_ws.cell(row=r, column=c).value
                test_v = test_ws.cell(row=r, column=c).value
                if ref_v != test_v:
                    differences.append((r, c, ref_v, test_v))
        
        if differences:
            print(f"  ✗ {len(differences)} differences from reference:")
            for d in differences[:5]:
                print(f"    Row {d[0]}, Col {d[1]}: ref={d[2]!r}, test={d[3]!r}")
            return 1
        else:
            print("  ✓ Matches reference exactly")
    else:
        print(f"\nNote: Reference {REFERENCE_PATH} not found for comparison")
    
    print("\n=== Self-test PASSED ===")
    return 0


if __name__ == '__main__':
    sys.exit(main())
