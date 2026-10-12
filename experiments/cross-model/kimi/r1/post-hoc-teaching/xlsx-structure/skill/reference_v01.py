"""
Reference implementation for xlsx-structure v01 task.

Task: Block-local sort by evaluated column D in descending order.

- Rows form variable-length blocks separated by blank rows
- Within each block, sort by evaluated numeric value of column D descending
- Preserve original relative order for ties (Python sort is stable)
- Evaluate formulas in columns D-F and write numeric values
- Preserve everything outside A1:H85, all styles, workbook properties

STYLE NOTE: This uses position-attached preservation (write values only,
do not move styles with rows). Styles are captured only for non-uniformity
detection, not for re-application.
"""

import sys
from pathlib import Path

# Add scripts directory to path
scripts_dir = Path(__file__).parent / 'scripts'
if scripts_dir.exists():
    sys.path.insert(0, str(scripts_dir))

try:
    from xlsx_structure import (
        make_evaluator,
        capture_all_cell_values,
        detect_blocks,
        write_value,
        verify_preservation,
        assert_no_formula_strings,
        is_blank_row,
    )
except ImportError:
    # Fallback: define locally if scripts not available
    from openpyxl.utils import get_column_letter
    from copy import copy
    import re
    
    CELL_REF = re.compile(r"(?<![A-Za-z0-9_$])\$?([A-Z]{1,3})\$?(\d{1,7})(?![A-Za-z0-9_(])", re.VERBOSE)
    
    def make_evaluator(cell_values):
        memo, in_progress = {}, set()
        def to_number(v):
            if v is None: return 0.0
            if isinstance(v, bool): return float(v)
            if isinstance(v, (int, float)): return float(v)
            if isinstance(v, str):
                try: return float(v.strip())
                except ValueError: return 0.0
            return 0.0
        def eval_cell(addr):
            if addr in memo: return memo[addr]
            v = cell_values.get(addr)
            if isinstance(v, str) and v.startswith('='):
                if addr in in_progress:
                    raise ValueError(f"circular reference at {addr}")
                in_progress.add(addr)
                expr = v[1:]
                expr = CELL_REF.sub(lambda m: repr(eval_cell(f"{m.group(1)}{m.group(2)}")), expr)
                try:
                    result = float(eval(expr, {"__builtins__": {}}, {}))
                except Exception as e:
                    raise ValueError(f"cannot evaluate {v!r} at {addr}: {e}")
                in_progress.discard(addr)
            else:
                result = to_number(v)
            memo[addr] = result
            return result
        return eval_cell
    
    def capture_all_cell_values(ws):
        from openpyxl.utils import get_column_letter
        cell_values = {}
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                col_letter = get_column_letter(col)
                cell_values[f"{col_letter}{row}"] = ws.cell(row=row, column=col).value
        return cell_values
    
    def is_blank_row(ws, row_num, max_col):
        for col in range(1, max_col + 1):
            if ws.cell(row=row_num, column=col).value is not None:
                return False
        return True
    
    def detect_blocks(ws, min_row, max_row, max_col):
        blocks, current_block = [], []
        for row_num in range(min_row, max_row + 1):
            if is_blank_row(ws, row_num, max_col):
                if current_block:
                    blocks.append(current_block)
                    current_block = []
            else:
                current_block.append(row_num)
        if current_block:
            blocks.append(current_block)
        return blocks
    
    def write_value(cell, value, eval_cell=None, addr=None):
        if isinstance(value, str) and value.startswith('=') and eval_cell is not None and addr is not None:
            cell.value = eval_cell(addr)
        else:
            cell.value = value
    
    def verify_preservation(input_path, output_path, sheet_name, transform_region):
        from openpyxl import load_workbook
        wi = load_workbook(input_path)
        wo = load_workbook(output_path)
        assert wi.sheetnames == wo.sheetnames, "Sheet list changed"
        si, so = wi[sheet_name], wo[sheet_name]
        assert si.freeze_panes == so.freeze_panes, "Freeze panes changed"
        r1, r2, c1, c2 = transform_region
        for row in range(1, si.max_row + 1):
            for col in range(1, si.max_column + 1):
                if r1 <= row <= r2 and c1 <= col <= c2:
                    continue
                assert si.cell(row, col).value == so.cell(row, col).value, \
                    f"Outside-region change at R{row}C{col}"
        print("Verification passed")
        return True
    
    def assert_no_formula_strings(ws, min_row, max_row, min_col, max_col):
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                val = ws.cell(row, col).value
                if isinstance(val, str) and val.startswith('='):
                    raise AssertionError(f"Cell R{row}C{col} still has formula: {val!r}")


from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


# Configuration for this specific task
CONFIG = {
    'input_path': '/app/input.xlsx',
    'output_path': '/app/output.xlsx',
    'sheet_name': 'Sheet1',
    'transform_region': (1, 85, 1, 8),  # A1:H85
    'sort_column': 4,  # Column D
    'sort_descending': True,
    'formula_columns': [4, 5, 6],  # Columns D-F
}


def check_style_uniformity(ws, rows, cols):
    """Check if all cells in region have identical styles.
    
    Returns True if uniform, False if non-uniform.
    Used for decision: position-attached vs move-with-row.
    """
    style_sigs = set()
    for row in rows:
        for col in cols:
            cell = ws.cell(row=row, column=col)
            if cell.has_style:
                # Simple signature: fill + font name
                sig = (
                    cell.fill.start_color.rgb if cell.fill and cell.fill.start_color else None,
                    cell.font.name if cell.font else None,
                )
                style_sigs.add(sig)
            else:
                style_sigs.add(None)
    return len(style_sigs) <= 1


def transform(config):
    """Main transformation function with position-attached style preservation."""
    wb = load_workbook(config['input_path'])
    ws = wb[config['sheet_name']]
    
    min_row, max_row, min_col, max_col = config['transform_region']
    sort_col = config['sort_column']
    
    # Check style uniformity before transformation
    all_rows = list(range(min_row, max_row + 1))
    all_cols = list(range(min_col, max_col + 1))
    is_uniform = check_style_uniformity(ws, all_rows, all_cols)
    if not is_uniform:
        print("WARNING: Region has non-uniform styles. Position-attached preservation may not be appropriate.")
        print("Consider whether the task requires moving styles with rows.")
    
    # Capture all cell values before any modification
    cell_values = capture_all_cell_values(ws)
    
    # Create evaluator
    eval_cell = make_evaluator(cell_values)
    
    # Identify blocks
    blocks = detect_blocks(ws, min_row, max_row, max_col)
    
    # Capture row data only (not styles - position-attached preservation)
    data_rows = {}
    for row_num in range(min_row, max_row + 1):
        if not is_blank_row(ws, row_num, max_col):
            row_data = []
            for col in range(min_col, max_col + 1):
                cell = ws.cell(row=row_num, column=col)
                row_data.append(cell.value)
            data_rows[row_num] = row_data
    
    # Build blocks with data
    blocks_with_data = []
    for block in blocks:
        block_data = [(rn, data_rows[rn]) for rn in block if rn in data_rows]
        if block_data:
            blocks_with_data.append(block_data)
    
    # Sort each block by evaluated sort column
    sorted_blocks = []
    for block in blocks_with_data:
        rows_with_key = []
        for row_num, row_data in block:
            cell_addr = f"{get_column_letter(sort_col)}{row_num}"
            sort_key = eval_cell(cell_addr)
            rows_with_key.append((row_num, row_data, sort_key))
        
        # Python sort is stable - preserves original order for ties
        rows_with_key.sort(key=lambda x: x[2], reverse=config['sort_descending'])
        sorted_blocks.append([(rn, rd) for rn, rd, _ in rows_with_key])
    
    # Clear transform region (preserves styles at positions)
    for row_num in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            ws.cell(row=row_num, column=col).value = None
    
    # Write sorted blocks using write_value (evaluates formulas)
    formula_cols = set(config['formula_columns'])
    output_row = min_row
    
    for block_idx, block in enumerate(sorted_blocks):
        for orig_row_num, row_data in block:
            for col_idx, value in enumerate(row_data):
                col = min_col + col_idx
                cell = ws.cell(row=output_row, column=col)
                
                addr = f"{get_column_letter(col)}{orig_row_num}"
                
                # Evaluate formulas in specified columns
                if col in formula_cols:
                    write_value(cell, value, eval_cell, addr)
                else:
                    write_value(cell, value)
            
            output_row += 1
        
        if block_idx < len(sorted_blocks) - 1:
            output_row += 1
    
    wb.save(config['output_path'])
    
    # Post-write verification
    assert_no_formula_strings(ws, min_row, output_row - 1, min_col, max_col)
    
    return output_row - 1


if __name__ == '__main__':
    last_row = transform(CONFIG)
    print(f"Transformation complete. Output written to row {last_row}")
    verify_preservation(
        CONFIG['input_path'],
        CONFIG['output_path'],
        CONFIG['sheet_name'],
        CONFIG['transform_region']
    )
