#!/usr/bin/env python3
"""
Family probe for the second executable semantic (filter + compaction).

Builds a synthetic in-family workbook in the v01 style (variable blocks
separated by one blank row, D='=B{r}+C{r}' / F='=D{r}+E{r}' formulas written by
openpyxl, hence NO cached values - the pipeline must evaluate from text),
runs scripts/filter_compact.py with a numeric keep predicate, and verifies the
output bit-exactly against an INDEPENDENT expectation computed from the input's
raw values (no pipeline functions are used for the expectation).

Assertions:
  * pipeline --check exits 0;
  * every kept row satisfies the predicate; dropped rows are gone;
  * survivors compacted from the span start with exactly one blank row between
    adjacent survivor groups; tail blank;
  * span cells match the independent expectation bit-exactly for floats and
    value-exactly otherwise; outside-span values unchanged;
  * data_only view of the output is fully numeric in eval columns (no formulas
    left, no None from missing caches).

Usage: python3 scripts/family_probe.py [--workdir DIR]
"""
import argparse
import os
import struct
import subprocess
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font


def build_synthetic(path):
    """Blocks: rows 1-4, blank 5, rows 6-9, blank 10, rows 11-14, blank 15.
    BETA/DELTA/ETA/IOTA/LAMBDA evaluate to clearly negative D and must be
    dropped by the 'D >= 0' predicate."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    data = [
        ("ALPHA", 10.5, 2.25), ("BETA", 8.0, -19.5), ("GAMMA", 11.0, 1.5),
        ("DELTA", 6.0, -12.0),
        ("EPSILON", 12.0, 3.0), ("ZETA", 7.5, 0.5), ("ETA", 5.0, -12.0),
        ("THETA", 9.0, 1.0),
        ("IOTA", 4.0, -11.0), ("KAPPA", 13.0, 2.0), ("LAMBDA", 6.5, -10.5),
        ("MU", 10.0, 1.0),
    ]
    rows = [(1, 4, 1), (6, 9, 2), (11, 14, 3)]
    idx = 0
    for (s, e, g) in rows:
        for i, (name, b, c) in enumerate(data[idx:idx + (e - s + 1)]):
            r = s + i
            ws.cell(row=r, column=1).value = name
            ws.cell(row=r, column=2).value = b
            ws.cell(row=r, column=3).value = c
            ws.cell(row=r, column=4).value = "=B%d+C%d" % (r, r)
            ws.cell(row=r, column=5).value = 100.0 + r
            ws.cell(row=r, column=6).value = "=D%d+E%d" % (r, r)
            ws.cell(row=r, column=7).value = g
            ws.cell(row=r, column=8).value = i + 1
        idx += e - s + 1
    # style graffiti: column widths + a bold label far below the data
    ws.column_dimensions["A"].width = 14.0
    ws.column_dimensions["B"].width = 9.0
    ws.cell(row=40, column=2).value = "TAIL"
    ws.cell(row=40, column=2).font = Font(bold=True)
    wb.save(path)


def expected_span(path):
    """Independent expectation from raw values: blocks -> keep D>=0 -> compacted
    layout with one blank row between groups. No pipeline functions used."""
    wb = load_workbook(path, data_only=False)
    ws = wb.active
    used = sorted({c.column for row in ws.iter_rows(min_row=1, max_row=15)
                   for c in row if c.value is not None})
    def blank(r):
        return all(ws.cell(row=r, column=c).value is None for c in used)
    blocks, start, exp, cursor, group = [], None, {}, 1, 0
    for r in range(1, 16):
        if blank(r):
            if start is not None:
                blocks.append((start, r - 1)); start = None
        elif start is None:
            start = r
    if start is not None:
        blocks.append((start, 15))
    for bi, (s, e) in enumerate(blocks):
        picked = []
        for r in range(s, e + 1):
            d = ws.cell(row=r, column=2).value + ws.cell(row=r, column=3).value
            if d >= 0.0:
                picked.append((r, {c: ws.cell(row=r, column=c).value for c in range(1, 9)}))
        for off, (src, payload) in enumerate(picked):
            tr = cursor + off
            for c, v in payload.items():
                e_val = v
                if c == 4:
                    e_val = ws.cell(row=src, column=2).value + ws.cell(row=src, column=3).value
                elif c == 6:
                    e_val = (ws.cell(row=src, column=2).value + ws.cell(row=src, column=3).value) \
                        + ws.cell(row=src, column=5).value
                exp[(tr, c)] = e_val
        cursor += len(picked)
        if bi < len(blocks) - 1:
            cursor += 1
    return exp, blocks


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--workdir", default=None)
    args = ap.parse_args(argv)
    workdir = args.workdir or tempfile.mkdtemp(prefix="xlsx-family-probe-")
    os.makedirs(workdir, exist_ok=True)
    script_dir = os.path.dirname(os.path.abspath(__file__))
    synth = os.path.join(workdir, "synth_input.xlsx")
    out = os.path.join(workdir, "synth_output.xlsx")
    build_synthetic(synth)

    r = subprocess.run(
        [sys.executable, os.path.join(script_dir, "filter_compact.py"),
         synth, out, "--span", "A1:H15", "--eval-cols", "D:F",
         "--keep-col", "D", "--keep-op", "ge", "--keep-value", "0", "--check"],
        capture_output=True, text=True)
    ok = True

    def check(cond, msg):
        nonlocal ok
        print(("  ok: " if cond else "FAIL: ") + msg)
        ok = ok and cond

    check(r.returncode == 0, "filter_compact --check exits 0")
    if r.returncode != 0:
        print(r.stdout)
        print(r.stderr)

    exp, blocks = expected_span(synth)
    wb = load_workbook(out, data_only=False)
    wbd = load_workbook(out, data_only=True)
    ws = wb.active
    bad = 0
    for (tr, c), e in exp.items():
        o = ws.cell(row=tr, column=c).value
        if isinstance(e, float) and isinstance(o, float):
            if struct.pack("<d", e) != struct.pack("<d", o):
                bad += 1
        elif e != o:
            bad += 1
    check(bad == 0, "span payloads bit-exact vs independent expectation (%d expected cells)" % len(exp))

    # dropped rows: names BETA, DELTA, ETA, IOTA, LAMBDA must be absent in span
    names = [ws.cell(row=r, column=1).value for r in range(1, 16)]
    check(not any(n in names for n in ("BETA", "DELTA", "ETA", "IOTA", "LAMBDA")),
          "predicate-dropped names absent in span")
    kept_names = [n for n in names if n]
    check(sorted(kept_names) == ["ALPHA", "EPSILON", "GAMMA", "KAPPA", "MU", "THETA", "ZETA"],
          "exactly the D>=0 rows remain: %s" % sorted(kept_names))
    # compaction geometry: one blank row between survivor groups; tail blank
    blank_rows = [r for r in range(1, 16) if all(ws.cell(row=r, column=c).value is None for c in range(1, 9))]
    check(blank_rows == [3, 7, 10, 11, 12, 13, 14, 15],
          "compaction blanks: groups then blank tail %s" % blank_rows)
    # tail clearing: old block rows beyond the compacted structure are blank now
    check(all(ws.cell(row=r, column=4).value is None for r in range(10, 16)),
          "vacated rows cleared (D cells None in tail)")
    # no formula text anywhere in span; eval columns numeric incl. data_only
    left = [(r, c) for r in range(1, 16) for c in range(1, 9)
            if isinstance(ws.cell(row=r, column=c).value, str)
            and str(ws.cell(row=r, column=c).value).startswith("=")]
    check(not left, "no formula strings left in span")
    nonnum = [(r, c) for r in range(1, 16) for c in (4, 5, 6)
              if ws.cell(row=r, column=c).value is not None
              and not isinstance(ws.cell(row=r, column=c).value, (int, float))]
    check(not nonnum, "eval columns numeric in formula view")
    none_cached = [(r, c) for r in range(1, 16) for c in (4, 5, 6)
                   if ws.cell(row=r, column=c).value is not None
                   and wbd.active.cell(row=r, column=c).value is None]
    check(not none_cached, "data_only view has no None in eval columns (evaluated from text, no caches needed)")
    # outside-span preservation: TAIL at row 40 col B
    ob = load_workbook(out, data_only=False)
    check(ob.active.cell(row=40, column=2).value == "TAIL", "outside-span value preserved")
    check(ob.active.column_dimensions["A"].width == 14.0, "column width preserved")

    print("\nFAMILY PROBE (%s) %s" % (workdir, "PASSED" if ok else "FAILED"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())