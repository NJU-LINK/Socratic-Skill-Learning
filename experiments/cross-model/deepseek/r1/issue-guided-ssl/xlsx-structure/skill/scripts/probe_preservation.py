#!/usr/bin/env python3
"""
Preservation probe: proves that untouched cells survive the pipeline by
CONSTRUCTION (coordinate-scoped restoration from the input zip), not by
openpyxl round-trip faith.

Builds a probe input from a base workbook by zip-splicing two cells outside the
transform span into the target sheet:
  * H88: an untouched formula cell with a CACHED numeric result
    <c r="H88" s="1"><f>B1*2</f><v>231.2</v></c>
    (no t attribute, as Excel writes numeric formulas; openpyxl's data_only
    then reads the cached float 231.2, the formula view reads '=B1*2');
  * I88: an untouched rich inline-string cell with two runs "ri"+"ch".

Then runs the pipeline (sibling transform.py) on the probe input and asserts
that the output zip retains the cached <v> and the <r> runs, and that both
loads (data_only / formula view) see the expected values. A counterfactual
--no-restore run must FAIL the survival assertions, proving openpyxl alone
strips the cache and collapses the runs.

Usage: python3 probe_preservation.py BASE_INPUT.xlsx [--workdir DIR]
"""
import argparse
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile


def build_probe_input(base_path, probe_path):
    with zipfile.ZipFile(base_path) as z:
        items = {n: z.read(n) for n in z.namelist()}
    part = "xl/worksheets/sheet1.xml"
    xml = items[part].decode("utf-8")
    h88 = '<c r="H88" s="1"><f>B1*2</f><v>231.2</v></c>'
    i88 = '<c r="I88" s="1" t="inlineStr"><is><r><t>ri</t></r><r><t>ch</t></r></is></c>'
    m = re.search(r'<row r="88"[^>]*>.*?</row>', xml, re.S)
    if m:
        inner = m.group(0)[:-len("</row>")]
        xml = xml[:m.start()] + inner + h88 + i88 + "</row>" + xml[m.end():]
    else:
        xml = xml.replace("</sheetData>",
                          '<row r="88">' + h88 + i88 + "</row></sheetData>")
    items[part] = xml.encode("utf-8")
    with zipfile.ZipFile(probe_path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items():
            z.writestr(name, data)


def run_pipeline(transform_script, probe_in, probe_out, extra=()):
    cmd = [sys.executable, transform_script, probe_in, probe_out, "--check", *extra]
    r = subprocess.run(cmd, capture_output=True, text=True)
    return r


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("base_input")
    ap.add_argument("--workdir", default=None)
    args = ap.parse_args(argv)
    workdir = args.workdir or tempfile.mkdtemp(prefix="xlsx-probe-")
    os.makedirs(workdir, exist_ok=True)
    transform_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "transform.py")

    probe_in = os.path.join(workdir, "probe_input.xlsx")
    probe_out = os.path.join(workdir, "probe_output.xlsx")
    build_probe_input(args.base_input, probe_in)

    ok = True

    def check(cond, msg):
        nonlocal ok
        print(("  ok: " if cond else "FAIL: ") + msg)
        ok = ok and cond

    # --- pipeline WITH restoration (the production path) ---
    r = run_pipeline(transform_script, probe_in, probe_out)
    check(r.returncode == 0, "pipeline --check on probe input exits 0")
    if r.returncode != 0:
        print(r.stdout)
        print(r.stderr)
    with zipfile.ZipFile(probe_out) as z:
        xml = z.read("xl/worksheets/sheet1.xml")
    check(b"<v>231.2</v>" in xml, "output zip retains cached <v>231.2</v> for H88")
    check(b"<r><t>ri</t></r><r><t>ch</t></r>" in xml,
          "output zip retains rich-text runs for I88")

    import openpyxl
    wbf = openpyxl.load_workbook(probe_out, data_only=False)
    wbd = openpyxl.load_workbook(probe_out, data_only=True)
    check(wbf["Sheet1"]["H88"].value == "=B1*2",
          "formula view H88 reads '=B1*2'")
    check(wbd["Sheet1"]["H88"].value == 231.2,
          "data_only H88 reads 231.2 (cached result survived)")
    check(str(wbf["Sheet1"]["I88"].value) == "rich",
          "formula view I88 reads 'rich'")
    check(str(wbd["Sheet1"]["I88"].value) == "rich",
          "data_only I88 reads 'rich'")

    # --- counterfactual WITHOUT restoration: openpyxl alone strips the cache ---
    probe_out2 = os.path.join(workdir, "probe_output_norestore.xlsx")
    r2 = run_pipeline(transform_script, probe_in, probe_out2, ("--no-restore",))
    wbd2 = openpyxl.load_workbook(probe_out2, data_only=True)
    stripped = wbd2["Sheet1"]["H88"].value is None
    check(stripped,
          "counterfactual: without --restore openpyxl strips the cached result "
          "(data_only H88 is None) - proves restoration is the mechanism")
    with zipfile.ZipFile(probe_out2) as z:
        xml2 = z.read("xl/worksheets/sheet1.xml")
    check(b"<r><t>ri</t></r><r><t>ch</t></r>" not in xml2,
          "counterfactual: without restore the rich runs are collapsed")

    print("\nPROBE %s (%s)" % ("PASSED" if ok else "FAILED", workdir))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())