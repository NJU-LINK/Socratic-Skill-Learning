#!/usr/bin/env python3
"""
xlsx-structure verifier: full preservation + transformation battery.

Standalone: imports the evaluator AND the production evaluation pipeline from
solve.py (same directory) so the battery certifies exactly the code path the
solver ran.  Run after solve.py has produced OUTPUT.  Exits non-zero on any
failure.

Checks:
  A. Evaluator unit probes (Excel semantics): -2^2==4, 2^3^2==64 (left fold),
     50%^2==0.25, 2^-2==0.25.  Any failure = grammar regression.
  B. Hard bit-equality THROUGH THE PRODUCTION PATH, split by formula class:
     bare arithmetic (+, -, *, /, ^, % over stored doubles) must equal the
     cached value with == (identical IEEE-754 op sequence is deterministic);
     function-bearing formulas (SUM/AVERAGE/... over ranges) have unspecified
     accumulation order and possibly quantized engine caches, so they are
     checked with 1e-9-relative tolerance (B2), never bitwise.
  C. Answer-range hygiene: zero formula artifacts inside EVAL_COLS — no
     data_type=='f' cells, no '='-prefixed or formula-like strings.  Plain
     preserved data text ("N/A") is NOT a violation.
  D. Block structure identical: same separator rows, same block sizes.
  E. Exact stable permutation identity per block: output rows must equal the
     input rows permuted by the stable key (sign * evaluated KEY_COL, original
     index) — order-blind multiset equality is NOT enough for tie stability.
     When the input's column 8 (H) holds a permutation of 1..n within a block,
     the output H must equal the expected stable permutation too (v01
     fingerprint).
  F. Everything outside the answer range identical: values, styled-cell
     inventory as (coordinate, style_id), merged ranges (also INSIDE the
     range, post unmerge/remerge), auto_filter.ref, freeze_panes, row
     dims TRAVELING with content (rekeyed through the permutation; unmoved
     rows positional; v01 vacuous — zero input row dims), column widths
     positional, sheet names, workbook defined names.
  G. Raw XML observations (informational): <dimension> tag survival,
     calcChain.xml dropped, string serialization mode.
"""
import math
import os
import re
import sys
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import solve

FAILS = []


def check(name, cond, detail=""):
    status = "PASS" if cond else "FAIL"
    print("%-58s %s %s" % (name, status, detail))
    if not cond:
        FAILS.append(name)


# --------------------------------------------------------------------------
# A. evaluator unit probes (Excel semantics)
# --------------------------------------------------------------------------
def _eval_simple(text):
    return solve.evaluate(text, lambda ref: None)


check("A: probe -2^2 == 4 (negation > ^)", _eval_simple("=-2^2") == 4.0, repr(_eval_simple("=-2^2")))
check("A: probe 2^3^2 == 64 (left-assoc ^)", _eval_simple("=2^3^2") == 64.0, repr(_eval_simple("=2^3^2")))
check("A: probe 50%^2 == 0.25 (% tightest)", math.isclose(_eval_simple("=50%^2"), 0.25), repr(_eval_simple("=50%^2")))
check("A: probe 2^-2 == 0.25 (signed exponent)", math.isclose(_eval_simple("=2^-2"), 0.25), repr(_eval_simple("=2^-2")))
check("A: probe 1+2*3^2 == 19", _eval_simple("=1+2*3^2") == 19.0, repr(_eval_simple("=1+2*3^2")))


class _StubResolver:
    """Plain-cell + single-row-range resolver for parser probes."""

    def __init__(self, table):
        self._t = table

    def cell(self, ref):
        return self._t.get(ref)

    def range(self, r1, r2):
        m1 = re.fullmatch(r"([A-Za-z]{1,3})(\d+)", r1)
        m2 = re.fullmatch(r"([A-Za-z]{1,3})(\d+)", r2)
        if not m1 or not m2 or m1.group(2) != m2.group(2):
            return None
        c1 = openpyxl.utils.column_index_from_string(m1.group(1))
        c2 = openpyxl.utils.column_index_from_string(m2.group(1))
        if c1 > c2:
            return None
        return [self._t.get(openpyxl.utils.get_column_letter(c) + m1.group(2))
                for c in range(c1, c2 + 1)]


stub = _StubResolver({"B1": 1.0, "C1": 2.0, "D2": 4.0, "E2": 6.0, "F2": 8.0})
check("A: probe =SUM(B1:C1) == 3 (range)", solve.evaluate("=SUM(B1:C1)", stub) == 3.0,
      repr(solve.evaluate("=SUM(B1:C1)", stub)))
check("A: probe =AVERAGE(D2:F2) == 6", solve.evaluate("=AVERAGE(D2:F2)", stub) == 6.0,
      repr(solve.evaluate("=AVERAGE(D2:F2)", stub)))
check("A: probe =ROUND(2.5,0) == 3 (half away)", solve.evaluate("=ROUND(2.5,0)", stub) == 3.0,
      repr(solve.evaluate("=ROUND(2.5,0)", stub)))
check("A: probe =INT(-1.5) == -2 (floor)", solve.evaluate("=INT(-1.5)", stub) == -2.0,
      repr(solve.evaluate("=INT(-1.5)", stub)))
check("A: probe =MOD(5,2) == 1", solve.evaluate("=MOD(5,2)", stub) == 1.0,
      repr(solve.evaluate("=MOD(5,2)", stub)))
check("A: probe =ABS(B1-C1)+1 == 2", solve.evaluate("=ABS(B1-C1)+1", stub) == 2.0,
      repr(solve.evaluate("=ABS(B1-C1)+1", stub)))
check("A: probe =SUM(B1:C1,2) == 5 (range+scalar)", solve.evaluate("=SUM(B1:C1,2)", stub) == 5.0,
      repr(solve.evaluate("=SUM(B1:C1,2)", stub)))
check("A: probe =COUNT(D2:F2) == 3", solve.evaluate("=COUNT(D2:F2)", stub) == 3.0,
      repr(solve.evaluate("=COUNT(D2:F2)", stub)))
check("A: probe bare range =B1:C1 -> None", solve.evaluate("=B1:C1", stub) is None)
check("A: probe cross-row range =SUM(B1:B2) -> None (boundary)",
      solve.evaluate("=SUM(B1:B2)", stub) is None)
check("A: probe =MOD(5,0) -> None (#DIV/0!)", solve.evaluate("=MOD(5,0)", lambda ref: None) is None)
check("A: probe =1/0 -> None (#DIV/0!)", solve.evaluate("=1/0", lambda ref: None) is None)
check("A: probe =(-2)^0.5 -> None (#NUM!)", solve.evaluate("=(-2)^0.5", lambda ref: None) is None)
check("A: probe =(-2)^2 == 4 (neg base, integral exp)",
      solve.evaluate("=(-2)^2", lambda ref: None) == 4.0)
check("A: probe data 'N/A' not formula", solve.is_formula_text("N/A") is False)
check("A: probe data '12:30' not formula", solve.is_formula_text("12:30") is False)
check("A: probe data 'E2/TAX' not formula", solve.is_formula_text("E2/TAX") is False)

# --------------------------------------------------------------------------
# load input (both modes) and output
# --------------------------------------------------------------------------
wb_f = openpyxl.load_workbook(solve.INPUT, data_only=False)
wb_c = openpyxl.load_workbook(solve.INPUT, data_only=True)
wb_o = openpyxl.load_workbook(solve.OUTPUT, data_only=False)
ws_f, ws_c, ws_o = wb_f[solve.SHEET], wb_c[solve.SHEET], wb_o[solve.SHEET]


def is_blank(ws_, r):
    return all(ws_.cell(row=r, column=c).value is None for c in range(solve.C1, solve.C2 + 1))


def blocks_of(ws_):
    blocks, cur = [], []
    for r in range(solve.R1, solve.R2 + 1):
        if is_blank(ws_, r):
            if cur:
                blocks.append(cur)
                cur = []
        else:
            cur.append(r)
    if cur:
        blocks.append(cur)
    return blocks


blocks_i, blocks_o = blocks_of(ws_f), blocks_of(ws_o)

# input snapshot + PRODUCTION evaluation (the exact function the solver ran)
snap_i = {}
for b in blocks_i:
    snap_i.update(solve.snapshot_block(ws_f, b, solve.C1, solve.C2))
all_rows = [r for b in blocks_i for r in b]
eval_i = solve.evaluate_block(snap_i, all_rows, ws_c, solve.EVAL_COLS, solve.C1, solve.C2,
                              abort_on_failure=False)

# --------------------------------------------------------------------------
# B. hard bit-equality through the production path
# --------------------------------------------------------------------------
mism_bare = mism_func = bare_n = func_n = 0
for r in all_rows:
    for c in solve.EVAL_COLS:
        if not (solve.C1 <= c <= solve.C2):
            continue
        cell = ws_f.cell(row=r, column=c)
        if not solve.is_formula_text(cell.value):
            continue
        cache = ws_c.cell(row=r, column=c).value
        if not solve.is_number(cache):
            continue  # formula cell without cache: nothing to compare
        got = eval_i.get((r, c))
        toks = solve._tokenize(cell.value.lstrip("=").strip())
        has_func = any(k == "ident" for k, _ in toks)
        if has_func:
            # aggregates/ranges: engine accumulation order is unspecified and
            # engine caches may be quantized (15 sig digits) -> tolerance
            func_n += 1
            if got is None or abs(got - cache) > 1e-9 * max(1.0, abs(cache)):
                mism_func += 1
                if mism_func <= 5:
                    print("  func-cache mismatch", cell.coordinate, "production:", got,
                          "cache:", cache)
        else:
            # bare arithmetic of stored doubles: identical op sequence is
            # deterministic -> bit-exact == is the guarantee
            bare_n += 1
            if got is None or got != cache:
                mism_bare += 1
                if mism_bare <= 5:
                    print("  bit-mismatch", cell.coordinate, "production:", got, "cache:", cache)
check("B: bare-arithmetic evaluate == cache (bit-exact)", mism_bare == 0,
      "%d bare cells, %d mismatches" % (bare_n, mism_bare))
check("B2: function-bearing evaluate ~ cache (<=1e-9 rel)", mism_func == 0,
      "%d function cells, %d outside tolerance" % (func_n, mism_func))

# --------------------------------------------------------------------------
# C/D: hygiene + block structure
# --------------------------------------------------------------------------
sep_i = [r for r in range(solve.R1, solve.R2 + 1) if is_blank(ws_f, r)]
sep_o = [r for r in range(solve.R1, solve.R2 + 1) if is_blank(ws_o, r)]
check("D: separator rows identical", sep_i == sep_o, str(sep_i))
check("D: block sizes identical", [len(b) for b in blocks_i] == [len(b) for b in blocks_o])

hygiene_bad = 0
# hygiene: NO formula artifacts in declared columns.  Formula artifacts are
# data_type=='f' cells, or strings that are '='-prefixed or formula-like.
# Preserved non-formula values (numbers, plain text like "N/A", None) are
# contract-satisfying: the task forbids formula strings / unevaluated
# formulas, not data text.
for r in [rr for b in blocks_o for rr in b]:
    for c in solve.EVAL_COLS:
        if not (solve.C1 <= c <= solve.C2):
            continue
        cell = ws_o.cell(row=r, column=c)
        if cell.data_type == "f":
            hygiene_bad += 1
        elif isinstance(cell.value, str) and (cell.value.strip().startswith("=")
                                              or solve.is_formula_text(cell.value)):
            hygiene_bad += 1
check("C: answer-range hygiene in EVAL_COLS (no formula artifacts)", hygiene_bad == 0, "%d bad" % hygiene_bad)

# --------------------------------------------------------------------------
# E. exact stable permutation identity per block
# --------------------------------------------------------------------------
def norm_input_row(r):
    """Input row with EVAL_COLS normalized to the PRODUCTION evaluated value.
    Non-EVAL columns use the raw input value (the resolver may memoize
    same-row referenced cells as floats; those must NOT leak into the
    expected row — write substitution is scoped to EVAL_COLS)."""
    out = []
    for c in range(solve.C1, solve.C2 + 1):
        if c in solve.EVAL_COLS and eval_i.get((r, c)) is not None:
            v = eval_i[(r, c)]
        else:
            v = ws_f.cell(row=r, column=c).value
        out.append(round(v, 10) if isinstance(v, float) else v)
    return tuple(out)


def out_row(r):
    out = []
    for c in range(solve.C1, solve.C2 + 1):
        v = ws_o.cell(row=r, column=c).value
        out.append(round(v, 10) if isinstance(v, float) else v)
    return tuple(out)


def key_of_input(r):
    ev = eval_i.get((r, solve.KEY_COL))
    if ev is not None:
        return ev
    v = ws_f.cell(row=r, column=solve.KEY_COL).value
    return float(v) if solve.is_number(v) else 0.0


sign = -1.0 if solve.DESCENDING else 1.0
perm_ok, multiset_ok, h_ok = True, True, True
for bi, bo in zip(blocks_i, blocks_o):
    n = len(bi)
    in_rows = [norm_input_row(r) for r in bi]
    out_rows = [out_row(r) for r in bo]
    perm = sorted(range(n), key=lambda i: (sign * key_of_input(bi[i]), i))
    if out_rows != [in_rows[perm[i]] for i in range(n)]:
        perm_ok = False
        for k in range(n):
            if out_rows[k] != in_rows[perm[k]]:
                print("  perm diff at out pos %d: got %s expected %s" % (k, out_rows[k], in_rows[perm[k]]))
    if sorted(in_rows, key=repr) != sorted(out_rows, key=repr):
        multiset_ok = False
    # H fingerprint: if input col 8 is a permutation of 1..n within the block
    hin = [ws_f.cell(row=r, column=8).value for r in bi]
    if sorted(hin) == list(range(1, n + 1)):
        hout = [ws_o.cell(row=r, column=8).value for r in bo]
        if hout != [hin[perm[i]] for i in range(n)]:
            h_ok = False
            print("  H mismatch block %d-%d: got %s expected %s"
                  % (bi[0], bi[-1], hout, [hin[perm[i]] for i in range(n)]))
check("E: exact stable permutation per block", perm_ok)
check("E: normalized multiset per block (safety net)", multiset_ok)
check("E: output H fingerprint == stable permutation", h_ok)

# --------------------------------------------------------------------------
# H. serialization-leg drift: reloaded output values vs in-memory evaluated
#    (the EVALUATE leg, check B, is bit-exact in memory; openpyxl writes
#    numerics with %.16g, so a 17-significant-digit double may lose its last
#    ulp IN THE FILE.  v01 real drift: F7/D2 at ~4e-17.  Anything above 1e-9
#    means the pipeline broke, not that %.16g misbehaved.)
# --------------------------------------------------------------------------
maxdrift = 0.0
drift_cells = []
for bi, bo in zip(blocks_i, blocks_o):
    n = len(bi)
    perm = sorted(range(n), key=lambda i: (sign * key_of_input(bi[i]), i))
    for k, dest in enumerate(bo):
        src = bi[perm[k]]
        for c in solve.EVAL_COLS:
            if not (solve.C1 <= c <= solve.C2):
                continue
            ev = eval_i.get((src, c))
            if ev is None:
                continue
            got = ws_o.cell(row=dest, column=c).value
            if not solve.is_number(got):
                continue
            d = abs(float(got) - float(ev))
            maxdrift = max(maxdrift, d)
            if d != 0.0:
                drift_cells.append((src, c, ev, got))
check("H: serialization drift <= 1e-9 over EVAL_COLS", maxdrift <= 1e-9,
      "maxdrift=%.3g, %d drift cells" % (maxdrift, len(drift_cells)))
print("INFO  serialize-leg drift cells (reloaded != evaluated in-memory): %d" % len(drift_cells))
for (src, c, ev, got) in drift_cells[:10]:
    print("       %s: evaluated %r reloaded %r"
          % (openpyxl.utils.get_column_letter(c) + str(src), ev, got))

# --------------------------------------------------------------------------
# F. outside-range and feature preservation
# --------------------------------------------------------------------------
check("F: sheet names identical", wb_o.sheetnames == wb_f.sheetnames)
vi = {(r, c): cell.value for (r, c), cell in ws_f._cells.items()
      if cell.value is not None and not (solve.R1 <= r <= solve.R2 and solve.C1 <= c <= solve.C2)}
vo = {(r, c): cell.value for (r, c), cell in ws_o._cells.items()
      if cell.value is not None and not (solve.R1 <= r <= solve.R2 and solve.C1 <= c <= solve.C2)}
check("F: values outside answer range identical", vi == vo, "%d vs %d cells" % (len(vi), len(vo)))


def styled_inv(ws_):
    return {(r, c): cell.style_id for (r, c), cell in ws_._cells.items()
            if cell.has_style and cell.value is None
            and not (solve.R1 <= r <= solve.R2 and solve.C1 <= c <= solve.C2)}


si, so = styled_inv(ws_f), styled_inv(ws_o)
check("F: styled-empty cells outside range identical", si == so, "%d cells" % len(si))
check("F: merged ranges identical", sorted(str(m) for m in ws_f.merged_cells.ranges)
      == sorted(str(m) for m in ws_o.merged_cells.ranges))
check("F: auto_filter.ref identical", ws_f.auto_filter.ref == ws_o.auto_filter.ref)
check("F: freeze_panes identical", ws_f.freeze_panes == ws_o.freeze_panes)

# Row dims TRAVEL with row content (declared Excel-row-sort semantics).  A
# positional rh_i == rh_o check is WRONG under this semantics (correct
# output has heights permuted); instead rekey input dims through the same
# stable permutation used by E/H (dest_of[src]), and require positional
# identity for rows that never move (separators, out-of-block rows).  v01
# evidence for this check is VACUOUS (input has zero row_dimensions
# entries); the property is toy-verified (T5/T8).
def dim_inv(ws_):
    out = {}
    for r, d in ws_.row_dimensions.items():
        t = (d.height, d.hidden, d.outlineLevel, d.s)
        if any(v not in (None, False, 0) for v in t):
            out[r] = t
    return out


dest_of = {}
# Conservation note: dest_of is derived from the BLOCK permutation because
# v01's transform moves rows only within blocks.  For a variant whose
# destination rows lie outside the blocks (append, extraction-into-a-tail),
# the dim rekey must come from that variant's plan — preservation checks
# over row-unit attributes compare through the transform's map, never by
# coordinates.
for bi, bo in zip(blocks_i, blocks_o):
    n = len(bi)
    perm = sorted(range(n), key=lambda i: (sign * key_of_input(bi[i]), i))
    for k, dest in enumerate(bo):
        dest_of[bi[perm[k]]] = dest
dim_in = dim_inv(ws_f)
dim_expect = {dest_of[src]: t for src, t in dim_in.items() if src in dest_of}
dim_expect.update({r: t for r, t in dim_in.items() if r not in dest_of})
check("F: row dims travel with content (permuted) / stay (unmoved)",
      dim_expect == dim_inv(ws_o),
      "%d dim entries in, %d in output" % (len(dim_in), len(dim_inv(ws_o))))

cw_i = {c: (d.width, d.customWidth) for c, d in ws_f.column_dimensions.items() if d.width is not None}
cw_o = {c: (d.width, d.customWidth) for c, d in ws_o.column_dimensions.items() if d.width is not None}
check("F: column widths identical (positional by design)", cw_i == cw_o)
try:
    dn_i = sorted((n, str(dn.value)) for n, dn in wb_f.defined_names.items())
    dn_o = sorted((n, str(dn.value)) for n, dn in wb_o.defined_names.items())
except AttributeError:
    dn_i = dn_o = []
check("F: defined names identical", dn_i == dn_o)

# --------------------------------------------------------------------------
# G. raw XML observations (informational)
# --------------------------------------------------------------------------
zi = zipfile.ZipFile(solve.INPUT)
zo = zipfile.ZipFile(solve.OUTPUT)
xio = zi.read("xl/worksheets/sheet1.xml").decode()
xoo = zo.read("xl/worksheets/sheet1.xml").decode()
m_i = re.search(r'<dimension[^>]*/>', xio)
m_o = re.search(r'<dimension[^>]*/>', xoo)
print("INFO  dimension: input %r -> output %r" % (m_i.group(0) if m_i else None, m_o.group(0) if m_o else None))
print("INFO  calcChain in output:", "xl/calcChain.xml" in zo.namelist(), "(absent is expected)")
print("INFO  output string cells inline:", xoo.count('t="inlineStr"'), " shared-ref:", xoo.count('t="s"'))
print("INFO  openpyxl writes numerics with %.16g and reads back decimal-free",
      "strings as int: reloaded-vs-evaluated comparisons are numeric (see H),",
      "never bit-exact/type-strict")

print()
if FAILS:
    print("VERIFICATION FAILED:", FAILS)
    sys.exit(1)
print("VERIFICATION PASSED")