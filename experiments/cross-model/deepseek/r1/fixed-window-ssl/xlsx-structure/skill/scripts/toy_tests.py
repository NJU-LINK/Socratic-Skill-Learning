#!/usr/bin/env python3
"""
toy_tests.py: parameterization edge-case unit tests for solve.py.

Runs synthetic workbooks through the production solve() pipeline and asserts
the intended behavior.  Catches regressions in:
  T1  KEY_COL is a plain numeric column OUTSIDE EVAL_COLS  -> still sorts
      correctly (value-first sort_key fallback).
  T2  KEY_COL holds non-numeric text                        -> solve aborts
      loudly with EvalError instead of silently sorting by a constant.
  T3  EVAL_COLS contains a formula cell with NO cached value -> evaluated
      to a real number and written as float, never leaked as a formula
      string (the old silent cache-fallback bug).
  T4  Function formulas (SUM/AVERAGE) as real formula cells without cache
      and as text formulas -> evaluated row-locally; a cross-row range
      formula without cache aborts loudly (unsupported boundary).
  T5  Pluggable transform 'filter_by_key' (synthetic): rows > threshold
      packed to block tops, vacated rows removed from the grid.
  T6  Pluggable transform 'append_by_copy' (synthetic): source rows copied
      into the first empty rows after a table; sources untouched.
  T7  Data text in a declared column ("N/A", "12:30") is NOT a formula:
      no crash, preserved verbatim (classification-by-parse, not by tokens).

Expectations below are computed from each toy's OWN inputs; when a toy
"fails", recompute the expectation from the toy inputs before touching
solve.py — the expectation string is the least-tested code in the loop.
"""
import os
import shutil
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import solve

FAILS = []


def make_wb(path, rows, ncols=3, extra=None, decorate=None):
    """rows: list of (col, value); extra: (row, col, value) one-off writes;
    decorate(ws): optional callback to set row dims/styles before save."""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    ws = wb.active
    for i, rvals in enumerate(rows, start=1):
        for c, v in rvals:
            if c <= ncols:
                ws.cell(row=i, column=c).value = v
    for (r, c, v) in (extra or []):
        ws.cell(row=r, column=c).value = v
    if decorate:
        decorate(ws)
    wb.save(path)
    return path


def run_solve(inp, out, nrows, ncols=3, key_col=2, eval_cols=(3,), descending=True,
              transform=None, threshold=None, src_rows=None, dest_start=None):
    s = solve
    s.INPUT, s.OUTPUT = inp, out
    s.SHEET = "Sheet1"
    s.R1, s.R2, s.C1, s.C2 = 1, nrows, 1, ncols
    s.KEY_COL, s.EVAL_COLS, s.DESCENDING = key_col, eval_cols, descending
    s.TRANSFORM = transform if transform is not None else "block_stable_sort"
    s.FILTER_THRESHOLD = threshold if threshold is not None else 100.0
    s.APPEND_SOURCE_ROWS = src_rows if src_rows is not None else ()
    s.APPEND_DEST_START = dest_start if dest_start is not None else None
    s.main()


def check(name, cond, detail=""):
    print("%-62s %s %s" % (name, "PASS" if cond else "FAIL", detail))
    if not cond:
        FAILS.append(name)


tmp = tempfile.mkdtemp(prefix="xlsx_toy_")

# --- T1: numeric KEY_COL outside EVAL_COLS sorts descending ----------------
# row1 B=5, row2 B=3, row3 B=4 -> desc [row1, row3, row2]; C only row1 has
# formula =B1+1 -> evaluated 6; rows 2-3 C are None.
t1_in = make_wb(os.path.join(tmp, "t1.xlsx"),
                [[(1, "row1"), (2, 5)], [(1, "row2"), (2, 3)], [(1, "row3"), (2, 4)]],
                extra=[(1, 3, "=B1+1")])
run_solve(t1_in, os.path.join(tmp, "t1o.xlsx"), nrows=3, key_col=2, eval_cols=(3,))
w1 = openpyxl.load_workbook(os.path.join(tmp, "t1o.xlsx"), data_only=False)["Sheet1"]
got_a = [w1.cell(row=r, column=1).value for r in range(1, 4)]
check("T1: numeric key outside EVAL_COLS sorts desc", got_a == ["row1", "row3", "row2"], repr(got_a))
check("T1: only formula-carrying row evaluated to number",
      [w1.cell(row=r, column=3).value for r in range(1, 4)] == [6.0, None, None])

# --- T2: text key column must abort loudly ---------------------------------
t2_in = make_wb(os.path.join(tmp, "t2.xlsx"),
                [[(1, "x1"), (2, "txt1")], [(1, "x2"), (2, "txt2")], [(1, "x3"), (2, "txt3")]])
try:
    run_solve(t2_in, os.path.join(tmp, "t2o.xlsx"), nrows=3, key_col=2, eval_cols=(3,))
    check("T2: text key aborts loudly", False, "did not raise")
except solve.EvalError:
    check("T2: text key aborts loudly", True, "EvalError raised")

# --- T3: formula cell with no cached value -> evaluated float --------------
t3_in = make_wb(os.path.join(tmp, "t3.xlsx"),
                [[(1, "n1"), (2, 10)], [(1, "n2"), (2, 20)], [(1, "n3"), (2, 30)]],
                extra=[(r, 3, "=B%d+1" % r) for r in range(1, 4)])
run_solve(t3_in, os.path.join(tmp, "t3o.xlsx"), nrows=3, key_col=2, eval_cols=(3,))
w3 = openpyxl.load_workbook(os.path.join(tmp, "t3o.xlsx"), data_only=False)["Sheet1"]
c3 = [w3.cell(row=r, column=3).value for r in range(1, 4)]
check("T3: no-cache formulas evaluate to floats", c3 == [31.0, 21.0, 11.0], repr(c3))
check("T3: no formula strings leaked",
      all(not (isinstance(v, str) and v.strip().startswith("=")) for v in c3))

# --- T4: function formulas (row-local ranges), with and without cache ------
t4_in = make_wb(os.path.join(tmp, "t4.xlsx"),
                [[(1, "f1"), (2, 1), (3, 5), (4, "=SUM(B1:C1)")],   # real formula, no cache
                 [(1, "f2"), (2, 2), (3, 6), (4, "=AVERAGE(B2:C2)")],  # text formula
                 [(1, "f3"), (2, 3), (3, 7), (4, "=B3+C3")]],          # text arithmetic
                ncols=4)
run_solve(t4_in, os.path.join(tmp, "t4o.xlsx"), nrows=3, ncols=4, key_col=2,
          eval_cols=(4,))
w4 = openpyxl.load_workbook(os.path.join(tmp, "t4o.xlsx"), data_only=False)["Sheet1"]
d4 = [w4.cell(row=r, column=4).value for r in range(1, 4)]
# Evaluations: row1 SUM(B1:C1)=1+5=6, row2 AVERAGE(B2:C2)=4, row3 B3+C3=10.
# B = [1,2,3] sorts DESCENDING -> dest row1<-src3(10), row2<-src2(4), row3<-src1(6).
check("T4: =SUM(B1:C1) no-cache -> 6 (at dest row 3)", d4[2] == 6.0, repr(d4[2]))
check("T4: =AVERAGE(B2:C2) text -> 4 (at dest row 2)", d4[1] == 4.0, repr(d4[1]))
check("T4: =B3+C3 text -> 10 (at dest row 1)", d4[0] == 10.0, repr(d4[0]))
check("T4: no formula cells left", not any(w4.cell(row=r, column=4).data_type == "f"
                                           for r in range(1, 4)))

# cross-row range text formula, no cache -> loud abort (unsupported boundary)
t4b_in = make_wb(os.path.join(tmp, "t4b.xlsx"),
                 [[(1, "g1"), (2, 1), (3, 5), (4, "=SUM(B1:B2)")]], ncols=4)
try:
    run_solve(t4b_in, os.path.join(tmp, "t4bo.xlsx"), nrows=2, ncols=4, key_col=2,
              eval_cols=(4,))
    check("T4: cross-row range without cache aborts loudly", False, "did not raise")
except solve.EvalError:
    check("T4: cross-row range without cache aborts loudly", True, "EvalError raised")

# --- T5: synthetic filter_by_key (pluggable transform, vacated rows) -------
# B = [10, 200, 30, 400], threshold 100 -> keep rows 2 and 4 packed to top.
# Row dims exercise: heights (10,20,30,40), row2 hidden, row4 banded (s=1)
# -> dest1<-src2 carries h20+hidden; dest2<-src4 carries h40+s=1; cleared
# rows 3-4 lose their dims with their cells (no ghost height/banding).
def decorate_t5(ws):
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 40
    ws.cell(row=4, column=1).fill = openpyxl.styles.PatternFill(
        start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    ws.row_dimensions[4]._style = ws.cell(row=4, column=1)._style


t5_in = make_wb(os.path.join(tmp, "t5.xlsx"),
                [[(1, "a"), (2, 10)], [(1, "b"), (2, 200)],
                 [(1, "c"), (2, 30)], [(1, "d"), (2, 400)]], ncols=2,
                decorate=decorate_t5)
run_solve(t5_in, os.path.join(tmp, "t5o.xlsx"), nrows=4, ncols=2, key_col=2,
          eval_cols=(), transform="filter_by_key", threshold=100.0)
w5 = openpyxl.load_workbook(os.path.join(tmp, "t5o.xlsx"), data_only=False)["Sheet1"]
# Check grid absence BEFORE materializing cells with ws.cell() probes.
vacated_absent = all((r, c) not in w5._cells for r in (3, 4) for c in (1, 2))
a5 = [w5.cell(row=r, column=1).value for r in range(1, 5)]
check("T5: filter packs survivors to block top", a5 == ["b", "d", None, None], repr(a5))
check("T5: vacated rows removed from grid", vacated_absent)
d1, d2, d3, d4 = (w5.row_dimensions.get(r) for r in (1, 2, 3, 4))
check("T5: row dims travel with content (dest1<-src2 h20+hidden)",
      (d1.height, d1.hidden) == (20.0, True), repr((d1.height, d1.hidden)))
check("T5: row banding style travels (dest2<-src4 s=1)",
      d2 is not None and d2.height == 40.0 and d2.s == 1,
      repr(None if d2 is None else (d2.height, d2.s)))
check("T5: cleared rows lose ghost dims", d3 is None and d4 is None)

# --- T6: synthetic append_by_copy (pluggable transform, out-of-block dest) -
# table rows 1-2, empty rows 3-4, source rows 5-6 -> copies into 3-4.
t6_in = make_wb(os.path.join(tmp, "t6.xlsx"),
                [[(1, "t1"), (2, 11)], [(1, "t2"), (2, 12)],
                 [(1, None), (2, None)], [(1, None), (2, None)],
                 [(1, "s1"), (2, 21)], [(1, "s2"), (2, 22)]], ncols=2)
run_solve(t6_in, os.path.join(tmp, "t6o.xlsx"), nrows=6, ncols=2, key_col=2,
          eval_cols=(), transform="append_by_copy", src_rows=(5, 6), dest_start=3)
w6 = openpyxl.load_workbook(os.path.join(tmp, "t6o.xlsx"), data_only=False)["Sheet1"]
a6 = [w6.cell(row=r, column=1).value for r in range(1, 7)]
b6 = [w6.cell(row=r, column=2).value for r in range(1, 7)]
check("T6: append copies sources into empty tail", a6 == ["t1", "t2", "s1", "s2", "s1", "s2"],
      repr(a6))
check("T6: appended values intact", b6 == [11, 12, 21, 22, 21, 22], repr(b6))

# --- T7: data text in a declared column -> preserved, no crash -------------
t7_in = make_wb(os.path.join(tmp, "t7.xlsx"),
                [[(1, "x1"), (2, 5), (3, "N/A")],
                 [(1, "x2"), (2, 3), (3, "12:30")],
                 [(1, "x3"), (2, 4), (3, "OK")]], ncols=3)
run_solve(t7_in, os.path.join(tmp, "t7o.xlsx"), nrows=3, ncols=3, key_col=2,
          eval_cols=(3,))
w7 = openpyxl.load_workbook(os.path.join(tmp, "t7o.xlsx"), data_only=False)["Sheet1"]
c7 = [w7.cell(row=r, column=3).value for r in range(1, 4)]
check("T7: data text in EVAL_COL survives sort (no crash, verbatim)",
      c7 == ["N/A", "OK", "12:30"], repr(c7))

# --- unit: error-domain probes (no escaping exceptions) ---------------------
ev_none = solve.evaluate
check("U: =MOD(5,0) -> None not exception", ev_none("=MOD(5,0)", lambda ref: None) is None)
check("U: =1/0 -> None not exception", ev_none("=1/0", lambda ref: None) is None)
check("U: =(-2)^0.5 -> None not exception", ev_none("=(-2)^0.5", lambda ref: None) is None)
check("U: =(-2)^2 == 4", ev_none("=(-2)^2", lambda ref: None) == 4.0)

# --- T8: sort moves row-level state with content ---------------------------
# B = [3, 1, 2] desc -> dest1<-src1, dest2<-src3, dest3<-src2.  Heights
# (10/20/30) must land on dest rows 1/3/2 respectively; row2 hidden and
# row3 banded (s=1) travel to dest3 and dest2.
def decorate_t8(ws):
    ws.row_dimensions[1].height = 10
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[2].hidden = True
    ws.row_dimensions[3].height = 30
    ws.cell(row=3, column=1).fill = openpyxl.styles.PatternFill(
        start_color="FF0000", end_color="FF0000", fill_type="solid")
    ws.row_dimensions[3]._style = ws.cell(row=3, column=1)._style


t8_in = make_wb(os.path.join(tmp, "t8.xlsx"),
                [[(1, "x1"), (2, 3)], [(1, "x2"), (2, 1)], [(1, "x3"), (2, 2)]],
                ncols=2, decorate=decorate_t8)
run_solve(t8_in, os.path.join(tmp, "t8o.xlsx"), nrows=3, ncols=2, key_col=2,
          eval_cols=())
w8 = openpyxl.load_workbook(os.path.join(tmp, "t8o.xlsx"), data_only=False)["Sheet1"]
a8 = [w8.cell(row=r, column=1).value for r in range(1, 4)]
check("T8: sort order x1,x3,x2", a8 == ["x1", "x3", "x2"], repr(a8))
g1, g2, g3 = (w8.row_dimensions.get(r) for r in (1, 2, 3))
check("T8: height travels (dest rows 1,3,2 <- srcs 1,3,2)",
      (g1.height, g2.height, g3.height) == (10.0, 30.0, 20.0),
      repr((g1.height, g2.height, g3.height)))
check("T8: hidden travels with src2 to dest3", g3.hidden is True and g1.hidden is False)
check("T8: row banding travels with src3 to dest2", g2.s == 1, repr(g2.s))

shutil.rmtree(tmp, ignore_errors=True)
print()
if FAILS:
    print("TOY TESTS FAILED:", FAILS)
    sys.exit(1)
print("TOY TESTS PASSED")