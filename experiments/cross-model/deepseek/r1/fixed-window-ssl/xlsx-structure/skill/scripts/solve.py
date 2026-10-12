#!/usr/bin/env python3
"""
xlsx-structure solver: structural transform of a workbook with strict
preservation of everything outside the declared transformation.

The realized v01 transformation is a block-local stable sort by evaluated
key column (descending, ties keep original order), with formula-bearing
columns declared by the task evaluated row-locally and written as plain
numbers.  The transform itself is PLUGGABLE via the TRANSFORM setting
(default "block_stable_sort"; "filter_by_key" and "append_by_copy" are
synthetic alternatives exercised by toy_tests.py; a custom callable with
the same contract can be supplied).  Every transform shares the same
safe machinery:

  * Mutate ONLY a workbook loaded with data_only=False.  A data_only=True
    workbook is a read-only cached-value reference, never saved.
  * Snapshot the pristine rows BEFORE any write; write every destination
    row entirely from ONE snapshot source row (opaque row units).  Never
    read live cells during the write phase.
  * Evaluate only the task-declared formula columns (EVAL_COLS).  A string
    in another column is data, never arithmetic.
  * Evaluator precedence: the _Parser docstring is the AUTHORITATIVE grammar
    (Excel semantics: negation > ^, left-associative ^, % tightest);
    scripts/verify.py's unit probes are the executable guard.  Supported
    additions beyond v01's bare arithmetic: Excel functions (SUM, AVERAGE,
    MIN, MAX, PRODUCT, COUNT, ABS, INT, ROUND, MOD) and single-row ranges
    (B1:C1).  Cross-row ranges, whole-column ranges, and functions outside
    this list are NOT supported: they raise EvalError and fall back to the
    cached value, or abort loudly when no cache exists.
"""
import math
import re
import sys

import openpyxl

# --------------------------------------------------------------------------
# Per-case configuration (edit for the variant at hand)
# --------------------------------------------------------------------------
INPUT = "/app/input.xlsx"
OUTPUT = "/app/output.xlsx"
SHEET = "Sheet1"
R1, R2, C1, C2 = 1, 85, 1, 8        # declared answer range (1-based, inclusive)
KEY_COL = 4                          # column whose evaluated values drive the sort
DESCENDING = True                    # False for ascending
EVAL_COLS = (4, 5, 6)                # task-declared formula-bearing columns to evaluate
MERGE_POLICY = "travel_with_anchor"  # "travel_with_anchor" | "error"

# Pluggable transform semantics.  String names select built-ins:
#   "block_stable_sort" -> v01 (KEY_COL/DESCENDING)
#   "filter_by_key"     -> synthetic: keep KEY_COL > FILTER_THRESHOLD packed
#                          to block tops, vacated rows cleared (toy only)
#   "append_by_copy"    -> synthetic: copy APPEND_SOURCE_ROWS into the first
#                          empty rows starting at APPEND_DEST_START (toy only)
# or supply a callable transform(blocks, snap, evaluated, ws, ref_ws) ->
# (plan: dict dest_row->src_row, clears: list of dest rows to blank).
TRANSFORM = "block_stable_sort"
FILTER_THRESHOLD = 100.0            # used by "filter_by_key"
APPEND_SOURCE_ROWS = ()             # used by "append_by_copy"
APPEND_DEST_START = None            # used by "append_by_copy"
# --------------------------------------------------------------------------

_NUM = r"(?:\d+\.?\d*|\.\d+)(?:[eE][+-]?\d+)?"
_REF = r"[A-Za-z]{1,3}\d+"
_IDENT = r"[A-Za-z]+"
_TOKEN_RE = re.compile(r"\s*(?:(%s)|(%s)|(%s)|([+\-*/^()%%:,]))" % (_NUM, _REF, _IDENT))


class EvalError(ValueError):
    pass


def _tokenize(text):
    i, toks = 0, []
    while i < len(text):
        m = _TOKEN_RE.match(text, i)
        if not m:
            raise EvalError("bad token at %r" % text[i:])
        num, ref, ident, op = m.groups()
        if num is not None:
            toks.append(("num", float(num)))
        elif ref is not None:
            toks.append(("ref", ref.upper()))
        elif ident is not None:
            toks.append(("ident", ident.upper()))
        else:
            toks.append(("op", op))
        i = m.end()
    return toks


def _round_excel(x, n):
    """Excel ROUND: half away from zero."""
    p = 10 ** int(n)
    return math.copysign(math.floor(abs(x) * p + 0.5) / p, x)


def _is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


class _LenientResolver:
    """Classification-only resolver for is_formula_text: every ref/range
    resolves, so the test is PARSEABILITY, never arithmetic content."""

    @staticmethod
    def cell(ref):
        return 0.0

    @staticmethod
    def range(ref1, ref2):
        return [0.0, 0.0]


_LENIENT = _LenientResolver()


def _apply_function(name, args):
    """args: numbers and/or single-row range lists (numbers with possible None)."""
    flat = [v for a in args for v in (a if isinstance(a, list) else [a])
            if _is_num(v)]
    if name == "SUM":
        return sum(flat)
    if name == "PRODUCT":
        p = 1.0
        for v in flat:
            p *= v
        return p
    if name == "COUNT":
        return float(len(flat))
    if name in ("MIN", "MAX"):
        if not flat:
            raise EvalError("%s() of empty range" % name)
        return (min if name == "MIN" else max)(flat)
    if name == "AVERAGE":
        if not flat:
            raise EvalError("AVERAGE() of empty range")
        return sum(flat) / len(flat)
    nums = [a for a in args if not isinstance(a, list)]
    if name == "ABS":
        if len(nums) != 1 or not _is_num(nums[0]):
            raise EvalError("ABS() needs one numeric argument")
        return abs(nums[0])
    if name == "INT":
        if len(nums) != 1 or not _is_num(nums[0]):
            raise EvalError("INT() needs one numeric argument")
        return math.floor(nums[0])
    if name == "ROUND":
        if len(nums) != 2 or not all(_is_num(a) for a in nums):
            raise EvalError("ROUND() needs two numeric arguments")
        return _round_excel(nums[0], nums[1])
    if name == "MOD":
        if len(nums) != 2 or not all(_is_num(a) for a in nums):
            raise EvalError("MOD() needs two numeric arguments")
        if nums[1] == 0:
            raise EvalError("MOD() division by zero")  # Excel #DIV/0! -> fallback path
        return nums[0] % nums[1]
    raise EvalError("unsupported function %s()" % name)


class _Parser:
    """Excel-precedence recursive descent (AUTHORITATIVE grammar).

    Grammar (tightest first):
      postfix := primary ('%')*
      unary   := ('+'|'-') unary | postfix   # negation binds tighter than ^
      power   := unary (('^' unary))*        # LEFT-associative: 2^3^2 = (2^3)^2 = 64
      term    := power (('*'|'/') power)*
      expr    := term (('+'|'-') term)*
    Excel: -2^2 = (-2)^2 = 4;  2^3^2 = 64;  50%^2 = 0.25;  2^-2 = 0.25.
    (Divergence trap: Python's ** is right-associative; Excel's ^ is left.)

    Extensions: primary may be an IDENT'(...)'-call (functions listed in
    _apply_function) and a 'ref':'ref' single-row range (resolved via
    resolver.range; a range reaching arithmetic raises EvalError).
    """

    def __init__(self, toks, resolver):
        self.toks, self.pos = toks, 0
        self.cell = resolver.cell if hasattr(resolver, "cell") else resolver
        self.rng = getattr(resolver, "range", None)

    def peek(self):
        return self.toks[self.pos] if self.pos < len(self.toks) else (None, None)

    def _next(self):
        t = self.peek()
        self.pos += 1
        return t

    def parse(self):
        v = self.expr()
        if self.pos != len(self.toks):
            raise EvalError("trailing tokens")
        return v

    def expr(self):
        v = self.term()
        while True:
            k, op = self.peek()
            if k == "op" and op in ("+", "-"):
                self._next()
                v = self._arith(op, v, self.term())
            else:
                return v

    def term(self):
        v = self.power()
        while True:
            k, op = self.peek()
            if k == "op" and op in ("*", "/"):
                self._next()
                v = self._arith(op, v, self.power())
            else:
                return v

    def power(self):
        v = self.unary()
        while True:  # left fold: (a^b)^c, matching Excel
            k, op = self.peek()
            if k == "op" and op == "^":
                self._next()
                rhs = self.unary()
                if isinstance(v, list) or isinstance(rhs, list):
                    raise EvalError("range in ^")
                if v == 0 and rhs < 0:
                    raise EvalError("0 ^ negative")
                res = v ** rhs
                if isinstance(res, complex):  # (-2)^0.5 -> #NUM! in Excel
                    raise EvalError("negative base ^ fractional exponent")
                v = res
            else:
                return v

    def unary(self):
        k, op = self.peek()
        if k == "op" and op in ("+", "-"):
            self._next()
            v = self.unary()
            if isinstance(v, list):
                raise EvalError("range under unary sign")
            return v if op == "+" else -v
        return self.postfix()

    def postfix(self):
        v = self.primary()
        while True:
            k, op = self.peek()
            if k == "op" and op == "%":
                self._next()
                if isinstance(v, list):
                    raise EvalError("range under %%")
                v = v / 100.0
            else:
                return v

    def _arith(self, op, a, b):
        if isinstance(a, list) or isinstance(b, list):
            raise EvalError("range in arithmetic")
        if op == "+":
            return a + b
        if op == "-":
            return a - b
        if op == "*":
            return a * b
        if b == 0:
            raise EvalError("division by zero")  # Excel #DIV/0! -> fallback path
        return a / b

    def primary(self):
        k, v = self._next()
        if k == "num":
            return v
        if k == "ref":
            k2, v2 = self.peek()
            if k2 == "op" and v2 == ":":  # range: ref ':' ref (single row)
                self._next()
                k3, v3 = self._next()
                if k3 != "ref":
                    raise EvalError("bad range end after %s:" % v)
                if self.rng is None:
                    raise EvalError("range resolver unavailable")
                rng = self.rng(v, v3)
                if rng is None:
                    raise EvalError("unsupported range %s:%s" % (v, v3))
                return rng
            rv = self.cell(v)
            if rv is None:
                raise EvalError("unresolvable ref %s" % v)
            return rv
        if k == "ident":
            name = v
            k2, v2 = self._next()
            if not (k2 == "op" and v2 == "("):
                raise EvalError("unexpected ident %s" % name)
            args = []
            if not (self.peek() == ("op", ")")):
                while True:
                    args.append(self.expr())
                    k3, v3 = self._next()
                    if k3 == "op" and v3 == ",":
                        continue
                    if k3 == "op" and v3 == ")":
                        break
                    raise EvalError("bad argument list")
            return _apply_function(name, args)
        if k == "op" and v == "(":
            val = self.expr()
            k2, v2 = self._next()
            if not (k2 == "op" and v2 == ")"):
                raise EvalError("missing )")
            return val
        raise EvalError("unexpected %r" % (v,))


def evaluate(text, resolver):
    """Evaluate formula-like text; return float or None on failure.

    `resolver` is either a callable(ref)->float/None (plain-cell only) or an
    object with .cell(ref) and optionally .range(ref1, ref2)->list|None.
    """
    if not isinstance(text, str):
        return None
    s = text.strip()
    if not s:
        return None
    expr = s[1:] if s.startswith("=") else s
    try:
        toks = _tokenize(expr)
        if not toks:
            return None
        if len(toks) == 1 and toks[0][0] == "num":
            return toks[0][1]
        result = _Parser(toks, resolver).parse()
        if isinstance(result, list):
            return None  # bare range has no scalar value
        return result
    except EvalError:
        return None


def is_formula_text(v):
    """Formula-like within EVAL_COLS: '=' prefix, or a FULL lenient parse.

    Classification is by parseability, never by mere tokenizeability: a
    string is a formula iff a complete arithmetic parse succeeds (with every
    ref/range resolving via _LENIENT).  This keeps mixed letter/operator
    DATA such as "N/A", "12:30", "E2/TAX" as data (they fail to parse), and
    the explicit carve-outs keep bare references ("B2") and bare words
    ("ON TOP") as data too.  A bare number in a declared column is
    converted.  Only meaningful for columns the task declares
    formula-bearing.
    """
    if not isinstance(v, str):
        return False
    s = v.strip()
    if not s:
        return False
    if s.startswith("="):
        return True
    try:
        toks = _tokenize(s)
    except EvalError:
        return False
    if not toks:
        return False
    if len(toks) == 1 and toks[0][0] in ("ref", "ident"):
        return False
    return evaluate(s, _LENIENT) is not None


def is_number(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def is_blank_row(ws, r, c1, c2):
    return all(ws.cell(row=r, column=c).value is None for c in range(c1, c2 + 1))


def find_blocks(ws, r1, r2, c1, c2):
    """Maximal runs of non-blank rows (a row is blank when all range cells are None)."""
    blocks, cur = [], []
    for r in range(r1, r2 + 1):
        if is_blank_row(ws, r, c1, c2):
            if cur:
                blocks.append(cur)
                cur = []
        else:
            cur.append(r)
    if cur:
        blocks.append(cur)
    return blocks


def snapshot_block(ws, block, c1, c2):
    """Pristine cell state: {(row, col): (value, _style, _hyperlink, _comment)}."""
    snap = {}
    for r in block:
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            snap[(r, c)] = (cell.value, cell._style, cell._hyperlink, cell._comment)
    return snap


class _RowResolver:
    """Production row-local resolver: .cell(ref) and .range(ref1, ref2).

    Same-row references resolve recursively from the snapshot with memo and
    a cycle guard; cross-row references fall back to the cached-value
    workbook (ref_ws).  Ranges must be single-row and on the formula's own
    row (cross-row/whole-column ranges return None -> EvalError -> cache
    fallback or loud abort).
    """

    def __init__(self, row, snap, evaluated, visiting, ref_ws):
        self.row = row
        self.snap = snap
        self.evaluated = evaluated
        self.visiting = visiting
        self.ref_ws = ref_ws
        self._re = re.compile(r"([A-Za-z]{1,3})(\d+)")

    def _cell_rc(self, r, col):
        key = (r, col)
        if key in self.evaluated:
            return self.evaluated[key]
        if key in self.visiting:
            raise EvalError("cyclic ref r%d c%d" % (r, col))
        val = self.snap[key][0]
        if is_number(val):  # value-first: plain numeric cell
            self.evaluated[key] = float(val)
            return self.evaluated[key]
        if isinstance(val, str) and is_formula_text(val):
            self.visiting.add(key)
            try:
                v = evaluate(val, self)
            finally:
                self.visiting.discard(key)
            if v is None:
                cv = self.ref_ws.cell(row=r, column=col).value
                v = cv if is_number(cv) else None
            self.evaluated[key] = v
            return v
        self.evaluated[key] = None
        return None

    def cell(self, ref):
        mc = self._re.fullmatch(ref)
        if not mc:
            return None
        col = openpyxl.utils.column_index_from_string(mc.group(1))
        r = int(mc.group(2))
        if r != self.row:  # cross-row ref: cached value only
            v = self.ref_ws.cell(row=r, column=col).value
            return v if is_number(v) else None
        return self._cell_rc(r, col)

    def range(self, ref1, ref2):
        m1, m2 = self._re.fullmatch(ref1), self._re.fullmatch(ref2)
        if not m1 or not m2:
            return None
        col1 = openpyxl.utils.column_index_from_string(m1.group(1))
        col2 = openpyxl.utils.column_index_from_string(m2.group(1))
        r1, r2 = int(m1.group(2)), int(m2.group(2))
        if r1 != r2 or r1 != self.row or col1 > col2:
            return None  # only single-row ranges on the formula's own row
        return [self._cell_rc(r1, c) for c in range(col1, col2 + 1)]


def evaluate_block(snap, rows, ref_ws, eval_cols, c1, c2, abort_on_failure=True):
    """PRODUCTION row-local evaluation of EVAL_COLS cells (shared by solve and verify).

    Same-row references (and single-row ranges) resolve recursively from the
    snapshot with memo and a cycle guard; cross-row references fall back to
    the cached-value workbook (ref_ws).  A formula-like cell that cannot be
    evaluated and has no cached number is left None; with abort_on_failure=True
    (solver) that is an error rather than silently shipping a formula string.
    Returns {(row, col): float}.
    """
    evaluated = {}
    visiting = set()

    for r in rows:
        for c in eval_cols:
            if not (c1 <= c <= c2):
                continue
            key = (r, c)
            val = snap[key][0]
            if is_formula_text(val):
                resolver = _RowResolver(r, snap, evaluated, visiting, ref_ws)
                ev = evaluate(val, resolver)
                if ev is None:
                    cv = ref_ws.cell(row=r, column=c).value
                    ev = cv if is_number(cv) else None
                if ev is None and abort_on_failure:
                    raise EvalError("cannot evaluate %r at row %d col %d (no cache)"
                                    % (val, r, c))
                evaluated[key] = ev
            elif is_number(val):
                evaluated[key] = float(val)
    return evaluated


# --------------------------------------------------------------------------
# Pluggable transforms: (blocks, snap, evaluated, ws, ref_ws)
#   -> (plan: dict dest_row->src_row, clears: list of dest rows to blank)
# --------------------------------------------------------------------------
def block_stable_sort(blocks, snap, evaluated, ws, ref_ws):
    """v01: per block, stable descending (or ascending) order by evaluated KEY_COL."""
    def sort_key(r):
        ev = evaluated.get((r, KEY_COL))
        if ev is not None:
            return ev
        v = snap[(r, KEY_COL)][0]
        if is_number(v):
            return float(v)
        raise EvalError("key column %d row %d unevaluable" % (KEY_COL, r))

    sign = -1.0 if DESCENDING else 1.0
    plan = {}
    for block in blocks:
        order = sorted(block, key=lambda r: (sign * sort_key(r), r))
        for dest, src in zip(block, order):
            plan[dest] = src
    return plan, []


def _key_value(r, snap, evaluated):
    ev = evaluated.get((r, KEY_COL))
    if ev is not None:
        return ev
    v = snap[(r, KEY_COL)][0]
    return float(v) if is_number(v) else None


def filter_by_key(blocks, snap, evaluated, ws, ref_ws):
    """SYNTHETIC (toy exercise, not v01-verified): keep rows with evaluated
    KEY_COL > FILTER_THRESHOLD, packed to the top of each block; vacated rows
    are cleared (value+style) like deleted rows.  A real filter task would
    declare its own predicate and compaction boundary."""
    plan, clears = {}, []
    for block in blocks:
        keep = [r for r in block
                if (kv := _key_value(r, snap, evaluated)) is not None
                and kv > FILTER_THRESHOLD]
        for d, s in zip(block[:len(keep)], keep):
            plan[d] = s
        clears.extend(block[len(keep):])
    return plan, clears


def append_by_copy(blocks, snap, evaluated, ws, ref_ws):
    """SYNTHETIC (toy exercise, not v01-verified): copy APPEND_SOURCE_ROWS
    (original row numbers) into the first empty rows starting at
    APPEND_DEST_START.  Source rows are left untouched (copy semantics);
    destination rows need not belong to any block."""
    plan = {}
    for dest, src in zip(range(APPEND_DEST_START,
                               APPEND_DEST_START + len(APPEND_SOURCE_ROWS)),
                         APPEND_SOURCE_ROWS):
        plan[dest] = src
    return plan, []


_TRANSFORMS = {
    "block_stable_sort": block_stable_sort,
    "filter_by_key": filter_by_key,
    "append_by_copy": append_by_copy,
}


def main():
    wb = openpyxl.load_workbook(INPUT, data_only=False)   # the ONLY object mutated/saved
    ref = openpyxl.load_workbook(INPUT, data_only=True)   # cached-value reference
    ws, ref_ws = wb[SHEET], ref[SHEET]

    blocks = find_blocks(ws, R1, R2, C1, C2)
    print("blocks:", [len(b) for b in blocks], "transform:", TRANSFORM)

    transform = _TRANSFORMS.get(TRANSFORM) if isinstance(TRANSFORM, str) else TRANSFORM
    if transform is None or not callable(transform):
        raise ValueError("unknown TRANSFORM %r" % (TRANSFORM,))

    # Merged ranges intersecting the write range (unverified strategy on v01:
    # v01 has no merges; handle explicitly when a variant has them).
    affected_merges = []
    if ws.merged_cells.ranges:
        for mr in ws.merged_cells.ranges:
            if mr.min_row <= R2 and mr.max_row >= R1 and mr.min_col <= C2 and mr.max_col >= C1:
                affected_merges.append(mr)
        if affected_merges and MERGE_POLICY == "error":
            raise RuntimeError("merges intersect write range: %s" % affected_merges)
        for mr in affected_merges:
            ws.unmerge_cells(str(mr))

    # One pristine snapshot + one production evaluation for the whole range;
    # resolvers are row-bound, so evaluating all rows at once is equivalent
    # to per-block evaluation (bit-exactness is re-asserted in verify.py).
    data_rows = [r for b in blocks for r in b]
    snap = {}
    for b in blocks:
        snap.update(snapshot_block(ws, b, C1, C2))
    evaluated = evaluate_block(snap, data_rows, ref_ws, EVAL_COLS, C1, C2,
                               abort_on_failure=True)

    plan, clears = transform(blocks, snap, evaluated, ws, ref_ws)

    # --- row-level state travels with row content --------------------------
    # RowDimension entries (height, hidden, outlineLevel, collapsed, row
    # style s, thickBot/top) are ROW-UNIT properties: they move with the
    # content, they do not stay positional.  Snapshot BEFORE any write — a
    # live read mid-write sees already-traveled state (the checkpoint-2
    # corruption class at row granularity).  customHeight/customFormat are
    # computed properties (derive from ht/s) and travel implicitly.
    dim_rows = set(plan.values()) | set(plan.keys()) | set(clears)
    rowdims = {}
    for r in dim_rows:
        d = ws.row_dimensions.get(r)  # non-side-effecting read (no auto-create)
        if d is None:
            rowdims[r] = None
        else:
            rowdims[r] = (d.height, d.hidden, d.outlineLevel, d.collapsed,
                          d.s, d.thickBot, d.thickTop)
    for dest in clears:
        ws.row_dimensions.pop(dest, None)  # no ghost height/banding on removed rows
    for dest, src in plan.items():
        sd = rowdims.get(src)
        if sd is None:
            ws.row_dimensions.pop(dest, None)
            continue
        rd = ws.row_dimensions[dest]  # auto-created with index=dest
        rd.height, rd.hidden, rd.outlineLevel, rd.collapsed = sd[0], sd[1], sd[2], sd[3]
        if sd[4] not in (None, 0):  # non-default row style travels (s=0 is implicit)
            rd._style = wb._cell_styles[sd[4]]  # `.s =` has no setter in 3.1.x
        if sd[5] is not None:
            rd.thickBot = sd[5]
        if sd[6] is not None:
            rd.thickTop = sd[6]

    # Write each dest row's full content from its snapshot src (opaque unit).
    for dest, src in plan.items():
        for c in range(C1, C2 + 1):
            val, style, hl, cm = snap[(src, c)]
            # Substitution is SCOPED to EVAL_COLS by definition.  The resolver
            # also memoizes same-row numeric cells referenced by formulas
            # (B/C when D = B+C), so `evaluated` may hold keys outside
            # EVAL_COLS; those must never replace the snapshot value —
            # non-declared columns are written verbatim (value AND type).
            ev = evaluated.get((src, c)) if c in EVAL_COLS else None
            if ev is not None:
                val = ev
            cell = ws.cell(row=dest, column=c)
            cell.value = val  # public setter infers data_type via _bind_value
            cell._style = style
            cell._hyperlink = hl
            cell._comment = cm

    # Vacated rows (filter/compaction): fully removed like deleted rows.
    for dest in clears:
        for c in range(C1, C2 + 1):
            ws._cells.pop((dest, c), None)  # cell leaves the serialized grid

    # Re-merge with each merge traveling with its anchor row's content.
    src_to_dest = {s: d for d, s in plan.items()}
    cleared = set(clears)
    for mr in list(affected_merges):
        if mr.min_row in src_to_dest:
            a = src_to_dest[mr.min_row]
            span = mr.max_row - mr.min_row
            ws.merge_cells(start_row=a, start_column=mr.min_col,
                           end_row=a + span, end_column=mr.max_col)
            affected_merges.remove(mr)
        elif mr.min_row in cleared:
            affected_merges.remove(mr)  # anchor row removed with the merge

    if affected_merges:
        print("merges with anchors outside the transform left untouched:", affected_merges)

    wb.save(OUTPUT)
    print("saved", OUTPUT)


if __name__ == "__main__":
    sys.exit(main())