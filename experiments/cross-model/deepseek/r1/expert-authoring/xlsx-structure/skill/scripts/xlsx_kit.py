"""xlsx_kit: helpers for in-place structural transformations of .xlsx workbooks.

Design goals (mirroring the family invariants):
  * Load with openpyxl and mutate in place; never rebuild the workbook from
    scratch, so styles, filters, merged cells, panes, widths and hidden state
    survive the round-trip automatically.
  * Only ever write to the task's declared answer range; everything outside it
    is preserved by not being touched.
  * Formulas are replaced by *evaluated numeric values* in the answer range
    (this environment has no formula engine: no LibreOffice, no pandas, no
    formulas/pycel/xlcalculator - only openpyxl is available), using the tiny
    row-local arithmetic evaluator below.
  * Provide dynamic region/block detection (blank-separated blocks, bounding
    boxes, header rows) rather than hard-coded addresses.

Depends only on: openpyxl (>=3.1), Python >=3.8 stdlib.
"""

import re
import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string


# --------------------------------------------------------------------------
# formula evaluation (row-local arithmetic only)
# --------------------------------------------------------------------------

class FormulaError(ValueError):
    """Raised when a formula cannot be evaluated by the row-local evaluator."""


_TOKEN_RE = re.compile(
    r"\s*(?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)"
    r"|(?P<ref>\$?[A-Za-z]{1,3}\$?\d+)"
    r"|(?P<op>[+\-*/^()])"
)

# Only these characters may appear in a formula expression that the
# evaluator can handle. Anything else (functions, string literals, commas,
# '&', comparisons) falls outside row-local arithmetic.
_ALLOWED_EXPR_CHARS = set("0123456789.+-*/^()=$ ABCDEFGHIJKLMNOPQRSTUVWXYZ")


def ref_col_row(ref):
    """'B12' or '$B$12' -> (col_number, row_number)."""
    m = re.fullmatch(r"\$?([A-Za-z]+)\$?(\d+)", ref)
    if not m:
        raise FormulaError(f"bad reference {ref!r}")
    return column_index_from_string(m.group(1).upper()), int(m.group(2))


def tokenize(expr):
    """Split an arithmetic expression into ('num'|'ref'|'op', value) tokens."""
    tokens, pos = [], 0
    while pos < len(expr):
        m = _TOKEN_RE.match(expr, pos)
        if not m:
            raise FormulaError(f"unsupported token near {expr[pos:pos + 20]!r}")
        pos = m.end()
        kind = m.lastgroup
        val = m.group()
        if kind == "num":
            tokens.append(("num", float(val)))
        elif kind == "ref":
            tokens.append(("ref", val.replace("$", "").upper()))
        else:
            tokens.append(("op", val))
    return tokens


def _eval(tokens, resolve):
    """Recursive-descent evaluator over tokenized arithmetic."""
    pos = [0]

    def peek():
        return tokens[pos[0]] if pos[0] < len(tokens) else None

    def parse_expr():
        v = parse_term()
        while peek() and peek()[0] == "op" and peek()[1] in "+-":
            op = peek()[1]
            pos[0] += 1
            rhs = parse_term()
            v = v + rhs if op == "+" else v - rhs
        return v

    def parse_term():
        v = parse_factor()
        while peek() and peek()[0] == "op" and peek()[1] in "*/":
            op = peek()[1]
            pos[0] += 1
            rhs = parse_factor()
            v = v * rhs if op == "*" else v / rhs
        return v

    def parse_factor():
        # Excel-like precedence: unary minus binds tighter than '^'
        # (-2^2 == 4), and '^' is left-associative like Excel (2^3^2 == 64).
        def parse_power():
            v = parse_unary()
            while peek() and peek()[0] == "op" and peek()[1] == "^":
                pos[0] += 1
                v = v ** parse_unary()
            return v

        def parse_unary():
            t = peek()
            if t is not None and t[0] == "op" and t[1] in "+-":
                pos[0] += 1
                v = parse_unary()
                return -v if t[1] == "-" else v
            return parse_atom()

        def parse_atom():
            t = peek()
            if t is None:
                raise FormulaError("unexpected end of formula")
            if t[0] == "op" and t[1] == "(":
                pos[0] += 1
                v = parse_expr()
                if not (peek() and peek()[0] == "op" and peek()[1] == ")"):
                    raise FormulaError("missing ')'")
                pos[0] += 1
                return v
            if t[0] == "num":
                pos[0] += 1
                return t[1]
            if t[0] == "ref":
                pos[0] += 1
                return resolve(t[1])
            raise FormulaError(f"unexpected token {t!r}")
        return parse_power()

    v = parse_expr()
    if pos[0] != len(tokens):
        raise FormulaError(f"trailing tokens {tokens[pos[0]:]!r}")
    return v


def looks_like_formula(text):
    """Heuristic: a string cell that is arithmetic text rather than a label."""
    s = text.strip()
    if not s:
        return False
    if s.startswith("="):
        s = s[1:]
    if set(s.upper()) - _ALLOWED_EXPR_CHARS:
        return False
    return bool(re.search(r"[A-Z]{1,3}\$?\d+", s, re.IGNORECASE))


def numeric_from_cell(ws, coord, cache=None, depth=0):
    """Return the numeric value of a single cell, evaluating formulas.

    `cache` maps coordinate -> already-evaluated numeric value; use it when a
    row's columns must be evaluated in dependency order (e.g. F = D + E where
    D itself is a formula): put D's evaluated value in the cache before
    resolving F.
    """
    if depth > 40:
        raise FormulaError("recursion too deep (cycle?)")
    coord = coord.upper()
    if cache is not None and coord in cache:
        return cache[coord]
    col = re.sub(r"\d", "", coord)
    row = int(re.sub(r"\D", "", coord))
    if cache is not None and col in cache:
        return cache[col]
    value = stored_value(ws, row, col)
    if value is None:
        raise FormulaError(f"{coord} is empty")
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if looks_like_formula(s):
            if s.startswith("="):
                s = s[1:]
            tokens = tokenize(s)
            for kind, val in tokens:
                if kind == "ref":
                    c, r = ref_col_row(val)
                    if r != row:
                        raise FormulaError(
                            f"cross-row reference {val} in row-local formula")
            resolver = lambda ref: numeric_from_cell(
                ws, ref, cache=cache, depth=depth + 1)
            return _eval(tokens, resolver)
        try:
            return float(s)
        except ValueError:
            raise FormulaError(f"{coord} is not numeric: {value!r}")
    raise FormulaError(f"{coord} is not numeric: {value!r}")


def eval_row_columns(ws, row, colnames):
    """Evaluate several columns of one row, in dependency order.

    colnames: list of column letters, e.g. ['D', 'E', 'F']; earlier columns
    are resolved first and their results feed later formulas (row-local).
    Returns a dict keyed BOTH by full coordinate ('D1', 'E1', 'F1') and by
    bare column letter ('D', 'E', 'F') for the given row.
    Raises FormulaError if any column is neither numeric nor evaluable.
    """
    cache = {}
    for col in colnames:
        coord = f"{col}{row}"
        val = numeric_from_cell(ws, coord, cache=cache)
        cache[coord] = val
        cache[col] = val
    return cache


def cached_value_of(input_path, coord, sheet=None):
    """Fallback: the value Excel cached for a formula cell (data_only=True).
    Returns None when the file stores no cached value."""
    wb = load_workbook(input_path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.active
    return ws[coord].value


# --------------------------------------------------------------------------
# blocks, regions, rows
# --------------------------------------------------------------------------


def parse_range(ref):
    """'A1:H85' -> {'cols': ['A',..,'H'], 'first_col': 'A', 'last_col': 'H',
    'top': 1, 'bottom': 85}. Handles absolute markers and either order."""
    m = re.fullmatch(r"\$?([A-Za-z]+)\$?(\d+):\$?([A-Za-z]+)\$?(\d+)", ref.strip())
    if not m:
        raise ValueError(f"not an A1-style range: {ref!r}")
    c1 = column_index_from_string(m.group(1).upper())
    c2 = column_index_from_string(m.group(3).upper())
    letters = [get_column_letter(c) for c in
               range(min(c1, c2), max(c1, c2) + 1)]
    return {"cols": letters,
            "first_col": letters[0],
            "last_col": letters[-1],
            "col_from": min(c1, c2),
            "col_to": max(c1, c2),
            "top": min(int(m.group(2)), int(m.group(4))),
            "bottom": max(int(m.group(2)), int(m.group(4)))}


def stored_extent(ws):
    """(min_row, min_col, max_row, max_col) over cells the file actually
    stores (value or style). Differs from data_bbox: includes styled-but-empty
    cells that keep the sheet dimension alive. Non-mutating: never calls
    ws.cell() on coordinates that may not exist, because cell() CREATES a new
    Cell, which grows ws._cells and therefore the dimension openpyxl writes."""
    if not ws._cells:
        return None
    rows = [c[0] for c in ws._cells]
    cols = [c[1] for c in ws._cells]
    return (min(rows), min(cols), max(rows), max(cols))

def stored_cell(ws, r, col):
    """Non-mutating cell access: returns the Cell stored at (row, col) or
    None. Never call ws.cell()/ws['A1'] for reads: both INSERT a new Cell
    into ws._cells when the coordinate is absent, which changes the stored
    extent and the dimension string openpyxl writes on save (observed:
    A1:L250 silently became A1:Q250)."""
    c = column_index_from_string(col) if isinstance(col, str) else col
    return ws._cells.get((r, c))


def stored_value(ws, r, col):
    """Non-mutating value read; None when the cell is absent/empty."""
    cell = stored_cell(ws, r, col)
    return cell.value if cell is not None else None


def read_row(ws, r, cols):
    return [stored_value(ws, r, c) for c in cols]


def write_row(ws, r, cols, values):
    """Write one row of plain Python values (numbers stay numbers, strings stay
    strings). Never pass a string starting with '=': openpyxl would store it
    as a formula."""
    for c, v in zip(cols, values):
        cell = ws[f"{c}{r}"]
        if v is None:
            cell.value = None
        elif isinstance(v, str) and v.startswith("="):
            raise ValueError(f"refusing to write formula string to {c}{r}: {v!r}")
        else:
            cell.value = v
    return


def is_blank_row(ws, r, cols):
    return all((v is None or (isinstance(v, str) and not v.strip()))
               for v in read_row(ws, r, cols))


def iter_blank_blocks_list(ws, top, bottom, cols):
    """Convenience wrapper returning a full list of blocks."""
    blocks, cur = [], None
    for r in range(top, bottom + 1):
        if is_blank_row(ws, r, cols):
            cur = None
        else:
            if cur is None:
                cur = (r, [])
                blocks.append(cur)
            cur[1].append(r)
    return blocks


def block_layout(ws, top, bottom, cols):
    """[(start_row, n_rows), ...] — the block structure of a range. Use the
    same function on input and output to compare topologies, and on the input
    to derive the group boundaries the hidden tests check group-by-group."""
    return [(s, len(rows)) for s, rows in
            iter_blank_blocks_list(ws, top, bottom, cols)]


def snapshot_rows(ws, rows, cols):
    """Pre-read (row_number -> [values]) for rows about to be permuted or
    compacted IN PLACE. Mandatory before any reordering writes: a later read
    of a source row can otherwise see already-rewritten content (a swap of
    two rows would silently duplicate one of them). Non-mutating."""
    return {r: read_row(ws, r, cols) for r in rows}


def values_equal(a, b, tol=1e-6):
    """Row/cell equality for verification: None == None, numbers compare with
    tolerance (repr differs across writers: 224.9 vs 224.89999999999998),
    everything else exact."""
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, bool) or isinstance(b, bool):
        return a is b if isinstance(a, bool) and isinstance(b, bool) else False
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float_close(a, b, tol)
    return a == b


def rows_equal(row_a, row_b, tol=1e-6):
    """List-of-values row equality using values_equal per position."""
    return len(row_a) == len(row_b) and all(
        values_equal(a, b, tol) for a, b in zip(row_a, row_b))


def blank_cells(ws, rows, cols):
    """Set every cell of the given rows to None (the compaction tail;
    writes are the only part of the kit that may create cells, and these
    rows exist within the declared answer range)."""
    for r in rows:
        write_row(ws, r, cols, [None] * len(cols))


def compact_rows(ws, R, kept_rows, dst_start=None, evaluated=None):
    """Write `kept_rows` (in the given order) consecutively starting at
    `dst_start` (default R['top']; pass R['top']+1 when a header row occupies
    R['top'] and must stay), then blank all remaining rows through R['bottom']
    so no stale rows survive. Header rows are never written or blanked — the
    caller places the start of the survivor block explicitly instead, since
    header placement (inside vs above the range) varies by task.
    `evaluated` optionally maps row -> {col_letter: number} (from
    eval_row_columns) to substitute formula columns. Uses snapshot_rows
    internally (in-place compaction must never re-read rewritten rows).
    Returns the written survivor value-lists in order (the expected 'answer
    group' for verification)."""
    cols = R["cols"]
    if dst_start is None:
        dst_start = R["top"]
    data = {r: list(v) for r, v in snapshot_rows(ws, kept_rows, cols).items()}
    for r, vals in data.items():
        if evaluated and r in evaluated:
            for col, num in evaluated[r].items():
                if col in cols:
                    vals[cols.index(col)] = num
    for i, src in enumerate(kept_rows):
        write_row(ws, dst_start + i, cols, data[src])
    blank_cells(ws, list(range(dst_start + len(kept_rows), R["bottom"] + 1)),
                cols)
    return [data[r] for r in kept_rows]


def transform_block_sort(ws, R, key_col, eval_cols, reverse=True):
    """The observed block-local sort, made mechanically safe:
    for every blank-separated block in range R, stable-sort rows by the
    evaluated value of key_col (eval_cols are evaluated per row in dependency
    order and substituted into the written rows; key values come from the
    same evaluation). Uses snapshot_rows BEFORE any write (a permutation read
    off the live sheet would corrupt swapped rows). Returns:
        (blocks, {start_row: expected sequence of value-lists})
    so the caller can verify group-by-group against the saved output."""
    cols = R["cols"]
    blocks = iter_blank_blocks_list(ws, R["top"], R["bottom"], cols)
    ev, snap = {}, {}
    for _, rows in blocks:
        for r in rows:
            ev[r] = eval_row_columns(ws, r, eval_cols)
            snap[r] = read_row(ws, r, cols)
    expected = {}
    for start, rows in blocks:
        order = [t[0] for t in stable_sort_rows(
            rows, lambda r: ev[r][key_col], reverse=reverse)]
        seq = []
        for i, src in enumerate(order):
            vals = list(snap[src])
            for col in eval_cols:
                vals[cols.index(col)] = ev[src][col]
            write_row(ws, start + i, cols, vals)
            seq.append(vals)
        expected[start] = seq
    return blocks, expected


def data_bbox(ws, max_row=None, max_col=None, include_styled_empty=False):
    """Bounding box of cells with a non-None value (or all stored cells when
    include_styled_empty=True, which can reveal styled-but-empty regions that
    extend the sheet dimension)."""
    min_r = min_c = 10 ** 9
    max_r = max_c = 0
    found = False
    if include_styled_empty:
        for (r, c) in ws._cells:
            found = True
            min_r, min_c = min(min_r, r), min(min_c, c)
            max_r, max_c = max(max_r, r), max(max_c, c)
    else:
        for (r, c), cell in ws._cells.items():
            if cell.value is not None:
                found = True
                min_r, min_c = min(min_r, r), min(min_c, c)
                max_r, max_c = max(max_r, r), max(max_c, c)
    return (min_r, min_c, max_r, max_c) if found else None


def header_row_index(ws, top, bottom, cols):
    """First row in range whose non-empty cells are mostly strings
    (a candidate header). Returns row number or None."""
    for r in range(top, bottom + 1):
        vals = [v for v in read_row(ws, r, cols)
                if v is not None and str(v).strip() != ""]
        if not vals:
            continue
        str_share = sum(1 for v in vals if isinstance(v, str)
                        and not looks_like_formula(v)) / len(vals)
        if str_share >= 0.5:
            return r
    return None


def stable_sort_rows(rows, key, reverse):
    """sorted(..., key=key, reverse=reverse) is stable in CPython:
    equal keys keep original relative order. `rows` = list of row numbers;
    returns list of (row_number, key_value) ordered per the sort."""
    keyed = [(r, key(r)) for r in rows]
    keyed.sort(key=lambda t: t[1], reverse=reverse)
    return keyed


# --------------------------------------------------------------------------
# preservation: snapshots and verification
# --------------------------------------------------------------------------

def sheet_meta(ws):
    """Everything about a sheet that must survive a transformation."""
    return {
        "freeze_panes": ws.freeze_panes,
        "auto_filter_ref": ws.auto_filter.ref,
        "merged": sorted(str(rng) for rng in ws.merged_cells.ranges),
        "col_dims": {
            k: (d.width, d.min, d.max, d.hidden, d.style,
                getattr(d, "bestFit", None),
                getattr(d, "customWidth", None))
            for k, d in ws.column_dimensions.items()
        },
        "row_dims": {k: d.height for k, d in ws.row_dimensions.items()},
        "topLeftCell": ws.sheet_view.topLeftCell,
        "showGridLines": ws.sheet_view.showGridLines,
        "tabSelected": ws.sheet_view.tabSelected,
        "selection": [(s.activeCell, s.sqref) for s in ws.sheet_view.selection],
        "defaultRowHeight": ws.sheet_format.defaultRowHeight,
        "sheet_state": ws.sheet_state,
        "sheet_properties": (ws.sheet_properties.codeName,
                             ws.sheet_properties.tabColor
                             if hasattr(ws.sheet_properties, "tabColor") else None),
    }


def assert_sheet_meta_equal(label, before, after):
    problems = []
    for key in before:
        if before[key] != after[key]:
            problems.append(f"{label}: {key} changed:\n  before={before[key]}\n  after={after[key]}")
    return problems


def snapshot_range(ws, top, bottom, cols):
    """Snapshot (value, data_type, semantic_style, number_format) for every
    cell the file actually STORES inside the rectangle. Non-mutating: only
    iterates ws._cells, never calls ws.cell() (calling it creates a Cell,
    which changes what openpyxl later writes back — observed: the sheet
    dimension silently grew from A1:L250 to A1:Q250).
    Cells that are not stored are implicitly (None, blank, default style) and
    likewise on the assertion side.
    """
    snap = {}
    for (r, c), cell in ws._cells.items():
        if top <= r <= bottom and get_column_letter(c) in cols:
            snap[(r, c)] = (cell.value, cell.data_type,
                            cell_style_key(cell), cell.number_format)
    return snap


_REGION_KEY = ("__outside_snapshot_region__",)


def snapshot_outside(ws, region):
    """Snapshot of every STORED cell outside the declared answer region.
    region: dict from parse_range() or tuple (col_from, top, col_to, bottom).
    Non-mutating (same reason as snapshot_range). The snapshot records the
    excluded region under _REGION_KEY so assert_outside_unchanged can ignore
    the answer cells that legitimately changed."""
    if isinstance(region, dict):
        c1, r1, c2, r2 = region["col_from"], region["top"], \
            region["col_to"], region["bottom"]
    else:
        c1, r1, c2, r2 = region
    snap = {_REGION_KEY: (c1, r1, c2, r2)}
    for (r, c), cell in ws._cells.items():
        if c1 <= c <= c2 and r1 <= r <= r2:
            continue  # the answer range: expected to change
        snap[(r, c)] = (cell.value, cell.data_type,
                        cell_style_key(cell), cell.number_format)
    return snap


def cell_style_key(cell):
    """Semantic style summary, stable across openpyxl style renumbering:
    font name/size/bold/italic, fill pattern + fg/bg rgb, number format,
    border sides, alignment."""
    f = cell.font
    fill = cell.fill
    al = cell.alignment
    b = cell.border
    def side(s):
        return (None if s is None else (s.style,
                                        None if s.color is None else s.color.rgb))
    return {
        "numfmt": cell.number_format,
        "font": (f.name, f.size, f.bold, f.italic),
        "fill": (fill.patternType,
                 None if not fill.fgColor else (fill.fgColor.rgb, fill.fgColor.theme),
                 None if not fill.bgColor else (fill.bgColor.rgb, fill.bgColor.theme)),
        "border": (side(b.left), side(b.right), side(b.top), side(b.bottom)),
        "align": (al.horizontal, al.vertical, al.wrap_text),
    }


def fill_key(cell):
    """(patternType, fg rgb, fg theme) for fill-colour based selection."""
    fill = cell.fill
    return (fill.patternType,
            None if not fill.fgColor else fill.fgColor.rgb,
            None if not fill.fgColor else fill.fgColor.theme)


def assert_range_unchanged(ws, snap, label="range"):
    """Compare a snapshot against the live sheet UNSAFELY when the sheet is the
    same object the snapshot came from. For after-save checks prefer
    assert_outside_unchanged on the reloaded workbook. This helper does not
    create new cells either."""
    problems = []
    for (r, c), (val, dt, style, fmt) in snap.items():
        cell = stored_cell(ws, r, c)
        if cell is None:
            if val is not None or fmt != "General":
                problems.append(
                    f"{label}: stored cell vanished at {get_column_letter(c)}{r}")
            else:
                continue
        if cell.value != val:
            problems.append(
                f"{label}: value changed at {get_column_letter(c)}{r}: {val!r} -> {cell.value!r}")
        if cell.data_type != dt:
            problems.append(f"{label}: data type changed at {get_column_letter(c)}{r}")
        if cell_style_key(cell) != style:
            problems.append(f"{label}: style changed at {get_column_letter(c)}{r}")
        if cell.number_format != fmt:
            problems.append(f"{label}: number format changed at {get_column_letter(c)}{r}")
    return problems


def assert_outside_unchanged(snap, ws):
    """Check that after the transform + save + reload, every cell the input
    stored outside the answer range still matches, and that no NEW non-empty
    or styled cell appeared outside the answer range.
    The answer range itself is ignored (it is the declared transform region)."""
    problems = []
    (c1, r1, c2, r2) = snap.pop(_REGION_KEY, (0, 0, 0, 0))
    for (r, c), (val, dt, style, fmt) in snap.items():
        if (r, c) in ws._cells:
            cell = ws._cells[(r, c)]
            if cell.value != val:
                problems.append(
                    f"outside {get_column_letter(c)}{r}: value {val!r} -> {cell.value!r}")
            if cell_style_key(cell) != style:
                problems.append(f"outside {get_column_letter(c)}{r}: style changed")
            if cell.number_format != fmt:
                problems.append(f"outside {get_column_letter(c)}{r}: number format changed")
        else:
            if val is not None:
                problems.append(f"outside {get_column_letter(c)}{r}: stored cell vanished")
    for (r, c), cell in ws._cells.items():
        if c1 <= c <= c2 and r1 <= r <= r2:
            continue  # declared answer region
        if (r, c) not in snap:
            if cell.value is not None or cell.has_style or cell.style_id != 0:
                problems.append(
                    f"outside {get_column_letter(c)}{r}: new stored cell appeared "
                    f"(value={cell.value!r})")
    return problems


def assert_no_formulas(ws, top, bottom, cols):
    """After writing evaluated numbers, no cell in the range may be a formula
    (data_type 'f') and every non-empty cell must be int/float/None/string.
    Non-mutating read."""
    problems = []
    for r in range(top, bottom + 1):
        for c in cols:
            cell = stored_cell(ws, r, c)
            if cell is None:
                continue
            if cell.data_type == "f":
                problems.append(f"formula left at {c}{r}: {cell.value!r}")
                continue
            if cell.value is not None and not isinstance(
                    cell.value, (int, float, str, bool)):
                problems.append(f"non-scalar value at {c}{r}: {cell.value!r}")
    return problems


def run_checks(problems_by_label, stream=sys.stdout):
    """Print and count a dict {label: [problem strings]}."""
    total = 0
    for label, probs in problems_by_label.items():
        if probs:
            print(f"[{label}] {len(probs)} problem(s)", file=stream)
            for p in probs[:25]:
                print("   ", p, file=stream)
            total += len(probs)
    print(f"verification: {total} problem(s) across {len(problems_by_label)} check group(s)",
          file=stream)
    return total


def reload_workbook(path):
    """Default load mode: formulas visible as strings (data_type 'f')."""
    return load_workbook(path)


# --------------------------------------------------------------------------
# inventory, alignment and matrix helpers
# --------------------------------------------------------------------------

def describe_workbook(wb_or_path):
    """One-call inventory dump (the Step-2 report). Prints per sheet:
    dimensions vs stored extent, freeze panes, auto filter, merged ranges,
    column/row dims, sheet view state; then the data types stored per column
    of the first data-bearing sheet. Safe to run read-only."""
    wb = load_workbook(wb_or_path) if isinstance(wb_or_path, str) else wb_or_path
    print("sheets:", wb.sheetnames)
    print("defined names:", list(wb.defined_names))
    print("hidden sheets:", {w.title: w.sheet_state for w in wb.worksheets
                              if w.sheet_state != "visible"})
    for ws in wb.worksheets:
        ext = stored_extent(ws)
        print(f"--- {ws.title!r}: dimension={ws.dimensions} stored_extent={ext}")
        print("  freeze:", ws.freeze_panes, "| autofilter:", ws.auto_filter.ref,
              "| merged:", sorted(str(r) for r in ws.merged_cells.ranges))
        print("  view:", f"topLeft={ws.sheet_view.topLeftCell}",
              f"tab={ws.sheet_view.tabSelected}",
              f"grid={ws.sheet_view.showGridLines}",
              "selection=", [(s.activeCell, s.sqref)
                             for s in ws.sheet_view.selection],
              f"defaultRowHeight={ws.sheet_format.defaultRowHeight}")
        print("  col_dims:", {k: (d.width, d.min, d.max, d.hidden,
                                 d.customWidth, d.bestFit, d.style)
                              for k, d in ws.column_dimensions.items()})
        if ws.row_dimensions:
            print("  row_dims:", {k: d.height for k, d in ws.row_dimensions.items()})
        # column data-type profile across stored cells
        profile = {}
        for (r, c), cell in ws._cells.items():
            if cell.value is not None:
                profile.setdefault(get_column_letter(c), set()).add(cell.data_type)
        if profile:
            print("  column types:", {k: sorted(v) for k, v in sorted(profile.items())})
    return wb


def header_map(ws, row, cols=None):
    """Map normalized header text -> column letter for header/ID-aligned
    writes. Normalization: strip + casefold. Ignores empty cells."""
    cols = cols or get_column_letter(ws.max_column)
    if isinstance(cols, str):
        cols = [get_column_letter(c) for c in
                range(1, column_index_from_string(cols) + 1)]
    out = {}
    for col in cols:
        v = stored_value(ws, row, col)
        if v is not None and str(v).strip():
            out[str(v).strip().casefold()] = col
    return out


def last_data_row(ws, col, top=1, bottom=None):
    """Last row number <=bottom whose single column has a non-empty value;
    the append boundary for 'append below the last row' geometry."""
    bottom = bottom if bottom is not None else ws.max_row
    last = top - 1
    for r in range(top, bottom + 1):
        v = stored_value(ws, r, col)
        if v is not None and str(v).strip() != "":
            last = r
    return last


def read_matrix(ws, rows, cols):
    """Read a rectangular region as a list of row-value lists."""
    return [read_row(ws, r, cols) for r in rows]


def write_matrix(ws, anchor_row, anchor_col, matrix):
    """Write a list-of-lists at a (row number, column letter) anchor,
    preserving each value's Python type and refusing formula strings.
    Only cells with non-None values are written."""
    col_idx = column_index_from_string(anchor_col)
    for i, row in enumerate(matrix):
        for j, v in enumerate(row):
            if v is None:
                continue
            if isinstance(v, str) and v.startswith("="):
                raise ValueError(
                    f"refusing to write formula string at "
                    f"{get_column_letter(col_idx + j)}{anchor_row + i}: {v!r}")
            ws.cell(row=anchor_row + i, column=col_idx + j).value = v
    return None


def float_close(a, b, tol=1e-6):
    """Tolerant float equality for verification (repr differs across writers:
    224.9 vs 224.89999999999998 are the same value)."""
    if a is None or b is None:
        return a is b
    return abs(a - b) <= tol