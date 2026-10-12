#!/usr/bin/env python3
"""Self-test for the xlsx-structure skill package.

Builds synthetic fixtures, patches each solver's config block (SRC/DST + the
semantics-specific keys), runs the patched solver copy in a subprocess against
the SHIPPED snapshot (PYTHONPATH = <skill>/scripts, default /skill — override
with SELFTEST_SKILL_DIR to test another snapshot), and checks every output
against LITERAL expectations typed from the fixture spec — never values
recomputed through the solver machinery, so a bug that is shared between a
solver and the battery cannot pretend to be correct.

Run (first line of Step 4 in SKILL.md):
    cd /skill/scripts && python3 selftest.py
"""
import os
import shutil
import subprocess
import sys

import openpyxl

SKILL_DIR = os.environ.get('SELFTEST_SKILL_DIR', '/skill')
SCRIPTS = os.path.join(SKILL_DIR, 'scripts')
WORK = '/app/.xlsx_selftest'
FAILURES = []


def check(cond, msg):
    if not cond:
        FAILURES.append(msg)
        print('FAIL:', msg)
    else:
        print('ok  :', msg)


def simplify(v):
    """Primitive normalizer for literal comparisons (no shared-solver logic)."""
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def build(path, fn):
    """Build a fixture workbook with openpyxl (formulas are strings WITHOUT
    cached values — the evaluator must produce the literal expectations)."""
    if os.path.exists(path):
        os.remove(path)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    fn(wb, ws)
    wb.save(path)
    return path


def patch_config(text, **kv):
    """Replace `KEY = <oldrepr>` config lines in a solver copy."""
    for key, val in kv.items():
        import re
        pat = re.compile(r'^' + re.escape(key) + r' = .*$', re.M)
        new_text, n = pat.subn(f'{key} = {val!r}', text)
        assert n == 1, f'config line {key} not found exactly once'
        text = new_text
    return text


def run_solver(name, config, expect_fail=False, expect_err=''):
    """Copy <script> from the shipped snapshot, patch its config, run it."""
    os.makedirs(WORK, exist_ok=True)
    target = os.path.join(WORK, name)
    with open(os.path.join(SCRIPTS, name)) as f:
        text = f.read()
    text = patch_config(text, **config)
    with open(target, 'w') as f:
        f.write(text)
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([SCRIPTS, WORK]))
    r = subprocess.run([sys.executable, target], cwd=WORK, env=env,
                       capture_output=True, text=True, timeout=120)
    if expect_fail:
        if r.returncode == 0:
            FAILURES.append(f'{name}: expected failure, but exited 0')
            print('FAIL:', f'{name}: expected failure, but exited 0')
            return r
        if expect_err and expect_err not in (r.stdout + r.stderr):
            FAILURES.append(f'{name}: expected {expect_err!r} in output')
            print('FAIL:', f'{name}: expected {expect_err!r} in output')
            return r
        print('ok  :', f'{name}: failed as expected ({expect_err or "non-zero"})')
        return r
    if r.returncode != 0:
        FAILURES.append(f'{name}: solver exited {r.returncode}\n{r.stdout}{r.stderr}')
        print('FAIL:', f'{name}: solver exited {r.returncode}\n{r.stdout}{r.stderr}')
    return r


def read_grid(path, r1, r2, c1, c2, sheet='Sheet1'):
    """Primitive value grid (floats normalized to ints when integral)."""
    ws = openpyxl.load_workbook(path)[sheet]
    return [[simplify(ws.cell(row=r, column=c).value) for c in range(c1, c2 + 1)]
            for r in range(r1, r2 + 1)]


def scenario_mini_v01():
    """Block sort with ties + row-local formulas; expectations are literals."""
    def fn(wb, ws):
        for r, (n, b, c, e, h) in enumerate([
                ('ALPHA', 1, 2, 10, 1), ('BETA', 5, 5, 20, 2), ('GAMMA', 4, 6, 5, 3)], start=1):
            ws.cell(row=r, column=1).value = n
            ws.cell(row=r, column=2).value = b
            ws.cell(row=r, column=3).value = c
            ws.cell(row=r, column=4).value = f'=B{r}+C{r}'
            ws.cell(row=r, column=5).value = e
            ws.cell(row=r, column=6).value = f'=D{r}+E{r}'
            ws.cell(row=r, column=7).value = 1
            ws.cell(row=r, column=8).value = h
        for r, (n, b, c, e, h) in enumerate([
                ('DELTA', 10, 20, 5, 1), ('EPSILON', 1, 1, 1, 2)], start=5):
            ws.cell(row=r, column=1).value = n
            ws.cell(row=r, column=2).value = b
            ws.cell(row=r, column=3).value = c
            ws.cell(row=r, column=4).value = f'=B{r}+C{r}'
            ws.cell(row=r, column=5).value = e
            ws.cell(row=r, column=6).value = f'=D{r}+E{r}'
            ws.cell(row=r, column=7).value = 2
            ws.cell(row=r, column=8).value = h
    inp = build(os.path.join(WORK, 'in_mini.xlsx'), fn)
    # tie: BETA and GAMMA both evaluate to D=10 -> stable order keeps BETA first
    run_solver('solve_block_sort.py',
               dict(SRC=inp, DST=os.path.join(WORK, 'out_mini.xlsx'),
                    TABLE='A1:H6', FORMULA_COLS=(4, 5, 6)))
    out = os.path.join(WORK, 'out_mini.xlsx')
    names = [row[0] for row in read_grid(out, 1, 6, 1, 1)]
    check(names == ['BETA', 'GAMMA', 'ALPHA', None, 'DELTA', 'EPSILON'],
          f'mini-v01 order is literal (desc 10,10,3; tie stable): {names}')
    dvals = [row[3] for row in read_grid(out, 1, 6, 1, 4)]
    check(dvals == [10, 10, 3, None, 30, 2], f'mini-v01 evaluated D literals: {dvals}')
    fvals = [row[5] for row in read_grid(out, 1, 6, 1, 6)]
    check(fvals == [30, 15, 13, None, 35, 3], f'mini-v01 evaluated F literals: {fvals}')
    hs = [row[7] for row in read_grid(out, 1, 6, 1, 8)]
    check(hs == [2, 3, 1, None, 1, 2], f'mini-v01 H column rides with rows: {hs}')


def scenario_evaluator_edges():
    """Aggregate semantics through a real sort; failures must raise."""
    def fn(wb, ws):
        # row1: AVERAGE(10, blank)  = 10 (blank IGNORED, n=1) -> D=10
        ws['A1'] = 'P'; ws['B1'] = 10; ws['C1'] = None
        ws['D1'] = '=AVERAGE(B1:C1)'
        ws['E1'] = 100
        # row2: AVERAGE(text, 20)   = 20 (text IGNORED, n=1) -> D=20
        ws['A2'] = 'Q'; ws['B2'] = 'abc'; ws['C2'] = 20
        ws['D2'] = '=AVERAGE(B2:C2)'
        ws['E2'] = 100
        # row3: MIN(blank, 7)       = 7 (blank IGNORED)      -> D=7
        ws['A3'] = 'R'; ws['B3'] = None; ws['C3'] = 7
        ws['D3'] = '=MIN(B3:C3)'
        ws['E3'] = 100
        for r in (1, 2, 3):
            ws.cell(row=r, column=6).value = f'=D{r}+E{r}'
    inp = build(os.path.join(WORK, 'in_eval.xlsx'), fn)
    out = os.path.join(WORK, 'out_eval.xlsx')
    run_solver('solve_block_sort.py',
               dict(SRC=inp, DST=out, TABLE='A1:F3', FORMULA_COLS=(4, 5, 6)))
    names = [row[0] for row in read_grid(out, 1, 3, 1, 1)]
    check(names == ['Q', 'P', 'R'],
          f'aggregate sort order follows Excel semantics (D 20>10>7): {names}')
    dvals = [row[3] for row in read_grid(out, 1, 3, 1, 4)]
    check(dvals == [20, 10, 7],
          f'aggregate evaluated D literals (blanks/text ignored): {dvals}')

    # scalar arithmetic on text MUST raise (implement-or-raise)
    def fn2(wb, ws):
        ws['A1'] = 'X'; ws['B1'] = 5; ws['C1'] = 'abc'
        ws['D1'] = '=B1+C1'; ws['E1'] = 1; ws['F1'] = '=D1+E1'
    inp2 = build(os.path.join(WORK, 'in_text.xlsx'), fn2)
    run_solver('solve_block_sort.py',
               dict(SRC=inp2, DST=os.path.join(WORK, 'out_text.xlsx'),
                    TABLE='A1:F1', FORMULA_COLS=(4, 5, 6)),
               expect_fail=True, expect_err='#VALUE')

    # AVERAGE over a range with no numbers MUST raise (#DIV/0!)
    def fn3(wb, ws):
        ws['A1'] = 'X'; ws['B1'] = None; ws['C1'] = None
        ws['D1'] = '=AVERAGE(B1:C1)'; ws['E1'] = 1; ws['F1'] = '=D1+E1'
    inp3 = build(os.path.join(WORK, 'in_div0.xlsx'), fn3)
    run_solver('solve_block_sort.py',
               dict(SRC=inp3, DST=os.path.join(WORK, 'out_div0.xlsx'),
                    TABLE='A1:F1', FORMULA_COLS=(4, 5, 6)),
               expect_fail=True, expect_err='#DIV/0')


def scenario_filter_compact():
    def fn(wb, ws):
        rows = [('ALPHA', 1, 2, 99, False), ('BETA', 10, 20, 45, True),
                ('GAMMA', 1, 2, 99, False), ('DELTA', 7, 8, 88, True),
                ('EPSILON', 5, 5, 50, False), ('ZETA', 9, 1, 10, True)]
        for i, (a, b, c, e, mark) in enumerate(rows, start=1):
            ws.cell(row=i, column=1).value = a
            ws.cell(row=i, column=2).value = b
            ws.cell(row=i, column=3).value = c
            ws.cell(row=i, column=4).value = f'=B{i}+C{i}'
            ws.cell(row=i, column=5).value = e
            ws.cell(row=i, column=6).value = f'=D{i}+E{i}'
            if mark:
                ws.cell(row=i, column=1).fill = openpyxl.styles.PatternFill(
                    start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    inp = build(os.path.join(WORK, 'in_fc.xlsx'), fn)
    out = os.path.join(WORK, 'out_fc.xlsx')
    run_solver('solve_filter_compact.py',
               dict(SRC=inp, DST=out, TABLE='A1:F6',
                    PREDICATES=[{'kind': 'fill', 'rgb': 'FFFF00'}],
                    FORMULA_COLS=(4, 5, 6), DST_ANCHOR=None, DELETE_TAIL=False))
    names = [row[0] for row in read_grid(out, 1, 6, 1, 1)]
    check(names == ['BETA', 'DELTA', 'ZETA', None, None, None],
          f'filter keeps exactly the filled rows, in source order: {names}')
    dv = [row[3] for row in read_grid(out, 1, 3, 1, 4)]
    check(dv == [30, 15, 10], f'kept rows carry evaluated D literals: {dv}')


def scenario_extract_into_merged():
    """Extraction whose anchor is INSIDE a merged range: no crash, merge kept."""
    def fn(wb, ws):
        rows = [('ALPHA', 1, 2, 99, False), ('BETA', 10, 20, 45, True),
                ('GAMMA', 1, 2, 99, False), ('DELTA', 7, 8, 88, True)]
        for i, (a, b, c, e, mark) in enumerate(rows, start=1):
            ws.cell(row=i, column=1).value = a
            ws.cell(row=i, column=2).value = b
            ws.cell(row=i, column=3).value = c
            ws.cell(row=i, column=4).value = f'=B{i}+C{i}'
            ws.cell(row=i, column=5).value = e
            ws.cell(row=i, column=6).value = f'=D{i}+E{i}'
            if mark:
                ws.cell(row=i, column=1).fill = openpyxl.styles.PatternFill(
                    start_color='FFFF00', end_color='FFFF00', fill_type='solid')
        ws.merge_cells('A12:B13')   # anchor A12 lies inside this merge
    inp = build(os.path.join(WORK, 'in_exm.xlsx'), fn)
    out = os.path.join(WORK, 'out_exm.xlsx')
    run_solver('solve_filter_compact.py',
               dict(SRC=inp, DST=out, TABLE='A1:F4',
                    PREDICATES=[{'kind': 'fill', 'rgb': 'FFFF00'}],
                    FORMULA_COLS=(4, 5, 6), DST_ANCHOR='A12',
                    DELETE_TAIL=False))
    ws_out = openpyxl.load_workbook(out)['Sheet1']
    # anchor cell holds the first kept record; A13/B12/B13 are shadowed by the
    # RESTORED merged region (unreadable in Excel's model too -> None here)
    check(str(ws_out['A12'].value) == 'BETA'
          and ws_out['A13'].value is None
          and ws_out['B12'].value is None,
          'merged anchor shows the first kept row; shadowed cells read None')
    check('A12:B13' in [str(r) for r in ws_out.merged_cells.ranges],
          'merged region at the anchor is restored')
    # record columns outside the merged box survive for BOTH kept rows
    rest = [[simplify(ws_out.cell(row=r, column=c).value) for c in range(3, 7)]
            for r in (12, 13)]
    check(rest == [[20, 30, 45, 75], [8, 15, 88, 103]],
          f'kept records beyond the merged box are intact: {rest}')
    src_ok = read_grid(out, 1, 4, 1, 1)
    check([g[0] for g in src_ok] == ['ALPHA', 'BETA', 'GAMMA', 'DELTA'],
          'source table untouched by extraction')


def scenario_reconcile():
    def fn(wb, ws):
        for i, v in [(2, 'ID3'), (3, 'ID1'), (4, 'ID5'), (5, 'ID7'), (6, 'ID2'), (7, 'ID1')]:
            ws.cell(row=i, column=2).value = v   # duplicate ID1 AND missing ID7?
        for r, sid, name, val in [(2, 'ID1', 'ALPHA', 10), (3, 'ID5', 'GAMMA', 30),
                                  (4, 'ID2', 'BETA', 20), (5, 'ID3', 'DELTA', 40)]:
            ws.cell(row=r, column=4).value = sid
            ws.cell(row=r, column=5).value = name
            ws.cell(row=r, column=6).value = val
    inp = build(os.path.join(WORK, 'in_rec.xlsx'), fn)
    out = os.path.join(WORK, 'out_rec.xlsx')
    run_solver('solve_reconcile_by_id.py',
               dict(SRC=inp, DST=out, ID_LIST='B2:B7', SPARSE='D2:F5',
                    SPARSE_ID_COL=1, DST_ANCHOR='A10', WRITE_ID=True))
    ids = [row[0] for row in read_grid(out, 10, 15, 1, 1)]
    check(ids == ['ID3', 'ID1', 'ID5', 'ID7', 'ID2', 'ID1'],
          f'reconciliation follows the ID list incl. duplicate + missing: {ids}')
    names = [row[1] for row in read_grid(out, 10, 15, 1, 2)]
    check(names == ['DELTA', 'ALPHA', 'GAMMA', None, 'BETA', 'ALPHA'],
          f'matched records land under their IDs; missing ID7 blank: {names}')


def scenario_append():
    def fn(wb, ws):
        for c, h in [(1, 'Name'), (2, 'Total'), (3, 'Rank')]:
            ws.cell(row=1, column=c).value = h
        ws['A2'] = 'EXISTING'; ws['B2'] = 100; ws['C2'] = 1
        src = wb.create_sheet('Sheet2')
        for c, h in [(1, 'Rank'), (2, 'Name'), (3, 'Total')]:
            src.cell(row=1, column=c).value = h
        for r, (rank, name, total) in enumerate([(2, 'AAA', 10), (1, 'BBB', 20)], start=2):
            src.cell(row=r, column=1).value = rank
            src.cell(row=r, column=2).value = name
            src.cell(row=r, column=3).value = total
    inp = build(os.path.join(WORK, 'in_app.xlsx'), fn)
    out = os.path.join(WORK, 'out_app.xlsx')
    # DEST_C1/DEST_C2 and SRC_C1/SRC_C2 each share one config line -> patch
    # with paired patterns, then patch the single-name keys
    text = open(os.path.join(SCRIPTS, 'solve_append_by_header.py')).read()
    import re
    text, n = re.subn(r'^DEST_C1, DEST_C2 = .*$', 'DEST_C1, DEST_C2 = 1, 3',
                      text, flags=re.M)
    assert n == 1
    text, n = re.subn(r'^SRC_C1, SRC_C2 = .*$', 'SRC_C1, SRC_C2 = 1, 3',
                      text, flags=re.M)
    assert n == 1
    text = patch_config(text, SRC=inp, DST=out, SRC_SHEET='Sheet2')
    with open(os.path.join(WORK, 'solve_append_by_header.py'), 'w') as f:
        f.write(text)
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([SCRIPTS, WORK]))
    r = subprocess.run([sys.executable, os.path.join(WORK, 'solve_append_by_header.py')],
                       cwd=WORK, env=env, capture_output=True, text=True, timeout=120)
    if r.returncode != 0:
        FAILURES.append(f'solve_append_by_header.py exited {r.returncode}\n'
                        f'{r.stdout}{r.stderr}')
        print('FAIL:', f'solve_append_by_header.py exited {r.returncode}\n'
                       f'{r.stdout}{r.stderr}')
    grid = read_grid(out, 1, 4, 1, 3)
    check(grid == [['Name', 'Total', 'Rank'], ['EXISTING', 100, 1],
                   ['AAA', 10, 2], ['BBB', 20, 1]],
          f'append maps by header name (shuffled source columns): {grid}')


def scenario_transpose():
    def fn(wb, ws):
        for r in (1, 2, 3):
            for c in (1, 2):
                ws.cell(row=r, column=c).value = (r - 1) * 2 + c
                if c == 2:
                    ws.cell(row=r, column=c).font = openpyxl.styles.Font(bold=True)
    inp = build(os.path.join(WORK, 'in_tr.xlsx'), fn)
    out = os.path.join(WORK, 'out_tr.xlsx')
    run_solver('solve_transpose.py',
               dict(SRC=inp, DST=out, SRC_RANGE='A1:B3', DST_ANCHOR='E5'))
    grid5 = read_grid(out, 5, 6, 5, 7)
    check(grid5 == [[1, 3, 5], [2, 4, 6]],
          f'transposed matrix literals at E5:G6: {grid5}')
    ws_out = openpyxl.load_workbook(out)['Sheet1']
    check(ws_out['E6'].font.bold and ws_out['G6'].font.bold,
          'styles transposed with cells')


def scenario_strict_keys():
    """Sort keys implement-or-raise: text/blank keys FAIL LOUDLY with the cell
    address; numeric text coerces (Excel scalar) and sorts by value."""
    def fn(wb, ws):
        ws['A1'] = 'P'; ws['B1'] = 10
        ws['A2'] = 'Q'; ws['B2'] = 5
        ws['A3'] = 'R'; ws['B3'] = 7
    inp = build(os.path.join(WORK, 'in_strict_ok.xlsx'), fn)
    out = os.path.join(WORK, 'out_strict_ok.xlsx')
    run_solver('solve_block_sort.py',
               dict(SRC=inp, DST=out, TABLE='A1:B3', SORT_COL=2, DESC=True,
                    FORMULA_COLS=()))
    names = [row[0] for row in read_grid(out, 1, 3, 1, 1)]
    check(names == ['P', 'R', 'Q'],
          f'numeric keys sort strictly (10>7>5): {names}')

    def fn2(wb, ws):   # text in the sort column -> must fail loudly
        ws['A1'] = 'P'; ws['B1'] = 10
        ws['A2'] = 'Q'; ws['B2'] = 'abc'
    inp2 = build(os.path.join(WORK, 'in_strict_text.xlsx'), fn2)
    run_solver('solve_block_sort.py',
               dict(SRC=inp2, DST=os.path.join(WORK, 'out_strict_text.xlsx'),
                    TABLE='A1:B2', SORT_COL=2, DESC=True, FORMULA_COLS=()),
               expect_fail=True, expect_err='non-numeric sort key')

    def fn3(wb, ws):   # blank key -> must fail loudly
        ws['A1'] = 'P'; ws['B1'] = 10
        ws['A2'] = 'Q'; ws['B2'] = None
    inp3 = build(os.path.join(WORK, 'in_strict_blank.xlsx'), fn3)
    run_solver('solve_block_sort.py',
               dict(SRC=inp3, DST=os.path.join(WORK, 'out_strict_blank.xlsx'),
                    TABLE='A1:B2', SORT_COL=2, DESC=True, FORMULA_COLS=()),
               expect_fail=True, expect_err='blank sort key')

    def fn5(wb, ws):   # numeric PREDICATE on a text cell -> fail loudly too
        rows = [('ALPHA', 10), ('BETA', 'abc')]
        for i, (n, d) in enumerate(rows, start=1):
            ws.cell(row=i, column=1).value = n
            ws.cell(row=i, column=4).value = d
    inp5 = build(os.path.join(WORK, 'in_strict_pred.xlsx'), fn5)
    run_solver('solve_filter_compact.py',
               dict(SRC=inp5, DST=os.path.join(WORK, 'out_strict_pred.xlsx'),
                    TABLE='A1:D2',
                    PREDICATES=[{'kind': 'numeric', 'col': 4, 'op': '>',
                                 'value': 5}],
                    FORMULA_COLS=(), DST_ANCHOR=None, DELETE_TAIL=False),
               expect_fail=True, expect_err='non-numeric sort key')

    def fn4(wb, ws):   # numeric text coerces (Excel scalar), sorts by value
        ws['A1'] = 'P'; ws['B1'] = 10
        ws['A2'] = 'Q'; ws['B2'] = ' 42 '
        ws['A3'] = 'R'; ws['B3'] = 7
    inp4 = build(os.path.join(WORK, 'in_strict_coerce.xlsx'), fn4)
    out4 = os.path.join(WORK, 'out_strict_coerce.xlsx')
    run_solver('solve_block_sort.py',
               dict(SRC=inp4, DST=out4, TABLE='A1:B3', SORT_COL=2, DESC=True,
                    FORMULA_COLS=()))
    names4 = [row[0] for row in read_grid(out4, 1, 3, 1, 1)]
    check(names4 == ['Q', 'P', 'R'],
          f'numeric text keys coerce and sort by value: {names4}')


def main():
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(WORK)
    for fn in (scenario_mini_v01, scenario_evaluator_edges,
               scenario_filter_compact, scenario_extract_into_merged,
               scenario_reconcile, scenario_append, scenario_transpose,
               scenario_strict_keys):
        fn()
        print('---')
    if FAILURES:
        print(f'SELFTEST FAILED: {len(FAILURES)} problem(s)')
        sys.exit(1)
    print('SELFTEST PASSED: shipped snapshot certifies '
          'block-sort/ties/formulas/filter/extract-into-merge/reconcile/'
          'append/transpose against literal expectations')


if __name__ == '__main__':
    main()