#!/usr/bin/env python3
"""Block-local sort solver for the xlsx-structure family (anchor semantics:
blank-separated variable-length blocks, stable sort of complete rows by the
evaluated numeric value of one column, formulas evaluated row-locally to
numbers, in-place output in the same range, everything else preserved).

Config block: adapt SRC/DST/SHEET/TABLE/SORT_COL/DESC/FORMULA_COLS per task
after inspecting the workbook (scripts/inspect_workbook.py). For other family
semantics (filter/append/reconcile/extract/transpose) use their matching
solve_*.py script; the primitives live in xlsx_common.py.
"""
import sys

from openpyxl.utils import get_column_letter

from xlsx_common import (assert_standard_parts, capture_merges, clean_num,
                         evaluate_formula, find_blocks, load_dual, parse_a1,
                         restore_merges, snapshot_rows, strict_key, verify_output,
                         write_rows)

# ---------------- config (per task) ----------------
SRC = '/app/input.xlsx'
DST = '/app/output.xlsx'
SHEET = 'Sheet1'
TABLE = 'A1:H85'        # declared transformed range ("answer range")
SORT_COL = 4            # 1-based column letter->index (D = 4); or header name
DESC = True             # descending (False -> ascending)
FORMULA_COLS = (4, 5, 6)  # columns that may hold formulas / '=...' text
# ----------------------------------------------------

# optional: resolve SORT_COL from a header name here, e.g.
# header_row = 1; SORT_COL = next(c for c in range(c1, c2+1)
#                                if str(ws.cell(row=header_row, column=c).value).strip() == 'D')


def main():
    # preflight 0: openpyxl cannot model drawings/charts/media/VBA -> refuse
    # in-place editing for such inputs (raw sheet-XML surgery would be needed)
    assert_standard_parts(SRC)
    wb, wbv = load_dual(SRC)
    ws = wb[SHEET]
    r1, r2, c1, c2 = parse_a1(TABLE)
    blocks = find_blocks(ws, r1, r2, c1, c2)
    print('blocks:', blocks)
    # preflight 1: non-anchor cells of merged ranges are read-only MergedCell
    # objects -> unmerge any merge touching the write footprint, restore after
    saved_merges = capture_merges(ws, range(r1, r2 + 1), range(c1, c2 + 1))
    if saved_merges:
        print('captured merges to restore:', len(saved_merges))
    cache = {}
    for rs, re in blocks:
        rows = list(range(rs, re + 1))
        # 1) evaluate keys + formula cells on ORIGINAL rows (row-local, in
        #    dependency order: D before F, so F sees evaluated D)
        keys, evaluated = {}, {}
        for r in rows:
            rv = {c: ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)}
            for c in sorted(FORMULA_COLS):  # ascending column order = dep order
                v = rv.get(c)
                if isinstance(v, str) and v.startswith('='):
                    rv[c] = evaluate_formula(v, ws, r, rv, cache)
            evaluated[r] = rv
            keys[r] = strict_key(rv[SORT_COL],
                                 coord=f'{get_column_letter(SORT_COL)}{r}')
        # 2) stable sort: Python's sorted preserves original order on ties
        order = sorted(rows, key=lambda r: keys[r], reverse=DESC)
        # 3) style-preserving in-place rewrite
        snap = snapshot_rows(ws, rows, c1, c2)
        write_rows(ws, rows, snap, order,
                   formula_cols=FORMULA_COLS, evaluated=evaluated,
                   start_col=c1)
    restore_merges(ws, saved_merges)
    wb.save(DST)
    print('saved', DST)

    # 4) full verification battery: spans + single blank runs, sort order +
    #    tie stability (mapped through row identity), row multisets, no
    #    formulas left, D-F numeric, preservation of everything outside the
    #    answer range, styles, widths, heights, dims, merges, panes, filter,
    #    print area, zip manifest
    problems = verify_output(SRC, DST, SHEET, r1, r2, c1, c2,
                             formula_cols=FORMULA_COLS, spans=blocks,
                             sort_col=SORT_COL, desc=DESC,
                             blank_separator=1)
    if problems:
        print('VERIFICATION FAILED:')
        for p in problems[:30]:
            print(' -', p)
        sys.exit(1)
    print('verification passed: blocks intact, keys sorted, ties stable, '
          'no formulas left, everything outside the answer range preserved')


if __name__ == '__main__':
    main()