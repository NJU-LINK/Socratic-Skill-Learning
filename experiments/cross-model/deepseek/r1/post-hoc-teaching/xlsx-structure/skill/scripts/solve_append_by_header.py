#!/usr/bin/env python3
"""Header-aligned append solver for the xlsx-structure family.

Destination table has a header row; source records (with their own header row
or positional) are appended below the destination's last data row, each value
placed under the matching destination header. Matching is by normalized header
name (strip/lower/collapse-space); values for destination columns that the
source does not provide stay None; extra source columns without a destination
header are skipped (or appended by position when the instruction says
positional correspondence).
"""
import re
from copy import copy

from openpyxl.utils import column_index_from_string

from xlsx_common import (assert_standard_parts, capture_merges, clean_num,
                         evaluate_formula, header_map, load_dual, parse_a1,
                         restore_merges, verify_manifest, verify_workbook_state)

# ---------------- config (per task) ----------------
SRC = '/app/input.xlsx'
DST = '/app/output.xlsx'
DEST_SHEET = 'Sheet1'
DEST_HEADER_ROW = 1      # header row of the destination table
DEST_C1, DEST_C2 = 1, 8  # destination header span
SRC_SHEET = 'Sheet1'
SRC_HEADER_ROW = 1       # source header row; None -> positional via SRC_COLS
SRC_FIRST_DATA = 2       # first source data row
SRC_C1, SRC_C2 = 1, 8    # source record span (used when SRC_HEADER_ROW is None)
POSITIONAL = False       # True -> ignore header names, map by column position
FORMULA_COLS = ()        # source columns holding formulas to evaluate
# ----------------------------------------------------


def main():
    assert_standard_parts(SRC)
    wb, wbv = load_dual(SRC)
    d_ws = wb[DEST_SHEET]
    s_ws = wb[SRC_SHEET]
    # source header -> source column map
    src_head = (header_map(s_ws, SRC_HEADER_ROW, SRC_C1, SRC_C2)
                if SRC_HEADER_ROW is not None else None)
    dst_head = header_map(d_ws, DEST_HEADER_ROW, DEST_C1, DEST_C2)
    # last destination data row
    last = DEST_HEADER_ROW
    for r in range(DEST_HEADER_ROW + 1, d_ws.max_row + 1):
        if any(d_ws.cell(row=r, column=c).value is not None
               for c in range(DEST_C1, DEST_C2 + 1)):
            last = r
        else:
            break
    print('destination header:', dst_head, '| last data row:', last)
    cache = {}
    n_src = max(1, s_ws.max_row - (SRC_FIRST_DATA - 1))
    saved = capture_merges(d_ws, range(last + 1, last + n_src + 2),
                           range(DEST_C1, DEST_C2 + 1))
    if SRC_HEADER_ROW is None:
        src_cols = list(range(SRC_C1, SRC_C2 + 1))
    else:
        src_cols = sorted(src_head.values())
    if POSITIONAL:
        col_map = {DEST_C1 + i: (src_cols[i] if i < len(src_cols) else None)
                   for i in range(DEST_C2 - DEST_C1 + 1)}
    else:
        col_map = {}
        for dc in range(DEST_C1, DEST_C2 + 1):
            h = d_ws.cell(row=DEST_HEADER_ROW, column=dc).value
            key = re.sub(r'\s+', ' ', str(h)).strip().lower() if h is not None else ''
            col_map[dc] = src_head.get(key) if src_head is not None and key else None
    if SRC_HEADER_ROW is not None:
        src_data_end = s_ws.max_row
        for r in range(SRC_FIRST_DATA, s_ws.max_row + 1):
            if all(s_ws.cell(row=r, column=c).value is None
                   for c in src_cols):
                src_data_end = r - 1
                break
    else:
        src_data_end = s_ws.max_row
    appended = 0
    appended_rows = []   # source row used for each appended destination row
    for r in range(SRC_FIRST_DATA, src_data_end + 1):
        # skip blank source rows
        if all(s_ws.cell(row=r, column=c).value is None
               for c in src_cols):
            continue
        dst_r = last + 1 + appended
        appended_rows.append(r)
        for dc, sc in col_map.items():
            if sc is None:
                continue
            s = s_ws.cell(row=r, column=sc)
            d = d_ws.cell(row=dst_r, column=dc)
            val = s.value
            if sc in FORMULA_COLS and isinstance(val, str) and val.startswith('='):
                rv = {c2_: s_ws.cell(row=r, column=c2_).value
                      for c2_ in range(SRC_C1, SRC_C2 + 1)}
                for c2_ in sorted(FORMULA_COLS):
                    v2 = rv.get(c2_)
                    if isinstance(v2, str) and v2.startswith('='):
                        rv[c2_] = evaluate_formula(v2, s_ws, r, rv, cache)
                val = clean_num(rv.get(sc, val))
            d.value = val
            d._style = copy(s._style)
        appended += 1
    restore_merges(d_ws, saved)
    wb.save(DST)
    print('appended', appended, 'rows after row', last, '->', DST)

    # verification: appended cells land under matching headers; the original
    # table (headers + pre-existing rows) is unchanged
    out, outv = load_dual(DST)
    ows = out[DEST_SHEET]
    problems = []
    for r in range(DEST_HEADER_ROW, last + 1):
        for c in range(DEST_C1, DEST_C2 + 1):
            a = d_ws.cell(row=r, column=c).value
            b = ows.cell(row=r, column=c).value
            if isinstance(a, float):
                a = round(a, 9)
            if isinstance(b, float):
                b = round(b, 9)
            if a != b:
                problems.append(f'original table cell changed at {r},{c}')
    for i, src_r in enumerate(appended_rows):
        r = last + 1 + i
        for dc, sc in col_map.items():
            got = ows.cell(row=r, column=dc).value
            if sc is None:
                if got is not None:
                    problems.append(f'appended cell at {r},{dc} should be empty')
                continue
            exp = s_ws.cell(row=src_r, column=sc).value
            if isinstance(exp, float):
                exp = round(exp, 9)
            if isinstance(got, float):
                got = round(got, 9)
            if got != exp and not (isinstance(exp, str) and exp.startswith('=')):
                problems.append(f'appended cell mismatch at {r},{dc}: '
                                f'{got!r} != {exp!r}')
    problems += verify_workbook_state(d_ws, ows, DEST_HEADER_ROW, last,
                                      DEST_C1, DEST_C2, issues=[])
    try:
        verify_manifest(SRC, DST)
    except Exception as e:
        problems.append(f'manifest: {e}')
    if problems:
        print('VERIFICATION FAILED:')
        for p in problems[:30]:
            print(' -', p)
        raise SystemExit(1)
    print('verification passed: appended rows are header-aligned, original '
          'table preserved')


if __name__ == '__main__':
    main()