#!/usr/bin/env python3
"""Layout transposition solver for the xlsx-structure family.

Reads the source matrix (values + styles) and writes it transposed to the
destination anchor: output cell (dst_r + j - 1, dst_c + i - 1) gets the
source cell (r_i, c_j). Formula columns are evaluated to numbers. With
SWAP_SIZE=True, source column widths become destination row heights and
source row heights become destination column widths (visual fidelity) — note
that changes row/column sizes by design.
"""
import re
from copy import copy

from openpyxl.utils import column_index_from_string, get_column_letter

from xlsx_common import (assert_standard_parts, capture_merges, clean_num,
                         evaluate_formula, load_dual, parse_a1, restore_merges,
                         verify_manifest, verify_workbook_state)

# ---------------- config (per task) ----------------
SRC = '/app/input.xlsx'
DST = '/app/output.xlsx'
SHEET = 'Sheet1'
SRC_RANGE = 'A1:F20'     # matrix to transpose
DST_ANCHOR = 'A1'        # top-left of the transposed output ('Sheet2!A1' ok)
FORMULA_COLS = ()        # source columns holding formulas to evaluate
SWAP_SIZE = False        # True -> widths <-> heights for the output block
# (True changes sizes by design; only enable when the task asks for visual
#  fidelity, and then relax the sizes part of verify_workbook_state)
# ----------------------------------------------------


def main():
    assert_standard_parts(SRC)
    wb, wbv = load_dual(SRC)
    ws = wb[SHEET]
    r1, r2, c1, c2 = parse_a1(SRC_RANGE)
    height = r2 - r1 + 1
    width = c2 - c1 + 1
    dst_sheet, anchor = SHEET, DST_ANCHOR
    if '!' in anchor:
        dst_sheet, anchor = anchor.split('!', 1)
    if dst_sheet not in wb.sheetnames:
        wb.create_sheet(dst_sheet)
    col_s, row_s = re.match(r'([A-Za-z]+)(\d+)', anchor).groups()
    dst_r0, dst_c0 = int(row_s), column_index_from_string(col_s.upper())
    d_ws = wb[dst_sheet]
    if dst_sheet == SHEET:
        d_ws = ws
    saved = capture_merges(d_ws, range(dst_r0, dst_r0 + width),
                           range(dst_c0, dst_c0 + height))
    cache = {}
    # evaluate formula columns on original rows first
    ev = {}
    for r in range(r1, r2 + 1):
        rv = {c: ws.cell(row=r, column=c).value for c in range(c1, c2 + 1)}
        for c in sorted(FORMULA_COLS):
            v = rv.get(c)
            if isinstance(v, str) and v.startswith('='):
                rv[c] = evaluate_formula(v, ws, r, rv, cache)
        ev[r] = rv
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            s = ws.cell(row=rr, column=cc)
            d = d_ws.cell(row=dst_r0 + (cc - c1), column=dst_c0 + (rr - r1))
            d._style = copy(s._style)
            v = s.value
            if cc in FORMULA_COLS and isinstance(v, str) and v.startswith('='):
                v = clean_num(ev[rr].get(cc, v))
            d.value = v
    if SWAP_SIZE:
        # source col widths -> dst row heights; source row heights -> dst col widths
        for j in range(width):
            w = ws.column_dimensions[get_column_letter(c1 + j)].width
            if w is not None:
                d_ws.row_dimensions[dst_r0 + j].height = w
        for i in range(height):
            h = ws.row_dimensions[r1 + i].height
            if h is not None:
                d_ws.column_dimensions[get_column_letter(dst_c0 + i)].width = h
    restore_merges(d_ws, saved)
    wb.save(DST)
    print(f'saved {DST} ({height}x{width} -> {width}x{height} at {anchor})')

    # verification: transposed values equal source values at mirrored coords
    out, outv = load_dual(DST)
    ows = out[dst_sheet]
    problems = []
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            exp = ws.cell(row=rr, column=cc).value
            if cc in FORMULA_COLS and isinstance(exp, str) and exp.startswith('='):
                exp = clean_num(ev[rr].get(cc, exp))
            got = ows.cell(row=dst_r0 + (cc - c1), column=dst_c0 + (rr - r1)).value
            if isinstance(exp, float):
                exp = round(exp, 9)
            if isinstance(got, float):
                got = round(got, 9)
            if got != exp:
                problems.append(f'transposed value mismatch at '
                                f'{dst_r0 + (cc - c1)},{dst_c0 + (rr - r1)}: '
                                f'{got!r} != {exp!r}')
    if dst_sheet == SHEET:
        problems += verify_workbook_state(ws, ows, min(r1, dst_r0),
                                          max(r2, dst_r0 + width - 1),
                                          min(c1, dst_c0),
                                          max(c2, dst_c0 + height - 1),
                                          formula_cols=FORMULA_COLS, issues=[])
    try:
        verify_manifest(SRC, DST)
    except Exception as e:
        problems.append(f'manifest: {e}')
    if problems:
        print('VERIFICATION FAILED:')
        for p in problems[:30]:
            print(' -', p)
        raise SystemExit(1)
    print('verification passed: transposed block matches the source matrix')


if __name__ == '__main__':
    main()